import re,csv,json
from io import BytesIO,StringIO
from datetime import datetime,date
from openpyxl import load_workbook

ALIASES={
    "ticker":["ticker","symbol","underlying","underlying symbol","security","security symbol"],
    "expiry":["expiry","expiration","expiration date","maturity","exp date"],
    "option_type":["option type","put/call","put call","call/put","cp","type"],
    "strike":["strike","strike price"],
    "quantity":["quantity","qty","position","contracts","contract quantity","net quantity"],
    "multiplier":["multiplier","contract multiplier"],
}

def norm(s):
    return re.sub(r"[^a-z0-9]+"," ",str(s).strip().lower()).strip()

def option_security(ticker,expiry,typ,strike):
    y,m,d=expiry.split("-")
    cp="C" if typ.upper().startswith("C") else "P"
    si=int(round(float(strike)*1000))
    return f"{ticker.upper()} {y[-2:]}{m}{d}{cp}{si:08d} Equity"

def detect_map(headers):
    nm={norm(h):h for h in headers}
    out={}
    for field,aliases in ALIASES.items():
        for a in aliases:
            if norm(a) in nm:
                out[field]=nm[norm(a)]
                break
    return out

def parse_upload(data,filename):
    if filename.lower().endswith(".csv"):
        txt=data.decode("utf-8-sig",errors="replace")
        reader=csv.DictReader(StringIO(txt))
        rows=list(reader)
        return list(reader.fieldnames or []),rows
    wb=load_workbook(BytesIO(data),read_only=True,data_only=True)
    ws=wb.active
    vals=list(ws.iter_rows(values_only=True))
    if not vals:return [],[]
    headers=[str(x).strip() if x is not None else "" for x in vals[0]]
    rows=[]
    for row in vals[1:]:
        if not any(x is not None and str(x).strip() for x in row):continue
        rows.append({headers[i]:row[i] for i in range(min(len(headers),len(row)))})
    return headers,rows

def parse_date(v):
    if v is None:return None
    if isinstance(v,datetime):return v.date().isoformat()
    if isinstance(v,date):return v.isoformat()
    s=str(v).strip()
    for fmt in ("%Y-%m-%d","%m/%d/%Y","%m/%d/%y","%d/%m/%Y","%d-%b-%Y","%b %d %Y"):
        try:return datetime.strptime(s,fmt).date().isoformat()
        except Exception:pass
    if re.match(r"^\d{4}-\d{2}-\d{2}",s):return s[:10]
    return None

def parse_type(v):
    s=str(v).strip().upper()
    if s.startswith("P"):return "P"
    if s.startswith("C"):return "C"
    return None

def normalized_import(rows,mapping):
    out=[];errors=[]
    for idx,r in enumerate(rows,start=2):
        try:
            t=str(r.get(mapping["ticker"],"")).strip().upper()
            e=parse_date(r.get(mapping["expiry"]))
            typ=parse_type(r.get(mapping["type"]))
            k=float(r.get(mapping["strike"]))
            q=float(str(r.get(mapping["qty"])).replace(",",""))
            m=100.0
            if mapping.get("mult") and r.get(mapping["mult"]) not in (None,""):
                m=float(r.get(mapping["mult"]))
            if not t or not e or not typ: raise ValueError("missing ticker/expiry/type")
            out.append({"ticker":t,"expiry":e,"option_type":typ,"strike":k,"quantity":q,"multiplier":m})
        except Exception as ex:
            errors.append(f"row {idx}: {ex}")
    return out,errors

def keypos(x):
    return (x["ticker"],x["expiry"],x["option_type"],round(float(x["strike"]),6))
