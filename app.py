from __future__ import annotations
import os,json,hmac,html,uuid,re,math
from datetime import datetime,timezone,date
from zoneinfo import ZoneInfo
from urllib.parse import quote,unquote
from io import BytesIO
from openpyxl import load_workbook,Workbook
from pypdf import PdfReader,PdfWriter
from fastapi import FastAPI,Request,Form,UploadFile,File,HTTPException
from fastapi.responses import HTMLResponse,RedirectResponse,Response

from db import SessionLocal,User,Setting,Position,Audit,Snapshot,ImportJob,DailyReturnBaseline,DailyReturnPosition,DailyReturnExposureBaseline,DailyReturnExposurePosition,get_setting,set_setting,audit,config_bump,utcnow
from security import ADMIN_USER,ensure_admin,hash_password,verify_password,sign_session,read_session
from helpers import option_security,detect_map,parse_upload,parse_nirvana_exposure_rows,normalized_import,keypos
from daily_return import parse_nirvana_pnl_pdf,calculate_mark_to_market_return,auto_refresh_allowed,ny_time_label


PNL_OPTION_RE=re.compile(
    r"^(?P<ticker>[A-Z0-9.\-]+)\s+(?P<expiry>\d{1,2}/\d{1,2}/\d{2})\s+"
    r"(?P<strike>[0-9.]+)\s+(?P<type>PUT|CALL)$",
    re.I,
)

def _pnl_xlsx_num(v):
    if v is None or v == "":
        return None
    if isinstance(v,(int,float)) and not isinstance(v,bool):
        try:return float(v)
        except Exception:return None
    s=str(v).strip()
    if not s:return None
    neg=s.startswith("(") and s.endswith(")")
    if neg:s=s[1:-1]
    s=s.replace("$","").replace(",","").replace("%","").strip()
    try:
        x=float(s)
        return -x if neg else x
    except Exception:
        return None

def _parse_nirvana_pnl_xlsx(data:bytes,filename:str="PNL.xlsx"):
    """Parse Nirvana PNL Report exported as XLSX.

    Baseline rules are identical to the existing PDF parser:
      Previous Mark = Excel Price column exactly
      Quantity      = Excel Quantity column
      NAV           = Excel Grand Total Market Value

    This parser does not change current-market Bid/Ask/Last logic.
    """
    wb=load_workbook(BytesIO(data),read_only=True,data_only=True)

    required={"security name","price","market value"}
    ws=None; vals=None; header_i=None; headers=None

    # Prefer a sheet whose name resembles PNL, but scan every sheet safely.
    sheets=sorted(
        wb.worksheets,
        key=lambda s: (0 if "pnl" in str(s.title or "").lower() else 1)
    )

    for candidate in sheets:
        candidate_vals=list(candidate.iter_rows(values_only=True))
        for i,row in enumerate(candidate_vals[:80]):
            norm=[str(x or "").strip() for x in (row or [])]
            lower={x.lower() for x in norm if x}
            if required.issubset(lower):
                ws=candidate;vals=candidate_vals;header_i=i;headers=norm
                break
        if header_i is not None:
            break

    if ws is None or vals is None or header_i is None:
        raise ValueError("Could not locate the PNL table header (Security Name / Price / Market Value).")

    hm={h.lower():i for i,h in enumerate(headers) if h}

    # Find Report Date / Run Date from metadata rows above the table.
    report_date=""
    run_date=""
    for row in vals[:header_i]:
        line=" ".join(str(x or "") for x in row if x not in (None,""))
        m=re.search(r"Report Date:\s*(\d{1,2}/\d{1,2}/\d{4})",line,re.I)
        if m and not report_date:
            report_date=datetime.strptime(m.group(1),"%m/%d/%Y").date().isoformat()
        m2=re.search(r"Run Date:\s*(.+?)(?:\s{2,}|$)",line,re.I)
        if m2 and not run_date:
            run_date=m2.group(1).strip()

    positions=[]
    ignored=[]
    nav=None

    for row in vals[header_i+1:]:
        if not row:
            continue
        try:name=str(row[hm["security name"]] or "").strip()
        except Exception:continue
        if not name:
            continue

        upper=name.upper().strip()

        if upper.startswith("GRAND TOTAL"):
            try: nav=_pnl_xlsx_num(row[hm["market value"]])
            except Exception: nav=None
            # Some exports may place Grand Total NAV in another numeric cell.
            if nav is None:
                nums=[_pnl_xlsx_num(x) for x in row]
                nums=[x for x in nums if x is not None]
                if nums: nav=nums[0]
            continue

        if upper.startswith("PAGE "):
            continue
        if upper in {"CLOSED","FGTXX","MPFXX","USD","CASH"}:
            continue

        try: price=_pnl_xlsx_num(row[hm["price"]])
        except Exception: price=None
        try: qty=_pnl_xlsx_num(row[hm["quantity"]]) if "quantity" in hm else None
        except Exception: qty=None
        try: market_value=_pnl_xlsx_num(row[hm["market value"]])
        except Exception: market_value=None

        if price is None:
            ignored.append(name)
            continue

        om=PNL_OPTION_RE.match(upper)
        if om:
            # Some Nirvana PNL Excel exports omit the Quantity column entirely.
            # For options, Market Value = Price × Quantity × 100, so recover the
            # exact quantity when Quantity is absent.
            if qty is None and market_value is not None and price != 0:
                qty=market_value/(price*100.0)
            if qty is None:
                ignored.append(name)
                continue
            expiry=datetime.strptime(om.group("expiry"),"%m/%d/%y").date().isoformat()
            typ="P" if om.group("type").upper()=="PUT" else "C"
            strike=float(om.group("strike"))
            positions.append({
                "security_name":name,
                "instrument_type":"OPTION",
                "ticker":om.group("ticker").upper(),
                "expiry":expiry,
                "option_type":typ,
                "strike":strike,
                "quantity":float(qty),
                "multiplier":100.0,
                "baseline_price":float(price),
                "display_price":float(price),
                "baseline_market_value":float(market_value if market_value is not None else price*qty*100.0),
            })
            continue

        # Same aliases used by the existing PDF parser.
        if upper.startswith("ISHARES 20+ YEAR TREASURY BOND"):
            ticker="TLT"
        elif upper.startswith("ZOOMINFO TECHNOLOGIES INC"):
            ticker="GTM"
        elif upper.startswith("POOL CORP"):
            ticker="POOL"
        else:
            ticker=upper.split()[0]

        if re.fullmatch(r"[A-Z][A-Z0-9.\-]{0,14}",ticker):
            # Equity Market Value = Price × Quantity.
            if qty is None and market_value is not None and price != 0:
                qty=market_value/price
            if qty is None:
                ignored.append(name)
                continue
            positions.append({
                "security_name":name,
                "instrument_type":"EQUITY",
                "ticker":ticker,
                "expiry":"",
                "option_type":"",
                "strike":None,
                "quantity":float(qty),
                "multiplier":1.0,
                "baseline_price":float(price),
                "display_price":float(price),
                "baseline_market_value":float(market_value if market_value is not None else price*qty),
            })
        else:
            ignored.append(name)

    if nav is None or nav<=0:
        raise ValueError("Could not read Grand Total / NAV from the PNL Excel report.")
    if not positions:
        raise ValueError("No live option/equity positions were parsed from the PNL Excel report.")

    return {
        "report_date":report_date,
        "run_date":run_date,
        "nav_usd":float(nav),
        "positions":positions,
        "ignored":ignored,
        "source_format":"XLSX",
    }

def _parse_pnl_upload(data:bytes,filename:str):
    lower=str(filename or "").lower()
    if lower.endswith(".pdf"):
        return parse_nirvana_pnl_pdf(data)
    if lower.endswith((".xlsx",".xlsm")):
        return _parse_nirvana_pnl_xlsx(data,filename)
    raise ValueError("Unsupported file type. Please upload PDF or XLSX.")

COLLECTOR_TOKEN=os.getenv("COLLECTOR_TOKEN","")
STALE_AFTER_SECONDS=max(int(os.getenv("STALE_AFTER_SECONDS","150")),150)
app=FastAPI(title="RPD Fund Management v6.15.1 — Automatic Expiration",docs_url=None,redoc_url=None)

def init_data():
    ensure_admin()
    with SessionLocal() as db:
        defaults={"portfolio_name":"RPD Fortress Fund","nav_usd":"46865138.06795","limit_pct":"15.0","config_version":"1","config_updated_at":utcnow().isoformat(),"opportunity_config_version":"1","opportunity_limit_pct":""}
        for k,v in defaults.items():
            if db.get(Setting,k) is None: db.add(Setting(key=k,value=v))
        if db.query(Position).count()==0:
            seeds=[
                ("CG","2026-11-20","P",35,-530),("CG","2026-11-20","P",37.5,-550),
                ("COMP","2026-10-16","P",8,-4400),("GTM","2026-10-16","P",3,-11660),
                ("GTM","2026-10-02","P",3,-11667),("MNDY","2026-11-20","P",65,-1090),
                ("POOL","2026-10-16","P",155,-230),("QXO","2026-11-20","P",10,-2279),
                ("REAL","2026-10-16","P",7.5,-4700)]
            for t,e,typ,k,q in seeds:
                db.add(Position(ticker=t,expiry=e,option_type=typ,strike=k,quantity=q,multiplier=100,
                    bloomberg_security=option_security(t,e,typ,k),underlying_security=f"{t} US Equity",active=True,source="seed"))
        db.commit()
init_data()

def current_user(request):
    username=read_session(request.cookies.get("session",""))
    if not username:return None
    with SessionLocal() as db:
        u=db.query(User).filter_by(username=username,active=True).first()
        return None if not u else {"id":u.id,"username":u.username,"role":u.role}

def require_user(request):
    u=current_user(request)
    if not u:raise HTTPException(401,"Login required")
    return u

def require_admin(request):
    u=require_user(request)
    if u["username"]!=ADMIN_USER or u["role"]!="admin":raise HTTPException(403,"Admin only")
    return u

def page(title,body,user=None,refresh=None):
    def nav_class(key):
        t=str(title or "").lower()
        active=False
        if key=="overview": active=(t=="overview")
        elif key=="fortress_bloomberg": active=(t=="delta monitor")
        elif key=="fortress_daily": active=(t=="bloomberg daily return")
        elif key=="fortress_yahoo": active=(t in {"yahoo daily return","estimated daily return"})
        elif key=="opp_bloomberg": active=(t=="opportunity bloomberg delta monitor")
        elif key=="opp_daily": active=(t=="opportunity bloomberg daily return")
        elif key=="opp_yahoo": active=(t=="opportunity yahoo daily return")
        elif key=="admin": active=t in {"admin","fortress setup","opportunity setup","users","audit","nirvana import","nirvana import preview","map import","import preview","daily return upload"}
        return "active" if active else ""

    admin_titles={"admin","fortress setup","opportunity setup","users","audit","nirvana import","nirvana import preview","map import","import preview","daily return upload"}
    is_admin_page=bool(user and user.get("username")==ADMIN_USER and str(title or "").lower() in admin_titles)

    if user:
        adminlink=(f'<a href="/admin/fortress" class="nav-short {nav_class("admin")}"><span class="nav-single">Admin</span></a>' if user["username"]==ADMIN_USER else "")
        display_user=html.escape(str(user["username"]).capitalize())
        auth=(
            f'<a href="/overview" class="nav-two nav-overview {nav_class("overview")}"><span>Overview</span><strong aria-hidden="true">&nbsp;</strong></a>'
            f'<a href="/opportunity" class="nav-two {nav_class("opp_bloomberg")}"><span>Opportunity Bloomberg</span><strong>Delta</strong></a>'
            f'<a href="/opportunity/bloomberg-daily-return" class="nav-two {nav_class("opp_daily")}"><span>Opportunity Bloomberg</span><strong>Daily Return</strong></a>'
            f'<a href="/opportunity/yahoo-daily-return" class="nav-two {nav_class("opp_yahoo")}"><span>Opportunity Yahoo</span><strong>Daily Return</strong></a>'
            f'<a href="/" class="nav-two {nav_class("fortress_bloomberg")}"><span>Fortress Bloomberg</span><strong>Delta</strong></a>'
            f'<a href="/bloomberg-daily-return" class="nav-two {nav_class("fortress_daily")}"><span>Fortress Bloomberg</span><strong>Daily Return</strong></a>'
            f'<a href="/daily-return" class="nav-two {nav_class("fortress_yahoo")}"><span>Fortress Yahoo</span><strong>Daily Return</strong></a>'
            f'{adminlink}<span class="nav-user">{display_user}</span><a href="/logout" class="nav-short"><span class="nav-single">Logout</span></a>'
        )
    else:
        auth=""

    admin_sidebar=""
    if is_admin_page:
        t=str(title or "").lower()
        def sactive(*names): return " active" if t in names else ""
        admin_sidebar=(
          '<aside class="admin-side">'
          '<div class="admin-side-title"><span>Admin</span></div>'
          '<a class="side-main" href="/overview">Dashboard / Overview</a>'
          '<div class="side-group"><div class="side-group-title">Fortress Fund</div>'
          f'<a href="/admin/fortress" class="{sactive("fortress setup","nirvana import","nirvana import preview","map import","import preview","daily return upload")}">PNL / Exposure Upload</a>'
          '<a href="/">Delta Positions</a><a href="/bloomberg-daily-return">Bloomberg Daily Return</a><a href="/daily-return">Yahoo Daily Return</a></div>'
          '<div class="side-group"><div class="side-group-title">Opportunity Fund</div>'
          f'<a href="/admin/opportunity" class="{sactive("opportunity setup")}">PNL / Exposure Upload</a>'
          '<a href="/opportunity">Delta Positions</a><a href="/opportunity/bloomberg-daily-return">Bloomberg Daily Return</a><a href="/opportunity/yahoo-daily-return">Yahoo Daily Return</a></div>'
          '<div class="side-group side-tools">'
          f'<a href="/admin/users" class="{sactive("users")}">User Management</a>'
          f'<a href="/admin/audit" class="{sactive("audit")}">Logs</a></div></aside>'
        )

    rf=f'<meta http-equiv="refresh" content="{refresh}">' if refresh else ""
    body_shell=(f'<div class="admin-layout">{admin_sidebar}<main class="admin-content">{body}</main></div>' if is_admin_page else body)
    return HTMLResponse(f"""<!doctype html><html><head><meta charset="utf-8">{rf}<meta name="viewport" content="width=device-width,initial-scale=1">
<title>{html.escape(title)}</title><style>
*{{box-sizing:border-box}}html{{background:#f4f7fb}}body{{font-family:Segoe UI,Arial,sans-serif;margin:0;background:#f4f7fb;color:#15223a}}.wrap{{max-width:1780px;margin:auto;padding:24px 28px 42px}}
.topbar{{background:#142b4d;color:#fff;box-shadow:0 2px 8px rgba(13,33,60,.18);position:relative;z-index:5}}.topinner{{max-width:1780px;margin:auto;padding:0 18px;display:flex;align-items:stretch;min-height:78px;gap:10px}}
.brand-logo{{display:flex;align-items:center;justify-content:center;flex:0 0 auto;text-decoration:none;margin:8px 12px 8px 0;background:#fff;border-radius:11px;padding:4px 8px;border:1px solid rgba(255,255,255,.55);overflow:hidden}}.brand-logo img{{width:188px;height:auto;display:block}}
nav{{display:flex;align-items:stretch;gap:2px;min-width:0;flex:1;overflow-x:auto;scrollbar-width:thin;scrollbar-color:#7f91aa transparent}}nav::-webkit-scrollbar{{height:4px}}nav::-webkit-scrollbar-thumb{{background:#71849f;border-radius:9px}}
nav a{{color:#fff;text-decoration:none;border-bottom:3px solid transparent;transition:background .15s ease,border-color .15s ease,color .15s ease}}nav a:hover{{background:rgba(255,255,255,.075)}}nav a.active{{background:#1d4574;border-bottom-color:#9fc5f5}}
.nav-two{{min-width:126px;padding:13px 8px 10px;display:flex;flex-direction:column;justify-content:center;align-items:center;text-align:center;white-space:nowrap;line-height:1.08}}.nav-two span{{font-size:11.5px;font-weight:500}}.nav-two strong{{font-size:12.5px;margin-top:5px;font-weight:700}}
.nav-short{{min-width:61px;padding:0 10px;display:flex;align-items:center;justify-content:center;font-size:12px;font-weight:600;white-space:nowrap}}.nav-overview{{min-width:84px;padding-left:12px;padding-right:12px}}.nav-single{{display:block}}.nav-user{{min-width:52px;padding:0 9px;display:flex;align-items:center;justify-content:center;font-size:12px;color:#dce7f5;white-space:nowrap}}
.top-clock{{display:flex;align-items:center;gap:9px;border-left:1px solid rgba(255,255,255,.22);padding-left:14px;margin:9px 34px 9px 4px;white-space:nowrap;flex:0 0 auto}}.top-clock-icon{{font-size:19px;line-height:1}}.top-clock-text{{display:flex;flex-direction:column;line-height:1.15}}.top-clock-label{{font-size:10.5px;color:#b9c8dc;font-weight:600}}.top-clock-value{{font-size:11.5px;color:#fff;font-weight:700;margin-top:3px;font-variant-numeric:tabular-nums}}
.card{{background:#fff;border:1px solid #dce4ef;border-radius:13px;padding:16px 18px;margin-bottom:14px;box-shadow:0 1px 2px rgba(15,39,70,.025)}}.grid{{display:grid;grid-template-columns:repeat(5,minmax(0,1fr));gap:12px;margin-bottom:16px}}.grid>.card{{margin-bottom:0;min-height:115px}}
.metric{{position:relative;overflow:hidden}}.metric:before{{content:"";position:absolute;left:0;top:0;bottom:0;width:4px;background:#d6e4f8}}.grid>.metric:nth-child(1){{background:linear-gradient(135deg,#f6fffb,#fff)}}.grid>.metric:nth-child(1):before{{background:#2aa66a}}.grid>.metric:nth-child(2){{background:linear-gradient(135deg,#fbf9ff,#fff)}}.grid>.metric:nth-child(2):before{{background:#7457d9}}.grid>.metric:nth-child(3){{background:linear-gradient(135deg,#fffafa,#fff)}}.grid>.metric:nth-child(3):before{{background:#e15f55}}.grid>.metric:nth-child(4){{background:linear-gradient(135deg,#f8fbff,#fff)}}.grid>.metric:nth-child(4):before{{background:#498ee7}}.grid>.metric:nth-child(5){{background:linear-gradient(135deg,#fffcf6,#fff)}}.grid>.metric:nth-child(5):before{{background:#e3a126}}
.metric .v{{font-size:25px;font-weight:750;margin-top:7px;letter-spacing:-.35px;font-variant-numeric:tabular-nums}}.muted{{color:#718099;font-size:12.5px;line-height:1.35}}.ok{{color:#087a4b}}.bad{{color:#c22a20}}.notice{{padding:11px 13px;background:#fff8e8;border:1px solid #f4d790;border-radius:9px;color:#684c00}}
.card:has(table){{overflow-x:auto;padding-top:10px}}table{{width:100%;border-collapse:separate;border-spacing:0;min-width:760px}}thead th{{background:#f6f9fd;color:#182642;font-weight:700;border-top:1px solid #e7ecf3}}th,td{{border-bottom:1px solid #e8edf4;padding:10px 11px;text-align:right;font-size:12.5px;white-space:nowrap;font-variant-numeric:tabular-nums}}th:first-child,td:first-child{{text-align:left}}tbody tr:hover td{{background:#f9fbfe}}tbody tr:last-child td{{border-bottom:0}}
input,select,button{{padding:8px 10px;border:1px solid #c8d3e2;border-radius:7px;font:inherit}}input:focus,select:focus{{outline:2px solid #a9c9ef;outline-offset:1px}}button{{background:#142b4d;color:white;cursor:pointer;border-color:#142b4d;font-weight:600}}button:hover{{background:#1e4473}}.danger{{background:#b42318;border-color:#b42318}}.secondary{{background:#52657d;border-color:#52657d}}.row{{display:flex;gap:9px;align-items:end;flex-wrap:wrap}}label{{display:flex;flex-direction:column;font-size:12px;color:#52657d;gap:5px}}h1{{font-size:31px;letter-spacing:-.45px;margin:3px 0 23px;color:#17233a}}h2{{color:#17233a}}h3{{color:#17233a}}
form[style*="display:flex"]{{align-items:center!important}}form[style*="display:flex"] input[type="number"]{{min-width:76px;background:#fbfdff}}
.admin-layout{{display:grid;grid-template-columns:235px minmax(0,1fr);gap:20px;align-items:start}}.admin-side{{background:linear-gradient(180deg,#102e52,#0d2746);color:#fff;border-radius:14px;padding:16px 10px;position:sticky;top:16px;box-shadow:0 5px 16px rgba(17,43,77,.10)}}.admin-side-title{{font-size:18px;font-weight:750;padding:8px 11px 14px;border-bottom:1px solid rgba(255,255,255,.16);margin-bottom:7px;display:flex;gap:9px;align-items:center}}.admin-side a{{display:block;color:#e9f1fb;text-decoration:none;padding:9px 11px;border-radius:8px;font-size:12.5px;margin:2px 0}}.admin-side a:hover{{background:rgba(255,255,255,.08)}}.admin-side a.active{{background:#2678dc;color:#fff;font-weight:700}}.side-main{{font-weight:650;margin-bottom:9px!important}}.side-group{{padding:7px 0;border-top:1px solid rgba(255,255,255,.09)}}.side-group-title{{font-size:12.5px;font-weight:750;color:#fff;padding:4px 11px 5px}}.side-group a{{padding-left:24px}}.side-tools a{{padding-left:11px}}.admin-content{{min-width:0}}
@media(max-width:1450px){{.brand-logo img{{width:170px}}.nav-two{{min-width:116px;padding-left:6px;padding-right:6px}}.nav-two span{{font-size:10.7px}}.nav-two strong{{font-size:11.7px}}.top-clock{{padding-left:10px}}.grid{{grid-template-columns:repeat(5,minmax(160px,1fr));overflow-x:auto}}}}
@media(max-width:1050px){{.topinner{{flex-wrap:wrap;align-items:center;padding-bottom:7px}}.brand-logo{{margin-bottom:0}}nav{{order:3;flex-basis:100%;min-height:58px}}.top-clock{{margin-left:auto;border-left:0}}.admin-layout{{grid-template-columns:1fr}}.admin-side{{position:static;display:grid;grid-template-columns:repeat(2,minmax(0,1fr));gap:4px}}.admin-side-title{{grid-column:1/-1}}.side-group{{border-top:0}}}}
@media(max-width:760px){{.wrap{{padding:16px 12px 30px}}.brand-logo img{{width:146px}}.top-clock-label{{display:none}}.top-clock-value{{font-size:10.5px}}.nav-two{{min-width:116px}}.grid{{grid-template-columns:1fr;overflow:visible}}.admin-side{{display:block}}h1{{font-size:26px}}}}

.report-export-bar{{display:flex;align-items:center;justify-content:flex-end;gap:8px;margin:-8px 0 12px;flex-wrap:wrap}}
.report-export-bar span{{color:#68758d;font-size:12px;font-weight:600}}
.export-btn{{display:inline-flex;align-items:center;justify-content:center;padding:7px 12px;border:1px solid #cbd8ea;border-radius:8px;background:#fff;color:#174f91!important;text-decoration:none;font-size:12px;font-weight:700}}
.export-btn:hover{{background:#eef5ff}}
@media(max-width:760px){{.report-export-bar{{justify-content:flex-start;margin:0 0 12px}}.export-btn{{min-height:40px;padding:8px 14px}}}}
.mobile-nav-toggle{{display:none}}
@media(max-width:760px){{
  html,body{{max-width:100%;overflow-x:hidden}}
  .topinner{{padding:8px 10px 10px;gap:8px;display:grid;grid-template-columns:minmax(0,1fr) auto;align-items:center;min-height:auto}}
  .brand-logo{{grid-column:1;grid-row:1;margin:0;justify-self:start;max-width:205px;padding:3px 7px;border-radius:10px}}
  .brand-logo img{{width:min(178px,42vw);max-width:100%;height:auto}}
  .top-clock{{grid-column:2;grid-row:1;margin:0!important;padding-left:10px;gap:7px;border-left:1px solid rgba(255,255,255,.22);justify-self:end;max-width:190px}}
  .top-clock-icon{{font-size:17px}}
  .top-clock-label{{display:block;font-size:9.5px}}
  .top-clock-value{{font-size:10px;white-space:normal;line-height:1.2;text-align:left}}
  .mobile-nav-toggle{{grid-column:1/-1;grid-row:2;display:grid;grid-template-columns:28px auto;grid-template-rows:repeat(3,4px);column-gap:12px;row-gap:3px;width:100%;height:52px;border:0;border-top:1px solid rgba(255,255,255,.15);background:transparent;color:#fff;align-content:center;justify-content:start;padding:10px 12px;cursor:pointer}}
  .mobile-nav-toggle span{{display:block;width:28px;height:3px;background:#fff;border-radius:2px;grid-column:1}}
  .mobile-nav-toggle::after{{content:"Menu";font-weight:700;font-size:14px;grid-column:2;grid-row:1/4;align-self:center;margin-left:0}}.mobile-nav-toggle span:nth-child(1){{grid-row:1}}.mobile-nav-toggle span:nth-child(2){{grid-row:2}}.mobile-nav-toggle span:nth-child(3){{grid-row:3}}
  nav#primary-nav{{grid-column:1/-1;grid-row:3;display:none;flex-direction:column;overflow:visible;width:100%;gap:2px;padding-top:4px}}
  nav#primary-nav.mobile-open{{display:flex}}
  nav#primary-nav a,.nav-user{{width:100%;min-width:0!important;min-height:44px;border-bottom:1px solid rgba(255,255,255,.08);border-left:3px solid transparent;padding:9px 12px!important;justify-content:center}}
  nav#primary-nav a.active{{border-left-color:#9fc5f5;border-bottom-color:rgba(255,255,255,.08)}}
  .nav-two{{display:flex;flex-direction:row;gap:6px;white-space:normal;text-align:center;line-height:1.15}}
  .nav-two span,.nav-two strong{{font-size:13px;margin:0}}
  .nav-two strong[aria-hidden="true"]{{display:none}}
  .nav-overview{{min-width:0}}
  .nav-short,.nav-user{{font-size:13px}}
  .wrap{{padding:14px 10px 28px}}
  h1{{font-size:24px!important;line-height:1.15;margin-top:8px}}
  h2{{font-size:20px}}
  h3{{font-size:17px}}
  .grid{{grid-template-columns:1fr!important;gap:10px;overflow:visible}}
  .grid>.card{{min-height:0}}
  .card{{padding:14px 13px;border-radius:12px;margin-bottom:12px;max-width:100%}}
  table{{display:block;width:100%;overflow-x:auto;-webkit-overflow-scrolling:touch;white-space:nowrap}}
  th,td{{padding-left:10px!important;padding-right:10px!important}}
  input,select,button,textarea{{max-width:100%}}
  input,select,textarea{{font-size:16px}}
  form{{max-width:100%}}
  .admin-layout{{display:block}}
  .admin-side{{position:static;display:block;margin-bottom:12px;border-radius:12px;padding:10px}}
  .admin-side-title{{padding-bottom:10px}}
  .side-group{{padding:5px 0}}
  .admin-content{{width:100%;min-width:0}}
}}
@media(max-width:420px){{
  .topinner{{grid-template-columns:minmax(0,1fr) minmax(135px,auto);padding-left:8px;padding-right:8px}}
  .brand-logo{{max-width:168px}}
  .brand-logo img{{width:min(150px,40vw)}}
  .top-clock{{padding-left:8px;gap:5px;max-width:160px}}
  .top-clock-icon{{font-size:15px}}
  .top-clock-label{{font-size:9px}}
  .top-clock-value{{font-size:9.5px}}
  .wrap{{padding-left:8px;padding-right:8px}}
  .card{{padding-left:11px;padding-right:11px}}
  h1{{font-size:22px!important}}
}}

</style></head><body>
<div class="topbar"><div class="topinner"><a class="brand-logo" href="/overview"><img src="data:image/png;base64,iVBORw0KGgoAAAANSUhEUgAAAeIAAADKCAIAAAD2A9QbAABR3klEQVR42u1dd1wURxue3av0coAg1QJ2RWzB2HvEEgv2aCxJbEnUxJLks8QYNUWjibGXxF6iRhQ19g6ogA0VEJAqHQ7uuL7z/TF3y3JcWeAoJvP8cJ2bnZ2Zd3fm2XfeeWeWgBACDAwMDIyGChLfAgwMDAxM0xgYGBgYmKYxMDAwME1jYGBgYGCaxsDAwMAoB/e/KzqEEGogRUFAQUihnwBCCCkAKQpSAFKQogCAkNJAAAGktKcoCmh/Mi5BySAEgIIQlkdCClK6/AFF5w8hJABKR0dCSMdrE1MQQlRRbekAoJTaSHQ5pQugswBASBF0AoqO18Zoswd0zlBRlg8hIAgAISQAgACi+8O8VegsfeMIQDCTQQAJbTKCTqY9q82WgBACgs4Wooz08keZoKowk+lKJ4ChQg1VEhAEAQCBQgBVDACiPIY+SUAIKp4FgBlDELpkkCAIoE2hzYQgCAjQ3QDMs3R6vUIJQvuLkYwuC1TMlj4LBNYuBEnqxCK1/xMkAIAgSF1WJCQIJDXjFCAIEuj9JLUB3S0iCYIgCBLSkQSJAhAQ6BQgCAKQqFBdAu2FgI4BjDBBVriK5KCzgCBIVB+CIAkOQJmQqAKMn4AgSA6quTYfgkOQZLn4mKZrD5RGBSk1pVFBSqXRqCClojQqqEFhNaRUGrUKUiqKUkONiqLUlEZJadSQUmmPlIrSqAGlpig1paGTqaBGTVFqADUUpYYUOqopDYrRQPSTUkNKjXgZv5wxMN56ECRJcgmSS5JcgsMhSS4gOOUxJIfQHbWRHB5JcgkOjyS5JIdHkDySwyVJHsE8ak9xSZJHcnm6GD5JckkuX3eWx9FeziNITh3JysZvGlJqjUqmVskotVytkmnQn1quUck1aplGJdeo5RT9U62g1HK1Sk6p5SisUSs0ajmlViBdDwMDA+Nf8abgcLhCkivgcIUcroDkCjlcAYdnxeEKOFwrUhsW0gEuz0oXI+TyrUiukMuz5vCszA4RDNB03uu7RVmxCmm+oixfIcmXS/MgpcaPBAMDA6M2wOEKBTauAlsXgbWLwMbFI2CwrXOTKmrTEKpVZRq1VolWq2SUSq5WySi1TK2iVWltQKOSqVVySiPXqOSUWqFWySikSmuU+GFgYGD8K0mW5Aq4PCuSKyC5Ai7XisMTklwhhyvk8IQcrhWHJ+RwhSTXisu34nCFSKHm8Kw5PCutfs2zsoDRo+aAkEJWaUqj1FmolZRGRVEqSq37SakhpaLUKopSQUpN0WGNWhujM20zrNIqpj0aQg3UqJn26HKDtUatnX3CwMD415gdCFJrhiY4BMklOVyC1FqotTEkl+BoAySHT3J4gOBo7dQkj+RwCZLH4fIIkmm55qN4kqMzT+sM0xwOn9DlQwfqQsz/1p4eyLtD59dBURQADHcOhr8HxfD6AABSlEbrNQEpCCCkNOWeG1o/EOSnQbtnUDr3DIgu0XqGUBoIIUHACmcZl0BIEeitxvDuMOryoT1VwccDXU1U8P0o9/qQlWYTtOtFBY8O2qfCvKuGYdcLFq4azEJN+JMwvE0AM1umI4c2D6OVLPfBqOg1QTC7OKjs46H7Ty+9vs+GNhmsmKySY0mlTComq5GjiJWdOwAE0wME6hw/CALJrfP6qOj+wXTh0HPw0F1LEgRBAD3fD50LRwWPDgP+HgTJ0Va0kuMHACRJkkDr/kESgCR0P1E8oTtq8yl3+SAIAv0k/4NeHwTeegkDAwOjIQMvb8HAwMDANI2BgYGBgWkaAwMDA9M0BgYGBgamaQwMDAwMTNMYGBgYmKYxMDAwMDBNY2BgYGCaxsDAwMDANI2BgYGBgWkaAwMDA9M0BgYGBgamaQwMDAxM0xgYGBgYmKYxMDAwMDBNY2BgYGCaxsDAwMDANI2BgYGBaRoDAwMDA9M0BgYGBgamaQwMDAxM0xgYGBgYmKYxMDAw/vPg4ltQ97j27M31uGwNBSufsrfieThZ87msXp9qDZVbIi+UKKCBnACPSzrbCER2AkdrnrOdQGQrdLLhuzkI61jYMoV697XE7GJZ5VMcknC1F4rsBCRBsMxNIldJFWqpXF2mVNNSkyRhxedY87nWAq6NgOtozfNwsnZzEHqLbCwujlSh3mMhceQqTU6xrESmYl86n0s62QicbflcDgkAIAjg62LrJbKuDUkxGg4IaLCLY9Qmeq+8cOtFdr2V3tp9RGefUV19mrjZ1UFxqXkSv3kn6kvYjn6id1u69Wnj/l6gl7XAAkrJ6zxJk/oTxwTcHa3cHKx8XWza+Ti393Xq3sINczemaYwaaWQPXuUBQAAACySKp6lF6QXS8zHp2WIZgCi6Lo6BfqL3u/jMHtSykYNVrcobl16cVyIDgKAgjE0pKChV3H+Vd/VpVp1JCiCwEfIm9Wg6d3CrQD/nGorzLL0ov0QBAKQgiEkuKJQqIhNyr8e9qVQuASA0VJ86im/l6TigfeOB7Rr3a9fYRoDHzZimMWoMmVL9/cnH3598BAgAIAQEoet4tR6e1idg5biOdaNc04hJzp+46UZClriO5R0S6LVucudAP5FlxYlOzp/4y/XE7JK6fNGyP/Zq5f5hX//Q4Ca2Qh7ua5imMWqET3bc2Xn5Zb0UvXhk+xWhHeuyG6fmSVp9fkKm1LC/pLWXo6uDFeKePLHseUZxtYVdPaGTkMexoDivc0ubzD3WwBvYxB7NpvcNGNjBE/c1TNMY1UROscx95gE2KYU8zqhuTXq0cm/l5UgAAIhysohNKYhKyI1KzK1q6R2biM7/7z13R+s6k/fzPfd+Pf+MTcr+7Tz//KyPp3MFe6tEropJzn+YlHchJv3Kk8wqFd3c3f6vxQM7WFSt/mzPvd/YiaN76zi19nb093DwcbFt6m6vN3UskakLJfKCUkWOuCw6Kf9hUl6xVGmRevq42K6e0Hla3wDc4zBNY1QHHb84+Sgl3+ww9p/lIYMCvUzkk5Rdsu9a/O4rL3KKZewHyB7O1v8sD2nn61w3wp6PSQv5/gKbul1dNaxfO0/T8q48+vDQ7UT2pgA7Ie/66uGdmrlaSpxzD1OHr71IvzK1ZRkKezhbR6x739e1aoamlNzS2OT8h0l5sSn50Un5eWIZm7KMhbv6u237pGdQUxfc6TBNY1QNUzdfO3Aj3tA0UYVw4u8Tm3s4mM1NKletOxX7/V8xxvKpHPZ0tonZEOpWy/OKCGl5Et9PDrKp2+ONoe1ZaL7/xKZP3Xw1t0TOUl5XB6uoH0Y3aWRvITNOqd/HB3WlmDouGdXxh6nv1LC4uPTC05Epp6NSYpLzzZZo7DhnSNv1H3Szt+bjrtdggZe3NDj4udkCAAEEAEAAofEwK9gIeWsmdU3ZPrljUxcWeUIAQWaBZNLGy3UjrI+rrdn6oDBL1+rBHb2jN4xt5+PE7h7CPLFsyOpzlhJHqx2jp6MtC1aM0cY7WIIW23g7/y+0U/TPY1N3TFn8fqCQR5ott3L8tovP/OcevvQoHXc9TNMYbOFozQeQApACEOoChsJVo367mA2hgzp4mslTF776OH3/9TqayXS24VlWXi+R7Z21o4KauJjPE1IAUgmZRT//HWu5x8c1UiLzD2op1XJvux+nBafu/ODzYW1Nl2swPrdYOnhV2NcHInHvwzSNwQpCPker7gGG6lc5XHWc/mpoV383U3kywkv/vNeg5K2Sdc7emn9tzfvt/ZzN3ENdeM3xB4USuUXEseJzGM/I6B+XQ1j8Tro5WG2a2TNz74eh3ZuxqYPe37qTD0euDccdENM0Bgva4nG0ag7FUHkqh6vO1NYCbvjyYT4uNmzyzy6UbAl/Uhfyckk29SGIqvGagzX//PLhNnyOqXuoC4ul8k1nHllEHAEtDl0KZaBELknU0v1s7GxzfMmQsK9D3B2sDEtNUcbiw6KSQlaH4T6IaRrDXD/ncRhWVBPH6sDF3mrTrF7s8oe/nXtUB/LyuSTL+lQ1Z0+R7aZZvVjlDOH2i5Z5J/E4BNL+AQAMhV0X1sWjTTlqD8O7Nnm5bcqMAa31yjVWHzr+/MOUYd9hpsY0jWESBGBlPq52/qPeadbNvxEbc3BCRkFMUm5ty8sh2MlbLYlnDWrzbksPNvnnFUsvxaZaQhzAZnBQa8o0czwh2PPpgKOL3xNwCfODM0Y4/H7yjM2XcE/ENI1h/JEgZymk45QrzpXCNcCKid3M5w8hAODvyFdvu7wbZvZilT+EZyKTLCEOQ101fuTUAU8DAAAY3zMg8ucJPq52bGpFH/ddidscFos7I6ZpDGOAWkWPguVKn4Fw9QsY2rlJc3d7c/lDAKm7cZm1LS3BUt7qolsL9wGB3izuJ4yKz7LwYMi40krUYXsKbOr66Ncp/dp5sVeoAaQW7LwWFf8G90ZM0xgmmboWbNM0JvZuyaaUyJdZdSFsLcs7vmcLNhbq6EQL7S4LaWM6NBauM20awclWeHVt6JBOfmzqRodD14UVSxW4N2KaxjDYySnzfzWze4zp7s+mlDK5oqRMWbviUhSbmtRkuezo7gGsbimkVGpNzeWpqKRTRgYH9bD69/T/RvZs42lkMGGgnum54kU7r+EeiWkawwhTG3RurRBfI3Ro6uZsKzSZv/Yvt1hau0YPgpW8RA20T2c7YTs/F3P3EwIA1RSs8aODDDsvrDBiYMQT9dGshHzu6f+938jRunJ9jNVz3+WnES+ycI/ENI1hRCOjoP6RGV9jvNPSw1T+uqNEVjfatBl5a7j5zDstPMzcTwoCaAH2JNit86xjowcNkb3V6eWjWK7PROHpG8/hHolpGkOfuMoV6sr2WQvZpgEATRo5mMxfe7Sr5U15IKwLef09nczlDwGEZI3ZE9a+qb2GCG7luWh0F5bu5ADA+PSCvZee4G6JaRqjIkuzsk3XFH6N7NkU5GhTu1+5regnXlvyNmnkwKYUC6w6oQwt9qt0JOq1ia2Z2tvd0ZpNPdHxh2P38E6amKYx9BQyNrbpmvabRo42bGzTIvva3dEUQlgHtninyob4SvI2crQmiZryJ6H1m2aafQ2E68vogWAl4P46dxCbeqK/hIyC03df4p6JaRqDSV20IgMZSg2sEF9jONsJTeVPUQBSHZu61X4ThGzkrSGrOdoKzdxPimrh5WyJR0dpHdogHTAUBvWsnYb2atWhiZv5eurCW848wP0S0zRGBauHpfabNsNc5vZ6DmzmXlcDiNqV195aYHbv6c4BHhYeDxk/EkT9t7P1M/uxt1DfeJySkV+COyemaQwAtIunmXPu0MhcfE0LUijVJvOHAFIDg5rUtrwEoFjJWzOBS8sU5u4n7NPe1zKDAxbWXi5Z/11vSJfmwa08WVqoIUUdvIInEjFNY2hpi6FTg8rGaMvYagEApTKFyfwhj0sMeyeggchbQ306q6DU3P2EvSxC0yRgs7mzBdbRWAIfhwSx35D64JXHuHtimsYAAAAuhzC7L7BFPD1KpHLT+xEPfyfAzlpQ+zQN2clbI57Oyi8xnf/gTk0dLOHTUr5FicmjUq1uCI1tQt92NgKu2dqiY1xKTnahBPdQTNMYgCAIA8bB8n2Bq7//sh6SswpN5z///W51IC+HZClvjZCaU2Q6//F921lSHG3mlY66eJWaagiNTcjnTujbzkQ99Y6XHr7CPRTTNAYgCWB4SRhFWWS/aRqvMgtM5N/Wz7VvYNN/jbyXHyaayN/RRjDBQjRNAp3rtEE7ry5e3TCMHgCA0N5tTNRT73g1GtN0fYyw8S1oaOCQJAAQAIZSZjhcU8SlZOvMjgbyXzmtf13JS7CRtyaeEblFkgcvM0zk//mY7lYCnsVsVqweUENZLjK4iz/7ykQ8T8U9FGvTGIAAQLtRmeljzSCRKR+/yjKW/+DOzcf2bltH8mr306hFef95kGAiZxd7qwVju1tMnnKF3WIfhq91pu7cnFWdIZWYnkfh9YiYpjG0y9gMfAnQkptC3HyUTFGUwfwdbYV/fBVaZ/Ia3QSjYnxNXD22nLpnIv9f5g13tLXYSkvd6kLzX8ZpOBjSNYB1nWFmnhh3Umz0+K8DGtO2oCVHzQcvReuXostv1+LR7s52dfdaYidvtdeDXIh6ed/gUB0CAEBIcKspgzpaUBwKWXXpAiA0GW4Q6NDcg32d03PF3m6OuJ9imv7PGz0MU7DFrNKFJWVHr8YazH/xxD5j+7Sv69cSC3mrTWtf7ThvLH93Z9uDyyfW1eOz5FvWsmjh7cq+Pm8K8FpETNNYm0a2WoLRnQmGeYCwQDf/fv+Vctpj5N+/k//6T0Lq/LXESt7qKdMbj918nJipf990+Z9aM82C5g7d46MamunZLBq72Ntb80ukcjaJWSbDsCCwbbqBcjWLb9ZVE8lZBRuP3qicZytft5NrPiTrfOc2CFl9lw9WXeTIuNQlW8MM5skliX82fBzc1q825GG8RfWWO4KKLskNCP5eLkY+NK5ff4kMfx0Ra9MYgGGrhUZGydXt5BQFx36zr3L+LXzcbv0+38FWWB/iUqzkraLISZn5I5bs0ui5J0MAAHC0tTq9fmafjs1rRfEhAKAohtufsWPD4ummHk7RL1LN1ZkAAEoxTWOaxqipRcMkZq47EpuQrhfZwsftzvbPXRxs/jXyxiZkDFqwNV9s4CuO7Zs1PvPDR34ezrUkCUVRRr4xCCy+K4sFoTP+VK6nfv1LsdED0zQGSRCAonT2UwJAaDCsqvqmEJ9u/OuP8Ai9fDo0b3z51/n1ydFokZs5ednnty88ct7PJ2RyZeV8Phrx7s5lE2rbiFPRNg3r9kVcTTjZWRkxqVdyMHrbLO+YpjFqSbWEABIAIHIxHP71+M1tS8azzDE1u3Dez8fD78bp5dMnyP/Mjx/b2wjrUVrCuIzMMBuHvDO3nvxv57lnSVlAuy9KeT4B3o22fBE6sGvLOnp89Tdgqj5Ns6uSjZCPuyim6f88STMnCY17sG4/devMrUcjerZv09SjuZerp6tjM09XG6sKXSgmPr1EKv/75uPNx65XUPcAAACO6Rv417qPGoLAbOQ1dnVuUWlGbvGV+y93/H0nOTNfT0YAQCvfRl9M6j9zRPc6FIfSX9Vf2czbwOBka6Ud0xisJyPGWsjDnRTTNAZbVetNvnjH6dvVK2DKkK4HVn3YQN5LVVU/xRLZ6GU7rz2MN5HaXWQ/tl/QmL6BfYIC6lIa7Y4stC4PGHo989jAqJrP4xipp36MrZUA909M01ibros9HyKeJElkiobR5djJy6A1CkJvN0e9q2ytBO4ie+9Gzv27tBzavW3HFt71NzhgrOgzdmxgCrVCqWK8L00dsdED0zSGrqtrVRmilsJJmbkTvtl1dsM8ot6/yoeMHmbrzByh21n/sWLaHyumFZZI41Nz3JzsGrs6WmqLuxqC4elhgQFTnUGuVLGskrvIHnfQOgZe3tIgOZrFt+lqfgy//XjFjrAGIC67b/EZMk8729sEt2vazMu1gXB0FR5fA4NCqWLZbHwaOeE+irXp/zq0X28xt/+yq6NtgK+7QqlSqNQKpVqhUiuUKhQokyvN7VWtDa/Zc65TS5/3+3SsT4EhrJv9tetI8SF0E5im7QcNTKgSiax8baTJo4+7M+6kmKb/8zQNABsr4T9bFnZs4WMskwKxZH/4vW9+PyVTqEznM+V/ux4eXN7Sz6OejR5sGOJtgG57WMB4zRgKNzCB3uQXM9xjCGP1b9LYhc/DpIGNHv95UJAysDlypaOdtSlnZ5GD7cJJgw5//7HZfKQy+bAFm8USWf0r1KaPbwlTk8jd2/CX4CtvkdFQkFtYwqbO9TYxi2kao+Fp02yMm+b7+ft9ghZOHGg2t6S07AnLttafvKyMufU/1VmFwQGb6YEGpk3nFbGxTffv2gr3UEzTGICiIDt1jBXWfxbaumljc2odvHjvyVe/nag3KwELeeFb8m0nDaV5GwcH8a/fmK8zhP27tsY9FNM0hm4VovaPMhpmBz6P++fqj0zlowuv33v27+vR9SEvZUF56//xUczKGz02qLfOm/xiaZncbJ3dnGxb+LrjHoppGgOwUWqqpIt1bt3k65nD2eT5wTfbUjLz6sNKYL5uxFv0+MxuNt3AVOmXKW/YbDY9sk8Q7pyYpjF0Rg9LGze/nx/atlljs3lKpLIxizbVubyaavtNN0htmlltY48SUg3JdTouKd1YPZkxU4e9i7snpmkMndEDsPmrGk79stBayDebbezLlE9W7/4XyFt/b1lWjjoNiqZvRb80W2ffxqIeHVvg7olpGgPRFgXK/6DRcBWZy9/Hfdv/ppvJE1IAUjv/urLv7xt1x2saDTt53xKariCO0b8GNTi4GvnUbIXnjhuA+yamaQyapullbLQ6YyhcdUwd3mt0/66m8tSFZ6zYfjvmZd3KbEZegnjLHqPpY8Oh6RfJmYXiUtO1FQp409/vjfsmpmkMLTQaDYAQoLGziWO1sPvbT9yd7dnkP+qzH1Oz6mI6Ua1Ss6kPRb0dCrVarSpXQinKWFhd9Y/v1BJOX71vop60Ku3qhHdcwjSNQdN0BT9io3/VM2462dscWDefTf4FxaXD5q2rg++TUpCVvG/NFKJ2HADKP2gAQcUPikMAgEbTUGzTf565YaKeAEIrAX/ZrPdxx8Q0jcGgLYpiZ9ysZv4Dgtt/MW0YmyKeJaSOW/Rzrcv7FhpzTb1l1ZoKzhJG9v/TaDQNobb3n75KSMk0UU8AqTnjBmJVGtM0ht6oWWPGKg0hAEBdg36+dsHkwJZ+bCzC529GL1i/t5blVbORt+Gon2bE0ajZDA5UDcPo8duh86br6eJkt3zOWNwrMU1jVFTHNBrT7qs19yPm87inNi+xEfLNuvcCSG3+8+yOY//UMk2bl1dDad6Kx6dWqcu9UwwPjCCgKJWq/mk69kXywbDrpuv5w6IpjnY2uFdimsaoAJVaXWkZmN7qNQgAVKtrRFtNvBrtXD3XRP7Mn7NXbb0d/bwWaZqFvG/NFKJGXT4OYH5vt4KXDlQ2AJqe/90u0/Xs0q75jNH9cZfENI2hD7lCaWhTBUovpuaj5knDeoUODjaWv1782M/WZeUW1oa8SvQVPnPyKlWqt+PxyRXmvXQgLNL6wNUb9vx1+V7McxM1tBHyD/64EPdHTNMYBlD+8VCTR7Ul5qB2r/nU213EZheR3ILi4bNXK5SW50qlipW8SqX6rXh8SpWKjW1aKpPXYyWT07M/W7PTdA33/7AgwK8x7o+YpjEM0bRCyWaPC4sYN+1trU9sXsbyI3gxzxInf/FTLWjT7ORVvx3atLhEwsZxRVpWn99hmPLlz2UymYnqLZg6fPTAYNwZMU1jGIakTMZm/2WZ3DIezd06tPhtxWwWnxqBAMCT/9yd9c1mC8srZSmv8u14fiy8VgCEhcUl9VXBkXO+i3j00kTdBvcI+uXrj3BPxDSNYRQymYLN/styhcVoa/6U4ZOG9Qbs9n3ec/zimt+PWKpoORo6sChXoXwLaFpcKtWvP2VYrpy8onqpYchHK8OuRBi4z7p6tm7mfXzzMtwNMU1jmEJRSSkbW21xicSChe5dv6hj66as9rmGcPmmP7ccCLOMsGIJy/21xaXSt4OmDXylxYA4b/IK6rpdiSV9Ji85f+O+8a/JwKZejS7u/c7e1hp3Q0zTGKZQUiphY6u1LE0L+Lxzu1a7uziwtFN/umrLtkNnLcNr7EoUW1TeWkJxiZSNYRpAqqi4pKAO7R4vktKCRs69GfXYRJVaNvG8c3SDt4cr7oOYpjHMdnUJG1eBjDcW3hepsZvo3O41QgGP3e7PcO6KX3/983QNCy0oLmFZXEZ23tvx7NiNSACAkbEv6qZW3/56oPWgWa8zsk3UZ3DPTneP/+Lh5ow7IKZpDPPIyy9i48iclpVj8aI7tfX/6/cVpstlHj//dsuaLQdrUmJ+oZil43ZqRk7Df3YFxWKzXxSkj+u3Ha3t+hw4fcW3x6RVm/ebrsm6xTMu7lvr7GiHex+maQxWiEt8XWn+XbdjGSM+OS2rNkoP6dvt4C9fmShXL375xn2frvqt2sU9T3xtOn86Pik1s+E/u+eJqboKmz/eefh09JxVlvLYYUJSJttx+FzT3lOmfrE+LSvXRB06tfWPDtu6bPYE3O8aMrj4FjQoPH6RpN1L2kCnAsyYZy+Ta6kOk0cOKCuTf/zVBoPlVq7Plj9O5eQVHv99ZTXKin4az1LeF4mvpWVyG2thQ3580U/jq/QR9NMXb7V/kbRpxbyQfu/UvPQSSdmVO9HHzl0/Hn7DbGJnB7u1S2Z9Mmk47nQNHwR8ez5f9F/Amt8OLN+wh2XiSwd+Htizcy3V5Jc9JxZ99zv79E19PH7+Zu6owT2rVIp9m6Gl0jK2pLZzzfuDejTkx2fXZqhEKgP0h2YgYBkWOdp/MHrQ3A9G+vt5VanE7LzCW/efRMY+vxn5KCYukU1ZVgLB/Gmjls2ZhK0cmKYxqgy5QtlmwFT21oyuga2izuyovfr8tOPIkrXbqnTJoF5ddqxb7Oflzibx3uPnZy5ezz7z9q2aPb64r8E+vl1Hzn381U9aWoQAEASAUBcGLOOdHO28Pdy8PFw93ETurs5cDkevFA1FFRaXFBaXFhaXvExKS8vKYZ+/lYA//8PRiz+Z4OrsiLsbpmmMaoxYpR8t/fH4uWsA6LoZAGbDQ3p327l+iXdjt1qq1bV7MR8uWpuelcOyPigc0i94xviho4eY+nrezchHw2csLZWWVUnevsFBu39c2tSnwe01cSMydviMr3QrKomGdmzm23jG+JCPJgxzFWGCxjSNwRpPXiTl5BemZmSnZ+XEJ6dfu/swr0DM6Fo6etLvcgbiO7b1D2rbwsvDlcfjdmjV3NbG2lXk2CagiaWqeuLctR+2HYp+8pJlfVC8g53t0H7B73ZpF9DUR6VS5xUU21gL07Jykl5nxiWk3IyMZZlP5fj2rZp1bt/Su7Ebj8dt26Kpg50tAKBLh1Z1ZrkulZRFP43PLxKnZeYkpWY+i0++FfW4Ys2ZFod6i7e3tZn0/sAPQ4d269ga9zhM0xhVVL4iYvqO+7RWu2tWdJiHm8iCdb54I/LY2asPH798Fp9c9foYO1pM3r7dg64d+61uHl/f8Z/eiIgt/82sGhP1FN+xjf97fYMH9era+51A3NcwTWNUH1GxcQVFYoOnhAK+u6uIxzPghwMhLC6RFIlLTX9Mz8/Lo7XltGk9qFTqjDe5OfmF7PcPEpdIS6RSbw8DxhmCIJwc7Jwd7QmCMFZcXmGx6S3lSJLs3L6lS12ZXBNT0o+FXUV7yYqc7Jv7eVkJBWavyisozsrJz80vysrJz84ryM0vyi0oysjKraFVw8nBrrmfl6e7a3M/r07tWg7s1UXk5ID7F6ZpDAwMi0GhVEnLZBKpTFomk8pk0jJ55U/zkCTB43J5PC6Xw+Hzeehoa23l7GTP42K3WkzTGBgYGBj1B7wKEQMDAwPTNAYGBgYGpmkMDAwMTNMYGBgYGHUNPEdcP4h8mnQnJoGiIADA0d66fYB3UEtfPu9f+zievcq4FZMgkWo/p83lcqaNeFfkYPs2yiKWyG7HxD9PqrCmv3Wzxv27trYS8nHbxrA4sKdHvWHk55vjkjKf/LXm5sOXn/1wkMfl3tv/P0c78983in+dLSmTB7XyNeZozESJVGYtFHA5VRs2lUrlz15lUBDa21i18/eyiLxLNx3fc/pW1MEVuYUl32w5mV9U+s/2Lz1cHNlcm19UGp+abWct9PEQGbxFUpniaWKGhqKshfzAFj5s7szrrPz84tKgln4kqU2soainiRl21sJm3mYW3xeIJS1GLOse6P/3L5+RJBH/OnvXqRt/nLnzTvtmB77/xMne6EOEEL5IeaNUqQNb+DDjlSr144R0pUrN5XCCWvnyuJzaaHIUBYtKpCJHW9z73i5wVq1ahe9CveD0tei8YsnCKYP9fRo5O9ievPKwuU8jxIlxSZkjP9/8zW8nb8cmHLkYtf3E9ekjewIAMnKKZqzcU1wqPX01evfpWx8M615cWvbb0SsjP9/815UHZ289Phh+j6JghwBvupS1e85BCjb1Kuedx/Fp/T764ac/L0hlimnLd12Net69Q3Oa+yRl8hVbT89fd6CwRFomU1y49/SjVfvsbISd2zR5mpgxeM5PK7ae7tWppVcjJwBAUnrujJV75q874O7i0KEi71TG7ZiEB3EpH4/p0z7Au01zz23Hrzdytu/aruntmIR3pqzef/bum3zxxKXbniSm9wpqYa1TSzNyij5csTv2ZWpxSdnFe0+Xbjr+MC6lR8cAGyvtchKFUrV+b/isVfvyikrL5Mrr9198vHqfUqXpGRTwIiUr9Mvfv9x49M6jxCMXozYfvjxrdG8AQE6BeMbKPTmFJWE3H/16+PL093sCAG5Fxy/ddLxEIt906FJqVn6fzi1NyAIh3HTwUpc2TUf06QgAcHG0HRTc1sHOev/ZuxGPX00c+g5JEBk5RROWbvv8x0N+ni5tmnkCABLTcmav+bNEIjt6IfLUlejQQV0Qde48eXPK19uTM/JkCmXMi9dzv9//JCF9ZN+g7Hzx9BV75n6/39HeulNrPwDAmRuxPaZ9H3YzNr9YsmZnWNqbgm7tmnM4ZNTT5K6TV+06datP55Zuzvbad0mxZO73+2d9u5fL5XRt15QkiNIy+bTluwe/285KoL29hSXSfX/fnrhse/9urd2c7VOz8ueu3f/Jd39cvf/8xOUHvx25Mqh7WwdbK9xb6xkQo54wfsnWliOXofCThHRRr3nH/omiz075ekeT975E4cgnSSgw7X+7fjnwDwqfv/0YBSKevBL1mrfv79sQQo2GinqaRGeiUKr83vtiytc79Ioe+fnmfh/9ACEskciahSwe/tkmFC9XKAd8/GO7Md/kFZbQiZ8kpP92+DIKj/1iS8CIpcysVm47Leo1L7+o1Ky8K7aeEvWal5SeCyHMLyoV9Zr3+7Gr6FRg6PIPl++CEL7OynftM3/+ugN0/YMmrNzx13U6k/Tsws4TV3WeuEqhVEMI1RrN6IW/Nh36ZeqbfDpNckbu97vCUHjWqj1ufT5FYfrOzFu7f83OMOZt1GiogOFL7z9LRpHnbj0yLYtUphD1mvfx6n2Vn6mo17wTl+6jn//bclLUa15xaRl92/efvaP3+L7YcMSl93zmUysqkS76+QgK//jHeVGveUzpfId8MWvVHgjh9QcvRL3m/XHmNooPGr+Cfo40dp+6Keo172liOvpZIJaIes3rOG5FckYuneZFclbvGevonz/sDadLfJaYIZMrcVetd+ApxPobyJAkbW9KzsxzsLUa+E7bCm9Q3cZD3do1RYGiEumRi5GlUjkA4L0e7VEkyRjgkyTRtW1T+uexf6L6dG55OeJZdn6FVekEAZCxy85G2Lm1n0qlRvG7T92KfZm6et5oF6fynYjb+XvRZekVR//kshikM/fkfPn6DZ/HDWFki+DrIfJr7KLWrYPfcvRqRnbh5KHBdAKvRk5r5o9OyczbfuIaAODoxaib0fHffDTcx71895Imnq7jB3fTSVrewuk7U1giPXH5fnFpGX0bKQjFkrKdJ2+gBCE9O7CRRVVpreDHY/oAAO7EJmhL1yYm6cf3Z9hduUJFlxv9/PUfZ+7MGtWL+dQc7aznjutXMQcO89kh9ApqweNy0rILGbUy3J1JUhuvVmsaiRwyc4sGz/459mUqiuRxOQI+t5L2BgAAbZp7CgU83FXrHZim6+/Wk4RKpb77KPHq/ec//XH+4NoKNk0uh1Qo1XcfJd59lLj371socvyQbknpuQM++TEuqfyLU8iOmZCaffdR4skrD1+mvKFPHQqP3Pr1VAc7692nb+q9IeixVEZu0dIZIejn5cg4AEBw+2Z6VaVttXweVypT/LjvPP13/2kysuqal5cgkFHlZnT8ut3nDq+f7dvYRXtKxy9KlaZEKv9y6nvo56mrD91dHGj7Bk1PAIB7jxIBANeingMA3mnf3FiFeVwOBeGd2IS7jxJ3/KVl4XGDumbkFPX/+MdH8Wn03R47sMvpq9ETlm4TS2Tmp3QI7d3Ti2/u4wYAyC4o0ZNL+/gGd30Un/bevA2vs/KZN/ydDkbrz6nEvPSzS88p5HI5YwZ0plNqNJSxxFqa1lDNvd12rZxeVFo29ostyRl5AAAej8M0haN57AfPku8+SvztyBXcTzFN/9ehUmvuxiZuPXo1I6eINinqKUcvU948iEtBMZPee+ebj4anZuUPmbvh4t2nOsogUC+VK1UnLj+gdfC7jxKDWvlaCfnjBnX948wdmVzJYBkiM7fox33npy3f1b9ba9oOixQ9a+O7CBEEEPC573b0p/88GzkZJCxjsLUW9u7U4tyWhUzjLwFAXFLmj/vOh365ZcmH79EkJZMrmdVGsBLyneytM3KLAABlChUAwNbaVIVJgiAIIik99/7TJBT5ft+g7+aNzikQD5234dTVh1oj/mdjR/XvdCUyru/M9UnpudV8oCoNAMDNyU7ProgCs0P7Lfpg8LNXmf1m/XA7JgEAUCZXAgBszPmHMG8vekM8Tcy49yjxzh/ftG7amNZ/Kz8D7TcldZer1BqlSj2iT8ddK6fLFKrRi35Nyczjc7mVtXUOSWbni9FbBAPT9H8XGoqythIsmT70wNpPbKwE4xb/XlRSpjeyfjfQf+aoXt/MGkFHLpwy+O9Nn5MEsfCnwxX0Ly+3/l1b71k1o7GrE4rZdfIm0nwlZfLi0rIjFyOZ3d7TzWlk36A3eeK9f9+mdXMfDxEAgB4OG+ILIODz3g30p/+01gYWLE0Zp3IIQOumnu/1aP86K/+nPy7QJppm3m6FJVI99ZaioFSmCPB1BwD4uDsDAKKfp5oyLnHIdwP9pw5/d/W80XTknHH9wjYvsBYKFm88htwi7W2sdq2YvvazsWnZBev3njMji5HRw9NXGQCAft1aGxs/fT1r+JEf5ihU6mWbj9P1j3mZWpXJJK0lauJ77/h6iCobQwy9qwja6KHWUOhFdf73RTKFauAnPz1NTOfzOHqPqVMbvzEDOm/7Zirup5im/9OgdENUayF/6zdT07ILpy3fRZs7EXcgeDVykpTJNRT18HkKAOCd9s0WT3uvuLQMdTnmUNfGSut7l5VX/CIla/PSyUumD/1l8aQmnq6/H7tGa1VKtYYgiBZ+7rtXzXCwtZq5am+BWAIACB3YGQCwYf9Fvaqi0TEAQKVW642sUa+mWGjTSDSDerdarSEI0M7f64/vZsmVqo9W70NKdOjArgCALRWH3vefJStVmknvvYPMFwCAXw9f0rO60BXWUBT9DvF0c5KUyQEAaHQS1Mr361nDSsvkCqXqcXyaUqUGAHw8ps/g7u3yiyVsZKlsm9518kbLJh7DegXSctE3p0QqS0jNBgAM6NZ6zri+qIjhvQOFAt6e07f0XkWMG67Ru71KlZrZNhhvfci8t1KZovxy3c3RUJRCqULhwBY+l7YvFvB505bvIhi7VjOfL33HMDBN/0chKVOgvgQA6N2pxQ8LQu89Spy39gCyPMgUSolMnpqVDwCIf5398PnrMply+/HrtNFzYHAbxMho4ExrxPvP3gUA7Dx5Y9rwdxnKY9/UrPwzN2LpPoyK9vUQndzwaU6BeO6a/aVS+cDgtlu/mfoiOWvV9r+VunnFm9HxV6Li9C5kSCEHAJRVjDQir5ymDz1IZQqZQgUA6NjS98/vPrr/NHnppuNKlSZ0UJfv5o3+7cjlg+ERdCZLNx9fMn3ogHfaAAA6tfY7tO6TnIKSJb8cp80jD5+n/H09BoXLZEqlSoP4MTUrP+LxK7WG+vXwZe1tJMjenVpaCflv8sVHL0bRaung7u1My4KkeJ2VTzNmbmHJnO//tBLwj/80j57KQ8nQsbi0bN/ft3VGHmJI93YAgEYih7DNC5ztbT7+dl+B7t2Qkpm378xtZg5lMiX98pYpVAbvoUyuzMorRtPLmblF1x+8oC+X6NKrVBrm+8DXQ3Rm02d2NlZMGzp6EHGvMtHdptsMRj0C+03XD+JeZd57nOjqZG9va+XdyJnH5XRs6dvYzelBXMr9Z8n2tlYxL1LdnO0v3H165GLU5ci4aSN62NtY3Yx+eSs64UFcSmZu0Y8Lxgv4XKVKc+7WI5WGKhBLjlyMOnIxqqmna0mZ/Pztxy6Odu8G+qMBr7OD7bNXmbmFJb4eLsWlZZFPXlkL+R6ujn6NXVwcbd8N9L8Tm/gkMaNV08bBHZpPHdHjeVLWH2fuPE/O+vXw5eSM3KnD3nWyt0lMy7n76JXI0c7aStDcuxFJEtn54qtRz+1trfk8bjNvNyHfqFdA6puC6w9eONrbAAA8GzkxXXHvP0t+kfJGwOc283Zzd3HwbezStrnn9QcvUrPyWzVp3Ltzy85tmhy/dP/CnSdxSZlPEtM/nzzo/b5B9OXNvRvNeL9n6pv8Xaduxqe8+f3YtUcv0z4Y1t3N2R4V6upsfyki7sjFqAt3n04eGuxsb3P30atr95/HPH+dlJH38xcTrAQ8giD2n737PCnzStTzJp6un04cYNLiAf++Hl1aprAS8h8npMW+SL37ODE5I29E746fTx5kZ6P90FdOgfhy5HNHexsel9PSz4PH5V64+yTqWXLE41dlcsV388Yg9xgPV8cPR/ZUqTU7T954npy19+9blyOfjxvU1bexS0Gx5FLEM1sbKy6HbOnnIeBzr0Y9z8orFgh4fo1dGrs60lWKepocl5Rpb2N16mr0kYtR4bcfjxnQmc/jXrr3TCjg8zicAD93Dkmevh6TXyxxdbZv6uWKLnR2sB3cvW2pVNY90B8AkFdUeikiztZa+PD56yMXo05ejR7WqwPTiwajXoBXIWJgYGBgowcGBgYGBqZpDAwMDEzTGBgYGBiYpjEwMDAwGMD7TTccQAghgBSEEAAIIQUghJACADAjIYSEdr8sSnsJgABSAABIUQAAZjwsj4cApdSWgv6ha3UBAAHUXl6+fg1ASNFnIdAlQgXpctCm1C69gGhBHAQQot1DdPXRnYJQG82og26fERSmr4D0nSEAwTyriwcEvflJ+SkCAJShdtUHAYjyIEDrPQiUTrcshCAIAtVCGwaAAARaGQJ1CXQrRcrPajPTpdQFkCOyLqz7qVtmgi4kdKUTBEFq05eXTzKuIkF5JGCmRzF0YhQg6DwBIAgSMuJ1WZH0WYIgICBQQFthgiQIEurOAkDgnolp2hQojQpCDaTUlEYFKW2AotQoDCl15TBFqRk/1RSlgZQaQAihBkIKUhRAAUhBSheAFIAUqBSJwuWnIAWhBlK6S3QpAaAqR9Ip0Vk6B9zgMP5VIAiC4CBmJ0htAJAkQXDQq4IRqUtGkARZfhYwInUJOIAgAEESBIcgSToShQHBIUj01uEQBElHkiSHILkkh0uQXILkEgSHIDkkqf1JkhySw6PD2gCHp01DcOlrG+htrplDHtSoFBSlpNQKjVqhPWoUlEalUSugRqlRKyiNkqIDaiX6icIajZLSKLSRlArFqGTFkFJjRsPAwKgfTiQ5BMERWDsTJI/k8jkcPsHhkRw+h8MnuXySg/4EKMzhCkgOn+TwOFwBofvJCAs4XAHJ5XM4ApIrsCRNQ0otLU5TK0qV8hK1QqJSlGqUEpWiVK2UqBUSlbJUrZCoFBKNUqpWleGHioGBgcHKdsG34QpseXw7Dt+GJ7DjCmy5fFuewI4nsOPybXlCe67Alie0t3H0NW/0IEguT2CP7FkkySFIjobLIzg8koveDAIuz5rDs1YrpVylVK0qUyul+AFgYGBgGCdoWy7fBtE0l2fN5dtw+bZcgS2Xb8MT2HJ4tlyBLU9gyxPYcfi2Fjd6AIDmd5BZg1Jq1AqoZlg5dH8adbmtg2nigNoEKkqtpCiVRq2ElJLSqJQyMaTUAGq01mdsA8HAwKglKwdBEiSXIDkEyeUL7AkOj+TwkClDZ+LgkRw+h8snSB7JFdBmEPTH4dIBAUkbRnTWD3SqpjV8SxaLQ0hpdPOH2mlDquIsIkWpoUZNUWoAKfoUQJOKWq8JDYQUoCgINQBASFHaGAgpqEHuEJDSaMtCV2nfEFBvapGeb9R5X+gmDCtMQmqQq4bBqUVUB0DhqUWMt4PJDE0VcnQ+JIwJQMasoN70oN6sICNZeSTJ4dGn6EK1KUkSAFI3bcjROaVoJxIBoZ0wJDlcZFwmSK5uClF3SveT4HBJgktTM5sPHNfz7cd7erxt0HnU6Xz1gNY3rtwbD+ic8HTee5TuOkYa5N1W8RRgONuVe9ExnPZ0rnuggpsdgDpfOmDAM0+bEtDedbpfkOFvB5i+dwBoXfK0Iuic+rT+djrPvYqeebrMAaF170PueLQPHNB+W6zcG0/nbVbuqIdc5SDtX0f76jFPMR3pCAAICCBB+9jRfnWEnh9euSMgHQ8AQZLoFEH7Beo+iEUwi9D56ZHaCpO6AGJJXSYEw72P6ZxHu/3RCSBjE2qMt+AtiWkaAwMDoyEDr0LEwMDAwDSNgYGBgYFpGgMDAwPTNAYGBgYGpmkMDAwMDEzTGBgYGJimMTAwMDAwTWNgYGBgmsbAwMDAwDSNgYGBgYFpGgMDAwPTNAYGBgYGpmkMDAwMTNMYGBgYGJimMTAwMDAwTWNgYGBgmsbAwMDAwDSNgYGBgWkaAwMDAwPTNAYGBgYGpmkMDAwMTNMYGGwRERHxVtf/9evX+CFWAw8fPqj3OoSHn/sX3EmusROJiQnnzp2jKKryKYIg9AIymSwhIb5Dh8CFCxfpJVYoFIcPH8rLy6MvIUmSvlCtVvfs2Ss4OFjvqrS0tEOHDvL5fL2CKIpq2rTZ6NGjTYi0ceOGqKjIY8dOsLwFq1d/e+zYUU9PT09PL09PT5FIBCFENZ84cZKfn1/lS5BQBQUFdMVQJQmCQNcy4eTk5KOFr0AgqMunu3HjhsTExG3btlfj2jlzZnfu3HnmzFk1rMPBgwfOng0LDj5R7RwyMtK7dw9OS8tg/1aIjKzwYkAPRaPR9OjRs3JjMw2pVNqrV48NGzaGho6z1HMRi8V37949fz48NfV1aWlpSUlpaWkJh8NxdhaJRM4ikUgkchGJRC4uLi4uriKRqHfv3iZyKywsDAxs7+zszOfzhQwIBAKhUCgUWjEjs7KyMjLS2feOmjShbdu2de/+mGUTWrdu7bFjRwmCFJqDRqMxxjaVsX//nwCAkJBhLOs8atTI1NQ0T8/GiA1sbGxQ1xaJRCEhw0QikV76W7duTpkyOTg42FEHgUDIpiClUpmd/ebNmze+vr4//vhzNWna3z8A3YWVK5fv27ePTcGjR4+pHCkQCKZPnwEA2Ldv78qVK5infvzxpwkTJhrMysfH56uvvs7ISF+5cuXly5d0xDH3q6++ZtlRJ0wYf/ToMTaJV6xY2a9fv7KyshUrll+9ekXXwrabeLS0UEePHlmyZDH7htu5c5cBAwZ06dKlS5eudUDTaWmp4eHnjh49Yuw+G8O1a9fCw8/Z29vVsAJlZWVbt27NyEgPDz/HvqtUkiIdAODn5/P6dRqb9MHBwa1bt05LS926dSutTC1cuGjx4iXVe80AAI4fP2YRmo6IiAgPPxcefq6goADFBAQEdOvWzd/fPz8/Pz8/LzHx1ePHj/WuOnnylIkG4+zsjN5hGzdu2LTpF7N1cHd3Z1/hlJSUajehy5cvubiIWKb/6quvv/rq64iIiE2bNrIZfg0dGsImW42GmjNntunuzMQff+y/c+d2QkLCL79sRDEhIcNMvKVevHhRw/GiUGie1rlmU3z77XenT58uLi5GP0NDx40dG6o3JPznnwvXrl3z82tiIp/p02ccOnQwISGBbp1mH7yXl/eCBQtomv7ww+nsHowGAHDv3t2pU6fs33+QzSU9evQEAKSnp3/77Sr0YFg+1AkTJu7evYsWytbW9tChI15eXq6urnS3fPLk8bVrV9GDfPjwARoJBgcHr1q1ulWrVnUwWj969GhV+9jRo4cBAMnJyTWswLlzZzMy0gEAZ8+erTZNazRqNJDy92+WmJjE5hIHB4d27drPmTMX0bRIJGKjeRnE+fPhAICoqKhbt2726tW7Jnfj66+/QqQPABgzZuyAAQMM3pOYmJgbN64zCdfa2ppN/osWfXH2bFhSkvYWLViwMDi4e8WXxL2wsDONGjViX2fUBuqsCQUHB2dkhDJZT0/xz8hIv3Pn9unTp729fdgprQo0NGTJ1A4ODiEhw0JCwJkzf6PKz50710T6ly9fIgLx8vJ0d/ewt7e3tbW1s7OztraJjY2hH2JwcPDcufMkEolYLBaLxTKZDAAQFRUZERHh4OBgAZoGALRq1Yq+cZ6ennrDxuDg4IkTJ44fHyoSOZvOhzlkcHNj1VbatWtPhz08PNhcolarUeDGjRszZ87Ys2cvyybi4dG4SgXRgtA03bFjUMeOHfVuTnBw8CefzI6IiDh16uSxY0dp+h48eOCmTZsNDkEsS9MxMdGnT58eNWoUy6tu37518eJFpEnVnKZpsouPj2/RokW1aJqibU2tW7d6/vwFyws9PRvrPdmq4tatm7GxsSh89uzZmtD0+PHl7PPdd2umTfvQWMqgoKCgoKDBgwfPmDH9zZs3AAArK2vWrdGNpmnU/PRaY/v27emHwqYrJScn1XET8vLy1quzHpOHho7Ly8tn+bJRqVS0EY+9To2YB9G0j4+vaZo2ZhCLjY1h/uzdu0/lNCdOHK88eKoMS04hOjuL2Cdmo+pXD0x7+uXLl2bP/oR1/bWvGUdHx+oJYkKo4ODgn376We+dsWDB54cPH6ql+5CenlZYWKhTSY6wv/DIEW3i3NxcqVRa7QpER0ffuHGD/lntyRyK0tBhiaQ0MLA9ywtFIhc0eVBt6w1SpRHCws6kp6dXL5/169fRHL1t23YTHE2jTZu2UVEPPD092WvTbNC+fQf2/fTVq1d0A6iXJmQMfn5+bm5ubFIqFEo6PGfObPaNsEmTJgAAkiRNa7txcc9qYg0LDR1nZ2e+cVqMpo8dO2FlZcU+fZUSV1H50uj1tM8//5TNhTweVxfgV08Qs0INHDhIbyps2bKltTQZzVRk7t27d/78eTZXRUZGMrWtlJTq2z30tLazZ8/qPZrqPdDCwsIuXTqxvJbD4QAASJJTjXJzc3PDw8tpWiaThYWdqUY+x48f27r1d93j/qpKxp/vv19rWZp2c3ObOJGt+QKp0vXYhIxh+vTptra27LRpJfNnlZiaTcthP7NtDO+9916datNVQm1q0xr6ZYhw+vTpL7/8wrwBiMtDAT6fVz1BWL579GYkfvllo0RSavH7kJqaWvE9ykobQiZFg1xfJZSUlKD+QA9ak5JeVe+FpFbrk3tOTk737qwcNrhcLgCAw6lOOz9/PlwsFjMH19Wg6fj4eLrttW3bbtasj6p0eb9+/S1L0wAAf/8A1jSdXI9NyCIiKJVKvZgqMXX1Wk5Vxzf/RZpGpsxhw4a7uLgwNZqvvlpmrktzdNp07dJ0SMiwyZOn0D8TEhJ27txp8fug5+17/fr1a9eumjNTPDx16pRF+ti5c2ezs7MBACtXfltDuwdt9Jg7dx4dmZGR3qdPLxbdjEvr1NWzeEydOq1Tp84o5sWLF1UVgcnsISEhtI8peyxcuAi9bOoeejRdx03IIkC26a1bt1WPqavXciyOfyFNoylEuVy+b9+fTMI9dOjgypXL2WjTVTJ6VIOmAQDz53/q5ORE/7x580Yt0XT//gMYao4ZbQiZFJlOEdV29jh79iwAYNq0D1u3bk1LeuHChefPn1fb6LF06bIpUz5g1m3QoAFsXr3VMHpERkZGRkYimp46dVr1FGq5XM5MP2TIkGrcyWr7qFiCppMqacp114QsAoVCAQBwdHRcv/4HPab+55+LNTd6YJqukdFDLpd36NDhzz/3M0/t27dvzZrvTA+Qa2L0YL96xdPTc8CAgfTP2NjYrKys2qDpXr3KVc6LFy/eu3fXWPonTx4fP34MADBixAhGJtVRhaKiou7evQMAmDFjJn2stkJNe3qo1eq1a9eNHPk+ferly5chIUNrw+iBVOkOHTo4ODgwPRwuXLgQF/eMZSZ3795hmp6aNWsO3ioghj1x4mTdNyHLatNqtWbSpMlff/0N89RHH826du1avRs96o6mIyIi/vqrasvMam80gXq1XC4HAPTo0XP79h3Mszt37vjhh/WmabpKRo9qC8LUcwEAL1++sOBNUKvVqHv4+voy3cgOHz5s7BJ0asaMGZ6eXrTTS/VGrGgGqV279miGgOnwHh5+jvaYrKrRA3W5337b0rdvX/rs06dPRo0aaY6mq/aMxGIxounx4yegGOZiCvYKdVpaGmhgiI2NpRdumEZ6enpxcTFBEH5+fky/2LppQiZQXFzMUgSgs00j1/vZs+fMn1/Bm+DDD6fevn3rX2j0OH8+fPz40Mp/MTExDaQhojEyomnUwTZt2sxM8PvvWzZu3GCSpvl1UE9f3wr+mLm5uRZVpVMQqfn4+DDt4GFhZ6Kjoyunf/HiBRrPfvrpZ0Kh0Ntb67taVFRUVFRUpaILCgoQTU+ePBnFODg4tGzZklbQkD2kqlYsZuDPPw/Q9mIAQHR09LhxY03SdNVsu+Hh4ehx0DYW5uKOsLCzEomk4dN05a7avXvwyJEV5mxMIDExEQDg7e3dqFEj5hKPOmhCTFSmmvbt2x44sJ9ly0HuuSqVtuUsWbJUzxty8uRJyLrVkI0eVZ6aoJdy6GHQoMENhqbVAAC0zgdh9Ogxcrl82bKldAxaHbRoUQX3D1qJrpLRo9rQ2x8gPz/f4sNVAICvr1/z5v7BwcG03+7hw4c6depUSQ86RFHU6NGjRSIXAICXlzftdZ+SksI0o7NRpdFK6EmTJtORs2Z9RDs8hIefY79Qgmn0oJcqAABOn/57wIB+dGuMjIwcPz608lYVSBuqqk6EVGmmp1SfPn0cHR3RWtzMzIyzZ8+ycWvLyMiox45grKuyXCyelPQK6BabhIQMa9eu/dOnT+qmCemN1A2JwGoBGt1gECcgfPfdmpIS8enTp+mYcePGnj0b3qFDh3+P0WPBgoVpaRnMvw0bNoKqmGVrn6YppjaNMGnS5G+/Xc2M2bTpFz2dmu7MVTJ6VBt6So1lHckRTYtEIiQLkzFPnDgeFxen1yGPHDmM9CAUQ6tC1bAtIlWaWSIAYNy48XT40qV/2Jt3mUYPPWvJlSvXmIwTERExceJ4g9o0SVahncfGxt66dRMAoLcpAvMnS7uHntOuWRw+fCgkZOi4cWNHjx4VGjqGqUKGho4JDR3D0nPZWFdF02gsF9m+epXIbAn02KgOmlDFEUkFEcLDz/fq1ZulCLQ3Hq1NI2ze/NvgwRWmc4cPDzH4Vvv3eHqEho6r9nYNLC0YaOBcPaMHjenTZ+jt3KTH1DQ7V9voUdlJ0wRKS0srKtcuFrxvqGP4+Gj3PRg58n2mlUBv6eOhQ4eUSuWQIUPoOS60+E1PMWeDu3fvREVF6fVqBOY2b1Wye9BtQK1W6Z26f/8h2sNMV/rdGTOmV+xmVbZNI1Xa3t5+4MBBxmj67t07bHbbcXSsmgo5adLk8PDzx4//9cEHU6OioiJ06NbtnRMnTp44cXLo0KE1aRWTJk3u3bs3y6Xz6LnTLWHSpMkVLdS11YRMo1279jNnzmJJ0/RrsnLL2bVrt96+gwMG9KtspPpXeXr07NmTzQJrpsLI0rpHJ3N1dWNZGdrTo/KpOXPmLliwUI+pf/99i96bs9pGD4OFGkNOTk5Fmna24HNF8zbM7QgmTJhAh48cOUzv/JCWlob0oNmz59AJmKpQlaaAaP5ds+Y7PZMi2p6CtnsgT6kq0bSeTqQzicYzt5O9cuXyvHlzGdp01YweMpkMrTxs06aNXv2//XYlMyUbhZqeRkNA21ywwcCBFaaXAwM7WqphBAV1qrwbpxGjR7Ie2zJfvbXXhMyiS5fO9P6CpkGvFDfYcg4cOKS3YciQIYPoLZffPqOHQSH13tJsNntjKhf0dhOsaZqtsokWrRljzEWLvtBzRP3hh/X79u2tqE1XgaaZS5mrRNOZmeWGSwcHh549e1nwuSJvPFqbBgCMHz+hffv2tPWA1oYOHToolUr79+8fFFRubfT09KpGH8vOzqZXCUdUAnNQmZqayt4zz4Q2rSOUFOYjO3s2bPHiL5naNHujx4UL59GWfhGGoEfTzBePQehtVs7GURfB1tauegYxNl2VTT5oV1W9ljBp0uS2bdvVahNieXNY7n9N26aNtZxjx04wmVoikYwYMZypQVrK6GH2uViApqvEPiyVi+JiVpO/9C1r2ZLtnp8mtGmEhQsX6TH1ypUrjh49Wj1PD2ZBcrmM/YXMnbEsu0+eVCpFDMKkaQDABx9MZWpDmZmZb968QZ2NeUpPh2K/J0N4+LmSkhIAQFjYOT2TIvpjblLD3u5hWpsGAHC53KdP45hcduzYUbSUqaraNFKlx4wZa7D+aWkZ9KC+tLTUrEKtZwwMCzuDDL61B7NdleXGcrSm7OXlxYzXG5NZvAlZFrQR0oQPqB5TZ2ZmhIaOodNbyuhRQwqtU5pmsgbLDxfRNM1+a2ZjtmnTTL1kyZdnz4bpjB7Vpukq3CjmPodjx461qMUjWXfDK/j8jR8/oXPnLvSNPXz40OHDh8RicY8ePdHeEQyFxZbeT1IqlbJ0FkSqdPfu3QMDAw0mGDZsOB2+evXKkyeP2WRLb3loTCcCAFhbW0dFPWBux7Nv375169ZWyTYdFxeHNjc34bY0ePBg9m8a9MkPpnKnt5C67mmaJdDrhMfjMdkWADBhwsSAgIDaa0KWBW2bZvoIGWRq5lg2Li5uypRJljV6vE00PWBAucUtOTnZ7JgRAEC7BHToEMiapik2da7M1J9+Or+GNM30AjSN58+f37x5kx5LMidnag6GN57+VrnMxdaHDx86dOiQnopEw9fXp0qD1hs3biB3Wr05dCb69u3LlJS5/1xNtGmdSc3x5s3bTMevbdu26haLs2rnFy6cBwA0adJk0KBBxtIwGfzJk8dm7RjMheYAgB07tkdHP6x3mjZLmq9evQIAeHt7673h+Hw+04Xcsk3I4qBt02ZXVB06dJg5Y3zv3r34+JcWNHq8TTTt5eXdr18/+uf169fYjKORjY+9Nm3W6MFkar0ZRZ3Rg1e9B8D+RtHd28vLe+bMmZZtnahLcLncxo315/RHjx7dvbv2ix4FBQX5+Xldu3YbMcLAKj7mFBCbmXqkSguFQhMchyrA1EbZvNgYy1tUplO6urqGh19wcXGlY5DbCZvORlEUem0MGjTYRPqgoCCmg61ZhTo4OJjZxlQq1aefflp7xMSyBdIfkTGtTRv8QsrEiRNpm7tlm1CVwEY9p5Vo09o0wp49e5mjvQsXLvxHjR4AgEGDhjDeYIfM3bjdaN6GuTjYLMRiMerSbPYjX7Toi8rfV6wSTTOJJi+P1RKVY8eO0qtd//zzT/a7MlbJ6KFnmKahZ0M0tkyDaTAxu5A9IyMd0fSgQYOZc0eVMXz4COZVbL6KgOzdAACJxPwD9fLyOnnylN4KDjY0ff58OFrQwTRrGASTksLCztCLPky0MebLKSMjffz4UJZT6DWhAz2PTyZu377NRpvWM0wj2NjYTpw4qapNqPIuTixBP309XL586f79KJbmMj3/DWPYunWbnrP8f9HoAQAYMmQIbdt6+vTJnDmzjaUsKMhHXybs27fvt99+x74I2lmH5bOZM2eu3sqXahs98vLMvOGfPn2ybNlS5Idgb29/7NgJi3M0ACA+Ph5UcjOgERIyrE+fPijcsWPHMWMMm8WZqpDZXn38+PGysjJQ0TnaINzc3Jg6C5vNGQoKtC8/5HtgFk2aNDlw4BDzW01sjB579uwBAAQEBNDme2PQ28fdrLYBANi06dePPy7/ilBERMSkSRNq43MQzNZobBViTEyMnsW5suaRmZmJRicGE4wfP4HOgWUTQp92rQYMLkxH95D9ysbExASWKTdu/IVpGHxrjB4ZGenMwebjx49qUp6zs/O2bdvbtGlL2zQGDx64Y8d2vR64d++ejh0DdfrmgSoVQa+6zs3NYXnJ9OkzmPscVkmbZpYiFosXL/6y8l53GRnpDx7cX7ZsaUjIUKQ/jhkzVm+K2VJ4/fo1WiEWEGD0w4O0Qj1+vNEVz8yt6F+9SqQ/CWhw+IncZoGBL9cZANPfo6SkxMSrWkfThVV67wIAWrRosXPnTvpFZbazbdy4AZmMg4KC2NjuunXrRv88fPiQWYUaAPC//y3fsGEj/fJ4/vz5nDmzv/zyC/qLzNWjFRNd9fbtW7t374qIiEAuhqht7Nu378MPp9rb25vIx6zm6+zsTG9KZaIJ2duXL0ljv4FtRMQ9vSF1ePi5uLg4Wq2+fv36559/GhZ2xsHBkWWext5YBrF27bqZM2fpXvDVp2kmVarVKvopVAPm9/Q4fvw48+f169f37ds3ffr0ahfp7x/w118nt2z5Da0refHixfffr/nppx/d3d09PT3fvHmTl5eH7BWmP71e2dZx585tpVJFmzI3bPh5wYJFAQEBbPz5J02aLBQKFyz4vEo0HR5+Tq8FHDt29OTJv5ycnJycnEQiUWZmZnZ2Nu0YJBKJ+vXrP3z4CFqftTjo/f9evXqVm5tr8JNxAwcOGjVq1IMHDydNmmQsHz1eXrz4i++++75Ro0ZNmzZlxl+9emXDhg30Up2MjEy9T45Whp4JKzz83IkTxw1+UO7hwwdqtYZu7j/8sN7Pzy8gIIDNEKRt23bbt++cP3/eq1eJJmi6oCD/6NGj9BegSZKTkZFuVoTg4O7I6o3w888/LVz4hTH/Fub7KTg4ePfu3Xfv3kEjnuPHjx0/fkwkcunYsaObm5ubm1t8/MvXr1/TE25+fn59+/Zjueju1i39zd5Wr/7WYErTeuiWLdoFX8+ePTPWhObNm//gwYOUlBQTTSghIZ6hPOWNHx+6YMGiyk1ID1evXq0o1E20dr8yTKxMTkpKKioqpFeuFRQUrFq1YtSoMf7+/mw+hbNy5SqhUPj771uqbfR4+vTJ9evXmTG7du2sklWACQJCaOzcgQP7d+3aaWNj07JlK5FIxOfzCwsLi4oKc3JyUlJSioqKLlz4p02bNjXR+65fv372bFixDkKh0NHR0dHRsVWrVhMmTOjSpSv73D74YDKfL7Czs3NwcHBwcBCLxSUlJWKxuLS0ZNWq1SzrefnypZkzZ7D5vtk//1z8+uuvBgwY0KiR1gxaXFxcWFhYVFRUXFwsFhejL717eno2buzp6dm4cWPP3r371Ib6TGPZsqXXrl1t1qx5q1atrKysxGJxcXFxaurr2bPnVF7Nn5iYIJFI9b6DrtfOLl26JNNBLpej48CBA9H+0egOdOnSxc7O3sHBQSAQZGdnI7374cMYY18UDQ8/t3v3rh49eqJsy8rKFAqFXC5/8eJFbm7OxYv/0BQcHn7uwIEDIpEINQmlUllSIi4uLhaLSwYNGsTcw9q0GrVz545u3boZfA388svGyMgId3d3JIJMJktOTn71KrFHj556u8jTePPmza5dO3fv3jV37jyZTKZQyOmbExcXN3Lk+ytWrGRTsdjY2NjY2Ojoh7du3USzKWgEY2dnZ2dnb29v16FDh6FDhzHVdhM4cGD/3r17PTw8OnfuDADQaDSoQxUVFTF3HaKxatVqg3PyX3yx6OrVKwMHDnJ2dgYAisUlxcXFr14lLly4qHITyshIl8lkJl6ZERERd+7cptsPDboJVbJlewUGBrZr1x7pVVKpFIkglRpYsezh0Vhv50vmM1qw4DMHBwcnJydHR0cAiNLSUnRDWrRosXDhItODCaYin5iYaKwlGENmZmZwcLdWrVp5eDRu0qSJnZ2dVCqVSiUSiVQqlcrlsunTZ1R1ozpTNI2BgYGBUe8g8S3AwMDAwDSNgYGBgYFpGgMDAwPTNAYGBgYGpmkMDAwMjHL8H4MjwusQKkVhAAAAAElFTkSuQmCC" alt="RPD Fund Management logo"></a><button type="button" class="mobile-nav-toggle" aria-expanded="false" aria-controls="primary-nav" aria-label="Open navigation"><span></span><span></span><span></span></button><nav id="primary-nav">{auth}</nav><div class="top-clock" aria-label="New York date and time"><div class="top-clock-icon">◷</div><div class="top-clock-text"><span class="top-clock-label">New York</span><span class="top-clock-value" id="global-ny-clock">Loading…</span></div></div></div></div>
<div class="wrap">{body_shell}</div>
<script>(function(){{const el=document.getElementById("global-ny-clock");if(el){{const fmt=new Intl.DateTimeFormat("en-US",{{timeZone:"America/New_York",weekday:"short",month:"short",day:"numeric",year:"numeric",hour:"numeric",minute:"2-digit",second:"2-digit",hour12:true}});function tick(){{el.textContent=fmt.format(new Date());}}tick();setInterval(tick,1000);}}const btn=document.querySelector(".mobile-nav-toggle");const nav=document.getElementById("primary-nav");if(btn&&nav){{btn.addEventListener("click",function(){{const open=nav.classList.toggle("mobile-open");btn.setAttribute("aria-expanded",open?"true":"false");btn.setAttribute("aria-label",open?"Close navigation":"Open navigation");}});nav.querySelectorAll("a").forEach(function(a){{a.addEventListener("click",function(){{if(window.innerWidth<=760){{nav.classList.remove("mobile-open");btn.setAttribute("aria-expanded","false");}}}});}});}}}})();</script></body></html>""")

def admin_tabs():
    return ''

@app.get("/health")
def health():return {"ok":True}

@app.get("/login")
def login_get():
    return page("Login","""<div class="card" style="max-width:420px;margin:80px auto"><h2>Sign in</h2>
<form method="post"><label>Username<input name="username"></label><br><label>Password<input type="password" name="password"></label><br><button>Sign in</button></form></div>""")

@app.post("/login")
def login_post(username:str=Form(...),password:str=Form(...)):
    name=username.strip().lower()
    with SessionLocal() as db:
        u=db.query(User).filter_by(username=name,active=True).first()
        if not u or not verify_password(password,u.password_hash):
            return page("Login failed",'<div class="card"><h2>Login failed</h2><a href="/login">Try again</a></div>')
        if name==ADMIN_USER:audit(db,name,"ADMIN_LOGIN","Admin signed in");db.commit()
    resp=RedirectResponse("/",303);resp.set_cookie("session",sign_session(name),httponly=True,secure=True,samesite="strict",max_age=604800)
    return resp

@app.get("/logout")
def logout():
    r=RedirectResponse("/login",303);r.delete_cookie("session");return r

def latest_snapshot(db):
    s=db.query(Snapshot).order_by(Snapshot.id.desc()).first()
    if not s:return None,None
    try:p=json.loads(s.payload)
    except Exception:return None,None
    try:
        ts=datetime.fromisoformat(p["timestamp_utc"].replace("Z","+00:00"))
        age=(datetime.now(timezone.utc)-ts.astimezone(timezone.utc)).total_seconds()
    except Exception:age=999999
    return p,age

def format_new_york_time(timestamp_utc):
    try:
        ts=datetime.fromisoformat(str(timestamp_utc).replace("Z","+00:00"))
        ny=ts.astimezone(ZoneInfo("America/New_York"))
        return ny.strftime("%d %b %Y, %H:%M:%S %Z")
    except Exception:
        return str(timestamp_utc or "")



def _json_setting(db,key,default=None):
    raw=get_setting(db,key,"")
    if not raw:
        return default
    try:
        x=json.loads(raw)
        return x
    except Exception:
        return default

def _set_json_setting(db,key,value):
    set_setting(db,key,json.dumps(value,separators=(",",":")))

def _opportunity_baseline(db):
    x=_json_setting(db,"opportunity_daily_baseline",None)
    return x if isinstance(x,dict) else None

def _opportunity_exposure(db):
    x=_json_setting(db,"opportunity_exposure_config",None)
    return x if isinstance(x,dict) else None

def _opp_position_identity(p):
    """Stable instrument identity for Opportunity history/change tracking."""
    inst=str(p.get("instrument_type") or "OPTION").upper()
    ticker=str(p.get("ticker") or "").upper().strip()
    if inst=="OPTION":
        expiry=str(p.get("expiry") or "")
        typ=str(p.get("option_type") or "").upper()[:1]
        try: strike=round(float(p.get("strike")),6)
        except Exception: strike=0.0
        return f"OPTION|{ticker}|{expiry}|{typ}|{strike:g}"
    return f"EQUITY|{ticker}"


def _opp_position_label(p):
    inst=str(p.get("instrument_type") or "OPTION").upper()
    ticker=str(p.get("ticker") or "").upper().strip()
    if inst=="OPTION":
        try: strike=f"{float(p.get('strike')):g}"
        except Exception: strike=str(p.get("strike") or "")
        return f"{ticker} {p.get('expiry') or ''} {str(p.get('option_type') or '').upper()[:1]}{strike}".strip()
    return ticker


def _opp_report_day(report_date=""):
    s=str(report_date or "").strip()
    if re.fullmatch(r"\d{4}-\d{2}-\d{2}",s):
        return s
    return date.today().isoformat()


def _opp_is_expired(p,as_of=None):
    expiry=str(p.get("expiry") or "").strip()
    if not expiry:
        return False
    day=str(as_of or date.today().isoformat())
    return expiry < day


def _opp_status(p,as_of=None):
    if _opp_is_expired(p,as_of):
        return "EXPIRED"
    return "ACTIVE" if bool(p.get("active",True)) else "ARCHIVED"


def _opp_exposure_history(db,exp=None):
    hist=_json_setting(db,"opportunity_exposure_portfolio_history",None)
    if isinstance(hist,dict) and isinstance(hist.get("positions"),list):
        return hist
    exp=exp if isinstance(exp,dict) else _opportunity_exposure(db)
    positions=[]
    for p in list((exp or {}).get("positions") or []):
        x=dict(p);x["active"]=True
        x["first_seen"]=(exp or {}).get("report_date") or ""
        x["last_seen"]=(exp or {}).get("report_date") or ""
        x["exit_status"]=""
        positions.append(x)
    return {"positions":positions,"report_date":(exp or {}).get("report_date") or "","filename":(exp or {}).get("filename") or ""}


def _opp_compare_positions(old_positions,new_positions,as_of):
    old_map={_opp_position_identity(p):dict(p) for p in (old_positions or [])}
    new_map={_opp_position_identity(p):dict(p) for p in (new_positions or [])}
    added=[];archived=[];expired=[];qty_changed=[];reactivated=[]
    for k,n in new_map.items():
        o=old_map.get(k)
        if o is None:
            added.append({"position":_opp_position_label(n),"quantity":n.get("quantity")})
            continue
        try: oq=float(o.get("quantity") or 0);nq=float(n.get("quantity") or 0)
        except Exception: oq=o.get("quantity");nq=n.get("quantity")
        if oq!=nq:
            qty_changed.append({"position":_opp_position_label(n),"old":oq,"new":nq})
        if o.get("active") is False:
            reactivated.append({"position":_opp_position_label(n),"quantity":n.get("quantity")})
    for k,o in old_map.items():
        if k not in new_map and o.get("active",True):
            item={"position":_opp_position_label(o),"quantity":o.get("quantity")}
            if _opp_is_expired(o,as_of): expired.append(item)
            else: archived.append(item)
    return {"added":added,"archived":archived,"expired":expired,"quantity_changed":qty_changed,"reactivated":reactivated}


def _opp_update_exposure_history(db,parsed,filename):
    previous_exp=_opportunity_exposure(db)
    hist=_opp_exposure_history(db,previous_exp)
    old_positions=list(hist.get("positions") or [])
    as_of=_opp_report_day(parsed.get("report_date"))
    changes=_opp_compare_positions(old_positions,parsed.get("positions") or [],as_of)
    registry={_opp_position_identity(p):dict(p) for p in old_positions}
    incoming={_opp_position_identity(p):dict(p) for p in (parsed.get("positions") or [])}
    for key,p in incoming.items():
        old=registry.get(key,{})
        x=dict(old);x.update(p)
        x["active"]=True
        x["first_seen"]=old.get("first_seen") or parsed.get("report_date") or as_of
        x["last_seen"]=parsed.get("report_date") or as_of
        x["archived_at"]="";x["exit_status"]=""
        registry[key]=x
    for key,old in list(registry.items()):
        if key in incoming: continue
        if old.get("active",True):
            old["active"]=False
            old["archived_at"]=parsed.get("report_date") or as_of
            old["exit_status"]="EXPIRED" if _opp_is_expired(old,as_of) else "ARCHIVED"
            registry[key]=old
    out={"filename":filename,"report_date":parsed.get("report_date") or "","updated_at":utcnow().isoformat(),"positions":list(registry.values())}
    _set_json_setting(db,"opportunity_exposure_portfolio_history",out)
    change_record={"filename":filename,"report_date":parsed.get("report_date") or "","created_at":utcnow().isoformat(),**changes}
    _set_json_setting(db,"opportunity_exposure_last_changes",change_record)
    return changes


def _opp_pnl_changes(db,parsed,filename):
    previous=_opportunity_baseline(db)
    old=[]
    for p in list((previous or {}).get("positions") or []):
        x=dict(p);x["active"]=True;old.append(x)
    new=list(parsed.get("positions") or [])
    as_of=_opp_report_day(parsed.get("report_date"))
    changes=_opp_compare_positions(old,new,as_of)
    rec={"filename":filename,"report_date":parsed.get("report_date") or "","created_at":utcnow().isoformat(),**changes}
    _set_json_setting(db,"opportunity_pnl_last_changes",rec)
    return changes


def _opp_changes_html(changes,title):
    if not isinstance(changes,dict):
        return f'<div class="notice"><strong>{html.escape(title)}:</strong> Tracking is ready. Changes will appear after the next upload.</div>'
    groups=[("Added","added"),("Reactivated","reactivated"),("Archived / Removed","archived"),("Expired","expired"),("Quantity Changed","quantity_changed")]
    total=sum(len(changes.get(k) or []) for _,k in groups)
    head=f'<div class="notice"><strong>{html.escape(title)}</strong> · {html.escape(str(changes.get("filename") or ""))} · Report {html.escape(str(changes.get("report_date") or ""))}'
    if total==0:
        return head+'<br><span class="ok"><strong>No position changes detected.</strong></span></div>'
    parts=[head+'<br>']
    for label,key in groups:
        items=changes.get(key) or []
        if not items: continue
        parts.append(f'<div style="margin-top:8px"><strong>{html.escape(label)} — {len(items)}</strong><ul style="margin:5px 0 0 20px">')
        for item in items:
            pos=html.escape(str(item.get("position") or ""))
            if key=="quantity_changed":
                try: detail=f'{float(item.get("old")):,.0f} → {float(item.get("new")):,.0f}'
                except Exception: detail=f'{item.get("old")} → {item.get("new")}'
                parts.append(f'<li>{pos}: {html.escape(detail)}</li>')
            else:
                parts.append(f'<li>{pos}</li>')
        parts.append('</ul></div>')
    parts.append('</div>')
    return ''.join(parts)


def _opp_portfolio_table(hist):
    rows=[];today=date.today().isoformat()
    positions=list((hist or {}).get("positions") or [])
    def sk(p):
        try: strike=float(p.get("strike") or 0)
        except Exception: strike=0.0
        return (0 if _opp_status(p,today)=="ACTIVE" else 1,str(p.get("ticker") or ""),str(p.get("expiry") or ""),strike)
    positions.sort(key=sk)
    for p in positions:
        inst=str(p.get("instrument_type") or "OPTION").upper();ticker=html.escape(str(p.get("ticker") or ""))
        expiry=html.escape(str(p.get("expiry") or "")) if inst=="OPTION" else "—"
        typ=html.escape(str(p.get("option_type") or "")) if inst=="OPTION" else "EQUITY"
        try: strike=f'{float(p.get("strike")):g}' if inst=="OPTION" else "—"
        except Exception: strike=html.escape(str(p.get("strike") or ""))
        try: qty=f'{float(p.get("quantity") or 0):,.0f}'
        except Exception: qty=html.escape(str(p.get("quantity") or ""))
        status=_opp_status(p,today);sec=html.escape(str(p.get("bloomberg_security") or ""))
        rows.append(f'<tr><td>{ticker}</td><td>{expiry}</td><td>{typ}</td><td>{strike}</td><td>{qty}</td><td>{status}</td><td>{sec}</td></tr>')
    if not rows:return '<p>No Opportunity portfolio positions available.</p>'
    return '<div style="overflow-x:auto"><table><thead><tr><th>Ticker</th><th>Expiry</th><th>Type</th><th>Strike</th><th>Qty</th><th>Status</th><th>Bloomberg Security</th></tr></thead><tbody>'+''.join(rows)+'</tbody></table></div>'


def _opp_pnl_positions_table(base):
    rows=[]
    for p in list((base or {}).get("positions") or []):
        name=html.escape(_opp_position_label(p))
        try: qty=f'{float(p.get("quantity") or 0):,.0f}'
        except Exception: qty=html.escape(str(p.get("quantity") or ""))
        try: mark=f'${float(p.get("baseline_price")):,.2f}'
        except Exception: mark='—'
        rows.append(f'<tr><td>{name}</td><td>{qty}</td><td>{mark}</td></tr>')
    if not rows:return '<p>No Opportunity P&amp;L baseline positions available.</p>'
    return '<div style="overflow-x:auto"><table><thead><tr><th>Position</th><th>Qty</th><th>Previous Mark</th></tr></thead><tbody>'+''.join(rows)+'</tbody></table></div>'


def _latest_opportunity_snapshot(db):
    p=_json_setting(db,"opportunity_bloomberg_snapshot",None)
    if not isinstance(p,dict):
        return None,None
    try:
        ts=datetime.fromisoformat(str(p.get("timestamp_utc") or "").replace("Z","+00:00"))
        age=(datetime.now(timezone.utc)-ts.astimezone(timezone.utc)).total_seconds()
    except Exception:
        age=999999
    return p,age

def _latest_yahoo_v2_fund_snapshot(db,fund_key):
    root=_json_setting(db,"yahoo_collector_snapshot_v2",None)
    if not isinstance(root,dict):
        return None
    funds=root.get("funds") or {}
    x=funds.get(fund_key)
    return x if isinstance(x,dict) else None

def _quote_data_from_snapshot(snapshot):
    out={}
    if not snapshot:
        return out
    for r in snapshot.get("positions",[]) or []:
        symbol=str(r.get("symbol") or "").upper().strip()
        if not symbol:
            continue
        out[symbol]={
            "mark":r.get("mark"),"bid":r.get("bid"),"ask":r.get("ask"),"last":r.get("last"),
            "source":r.get("source") or "","note":r.get("note") or "",
        }
    return out

def _dr_position_key(p):
    inst=str(p.get("instrument_type") or "OPTION").upper()
    if inst=="OPTION":
        return ("OPTION",str(p.get("ticker") or "").upper(),str(p.get("expiry") or ""),str(p.get("option_type") or "").upper(),round(float(p.get("strike") or 0),6))
    return ("EQUITY",str(p.get("ticker") or "").upper())

def _bloomberg_mark_map(snapshot):
    out={}
    if not snapshot:
        return out
    for r in snapshot.get("positions",[]) or []:
        try:
            inst=str(r.get("instrument_type") or "OPTION").upper()
            if inst=="OPTION":
                key=("OPTION",str(r.get("ticker") or "").upper(),str(r.get("expiry") or ""),str(r.get("option_type") or "").upper(),round(float(r.get("strike") or 0),6))
            else:
                key=("EQUITY",str(r.get("ticker") or "").upper())
            out[key]={
                "mark":r.get("market_mark",r.get("option_mark")),
                "source":r.get("market_source",r.get("option_mark_source")) or "",
                "bid":r.get("market_bid",r.get("option_bid")),
                "ask":r.get("market_ask",r.get("option_ask")),
                "last":r.get("market_last",r.get("option_last")),
            }
        except Exception:
            continue
    return out

def _calc_bloomberg_daily(positions,nav,snapshot,manual_field="bloomberg_manual_price"):
    marks=_bloomberg_mark_map(snapshot)
    total=0.0; valid=0; rows=[]
    for p in positions:
        q=marks.get(_dr_position_key(p),{})
        market=q.get("mark"); source=q.get("source") or ""; manual=p.get(manual_field)
        if source in ("PX_BID/ASK_MID","PX_ASK_HALF","EQUITY_BID/ASK_MID") and market is not None:
            current=float(market); eff="BLOOMBERG MID" if source!="PX_ASK_HALF" else "BLOOMBERG ASK/2"
        elif manual is not None:
            current=float(manual); eff="MANUAL OVERRIDE"
        elif market is not None:
            current=float(market); eff="BLOOMBERG LAST"
        else:
            current=None; eff="UNAVAILABLE"
        prev=p.get("baseline_price"); pnl=chg=contrib=None
        if current is not None and prev is not None:
            chg=current-float(prev); pnl=chg*float(p.get("quantity") or 0)*float(p.get("multiplier") or 1); total+=pnl; valid+=1
            contrib=pnl/float(nav)*100 if nav else None
        rows.append({**p,"market_mark":market,"market_source":source,"bid":q.get("bid"),"ask":q.get("ask"),"last":q.get("last"),"current_mark":current,"effective_source":eff,"change":chg,"estimated_pnl":pnl,"contribution_pct":contrib})
    n=len(positions)
    return {"rows":rows,"estimated_pnl":total,"estimated_return_pct":total/float(nav)*100 if nav else 0.0,"coverage_valid":valid,"coverage_total":n,"complete":bool(n and valid==n)}


# === v6.15 Intraday trade ledger / verified expiration (additive; no migration) ===
def _ny_now():
    return datetime.now(ZoneInfo("America/New_York"))

def _ledger_key(fund):
    if fund not in ("fortress","opportunity"):
        raise HTTPException(400,"Unknown fund")
    return f"intraday_trades_v615_{fund}"

def _trade_history(db,fund):
    value=_json_setting(db,_ledger_key(fund),[])
    return value if isinstance(value,list) else []

def _expiry_history(db,fund):
    value=_json_setting(db,f"expiration_confirm_v615_{fund}",[])
    return value if isinstance(value,list) else []

def _active_trades(db,fund,report_date,as_of=None):
    # Reported EOD positions already include trades on or before the report day.
    # Never count those trades twice when a new Nirvana baseline is uploaded.
    as_of=as_of or _ny_now().date().isoformat()
    return [t for t in _trade_history(db,fund)
            if t.get("status")=="ACTIVE" and report_date < t.get("trade_date","") <= as_of]

def _base_for_trade(db,fund):
    if fund=="fortress":
        b,pos=_fortress_daily_positions(db)
        return (str(b.report_date or "") if b else "",pos)
    b=_opportunity_baseline(db)
    return (str((b or {}).get("report_date") or ""),list((b or {}).get("positions") or []))

def _daily_report_date(db,fund):
    return _base_for_trade(db,fund)[0]

def _trade_to_position(t):
    ticker=t["ticker"]
    is_option=t["instrument_type"]=="OPTION"
    bbg=(option_security(ticker,t.get("expiry") or "",t.get("option_type") or "P",t.get("strike"))
         if is_option else f"{ticker} US Equity")
    return {"instrument_type":t["instrument_type"],"ticker":ticker,
            "expiry":t.get("expiry") or "","option_type":t.get("option_type") or "",
            "strike":t.get("strike"),"quantity":0.0,
            "multiplier":float(t.get("multiplier") or 1),
            "baseline_price":0.0,"security_name":_opp_position_label(t),
            "bloomberg_security":bbg,"underlying_security":f"{ticker} US Equity",
            "bloomberg_manual_price":None,"yahoo_manual_price":None}

def _overlay_positions(baseline_positions,trades,include_flat=False):
    result={_dr_position_key(p):dict(p) for p in baseline_positions}
    for t in trades:
        k=_dr_position_key(t)
        if k not in result: result[k]=_trade_to_position(t)
        p=result[k]
        if abs(float(p.get("multiplier") or 1)-float(t.get("multiplier") or 1))>1e-9:
            raise ValueError(f"Multiplier mismatch for {_opp_position_label(t)}")
        p["quantity"]=float(p.get("quantity") or 0)+float(t["signed_quantity"])
    return [p for p in result.values() if include_flat or abs(float(p.get("quantity") or 0))>1e-7]

# v6.15.1: automatic post-close OTM estimate from Bloomberg underlying spot.
# This is explicitly a PROVISIONAL valuation, NOT exercise/assignment confirmation.
AUTO_OTM_MIN_PCT = 2.0
AUTO_OTM_MIN_USD = 0.10

def _expiry_record_status(db, fund):
    # Most recent record per instrument; include CLEAR (= block auto) and RESUME.
    return {_dr_position_key(x):str(x.get("status") or "") for x in _expiry_history(db,fund)}

def _closing_bloomberg_snapshot(db, fund, expiry_day):
    snap = latest_snapshot(db)[0] if fund=="fortress" else _latest_opportunity_snapshot(db)[0]
    if not isinstance(snap, dict) or str(snap.get("mode") or "")!="MARKET_CLOSED":
        return None, "Bloomberg final closing snapshot is unavailable"
    try:
        ts=datetime.fromisoformat(str(snap.get("timestamp_utc") or "").replace("Z","+00:00"))
        if ts.tzinfo is None:return None,"Snapshot timestamp has no timezone"
        ts=ts.astimezone(ZoneInfo("America/New_York"))
    except Exception:
        return None,"Bloomberg closing timestamp is invalid"
    if ts.date().isoformat()!=expiry_day or not (ts.hour==16 and ts.minute<=45):
        return None,"No Bloomberg closing snapshot for this expiration date"
    return snap,""

def _auto_expiry_assessment(db,fund,report_date,baseline_positions,trades=None,as_of=None):
    """Read-only decision. Never silently write expirations, change the baseline,
    or treat a quote as official exercise/assignment confirmation.
    Fail closed if spot is missing, snapshot is wrong date, or underlying is near strike.
    """
    now=_ny_now();today=str(as_of or now.date().isoformat())
    latest=_expiry_record_status(db,fund)
    effective=_overlay_positions(baseline_positions, trades if trades is not None else _active_trades(db,fund,report_date,today))
    choices={}
    for p in effective:
        if str(p.get("instrument_type") or "OPTION").upper()!="OPTION":continue
        expiry=str(p.get("expiry") or "")
        if not (report_date < expiry <= today):continue
        key=_dr_position_key(p);state=latest.get(key)
        result={"status":"REVIEW","reason":"Pending Bloomberg closing quote", "underlying":None, "otm_pct":None, "snapshot_at":""}
        if state=="CONFIRMED_WORTHLESS":
            result.update(status="CONFIRMED",reason="Manually verified worthless")
            choices[key]=result;continue
        if state=="CLEARED":
            result.update(status="BLOCKED",reason="Excluded by admin (undo confirmation)")
            choices[key]=result;continue
        if expiry==today and (now.hour,now.minute)<(16,0):
            result.update(status="WAIT",reason="Expiration day: wait until 4:00 PM New York")
            choices[key]=result;continue
        snap,why=_closing_bloomberg_snapshot(db,fund,expiry)
        if snap is None:
            result["reason"]=why;choices[key]=result;continue
        result["snapshot_at"]=str(snap.get("timestamp_utc") or "")
        # Some Bloomberg feeds can duplicate the same instrument in the payload;
        # never infer spot from a different contract, symbol, Delta, or quote source.
        matching=[r for r in (snap.get("positions") or []) if _dr_position_key(r)==key]
        if not matching:
            result["reason"]="Instrument missing from Bloomberg final snapshot";choices[key]=result;continue
        spot_raw=matching[-1].get("underlying_spot")
        try:spot=float(spot_raw)
        except (ValueError,TypeError):spot=0.0
        src=str(matching[-1].get("spot_source") or "")
        if not (math.isfinite(spot) and spot>0 and src in ("OPT_UNDL_PX","UNDERLYING_BID_ASK_MID","UNDERLYING_PX_LAST")):
            result["reason"]="Underlying closing-window price unavailable/unverified (Collector v4.2 needed)"
            choices[key]=result;continue
        strike=float(p.get("strike") or 0)
        typ=str(p.get("option_type") or "").upper()
        if not (math.isfinite(strike) and strike>0 and typ in ("P","C")):
            result["reason"]="Invalid option contract";choices[key]=result;continue
        # Put: spot > strike. Call: spot < strike.
        distance=(spot-strike) if typ=="P" else (strike-spot)
        pct=distance/spot*100
        result["underlying"]=spot;result["otm_pct"]=pct
        if distance<=0:
            result.update(status="REVIEW",reason="ITM/ATM at Bloomberg closing-window price: do not zero")
        elif distance<max(AUTO_OTM_MIN_USD, spot*AUTO_OTM_MIN_PCT/100):
            result.update(status="REVIEW",reason="Near strike: assignment/exercise risk review")
        else:
            result.update(status="AUTO",reason="Provisional OTM = estimated $0; exercise/assignment not confirmed")
        choices[key]=result
    return choices

def _confirmed_expiry_keys(db,fund,baseline_day,as_of=None):
    as_of=as_of or _ny_now().date().isoformat()
    return {_dr_position_key(x) for x in _expiry_history(db,fund)
            if x.get("status")=="CONFIRMED_WORTHLESS" and baseline_day < x.get("expiry","") <= as_of}

def _daily_with_ledger(db,fund,source,positions,nav,snapshot,baseline_date):
    """v6.15 unchanged pricing plus intraday cash flows and v6.15.1 provisional expiry."""
    if source=="bloomberg":
        calc=_calc_bloomberg_daily(positions,nav,snapshot)
    else:
        quotes=_quote_data_from_snapshot(snapshot)
        calc=calculate_mark_to_market_return(positions,nav,quote_data=quotes)
    day=str(baseline_date or "")
    trades=_active_trades(db,fund,day)
    expiry_status=_auto_expiry_assessment(db,fund,day,positions,trades)
    zero_keys={k for k,r in expiry_status.items() if r["status"] in ("AUTO","CONFIRMED")}
    if not trades and not zero_keys:
        return {**calc,"auto_expiry_count":0,"expiry_review_count":sum(1 for x in expiry_status.values() if x["status"]=="REVIEW")}
    raw={_dr_position_key(p):dict(p) for p in positions}
    rows={_dr_position_key(r):dict(r) for r in calc.get("rows",[])}
    legs={}
    for t in trades:
        k=_dr_position_key(t);legs.setdefault(k,[]).append(t)
        if k not in raw: raw[k]=_trade_to_position(t)
    for k,p in raw.items():
        if k in rows: continue
        if source=="bloomberg":
            fresh=_calc_bloomberg_daily([p],nav,snapshot)
        else:
            fresh=calculate_mark_to_market_return([p],nav,quote_data=quotes)
        rows[k]=dict(fresh["rows"][0])
    newrows=[];total=0.0;valid=0
    for k,p in raw.items():
        r=rows[k];q0=float(p.get("quantity") or 0)
        ls=legs.get(k,[])
        dq=sum(float(t["signed_quantity"]) for t in ls)
        net=round(q0+dq,9)
        mark=r.get("current_mark")
        expiry_info=expiry_status.get(k)
        zero=k in zero_keys
        if zero:
            mark=0.0;r["current_mark"]=0.0
            r["effective_source"]=("VERIFIED EXPIRED WORTHLESS" if expiry_info["status"]=="CONFIRMED"
                                   else "AUTO OTM EXPIRY ($0 ESTIMATE)")
        if mark is None and abs(net)<1e-8:
            mark=0.0;r["current_mark"]=0.0;r["effective_source"]="CLOSED (TRADE CASH FLOWS)"
        previous=p.get("baseline_price")
        multiplier=float(p.get("multiplier") or 1)
        pnl=None
        if mark is not None and (previous is not None or not q0):
            pnl=q0*(float(mark)-float(previous or 0))*multiplier
            pnl+=sum(float(t["signed_quantity"])*(float(mark)-float(t["execution_price"]))*multiplier for t in ls)
            total+=pnl;valid+=1
        r["base_quantity"]=q0;r["quantity"]=net;r["intraday_trades"]=len(ls)
        r["expired_confirmed"]=bool(expiry_info and expiry_info["status"]=="CONFIRMED")
        r["expired_auto_estimate"]=bool(expiry_info and expiry_info["status"]=="AUTO")
        r["estimated_pnl"]=pnl
        r["contribution_pct"]=pnl/float(nav)*100 if pnl is not None and nav else None
        r["change"]=float(mark)-float(previous) if mark is not None and previous is not None else None
        r["baseline_price"]=previous;r["previous_mark"]=previous
        newrows.append(r)
    return {**calc,"rows":newrows,"estimated_pnl":total,
            "estimated_return_pct":total/float(nav)*100 if nav else 0.0,
            "coverage_valid":valid,"coverage_total":len(newrows),
            "complete":bool(newrows and valid==len(newrows)),
            "intraday_count":len(trades),"expired_confirmed_count":sum(1 for r in expiry_status.values() if r["status"]=="CONFIRMED"),
            "auto_expiry_count":sum(1 for r in expiry_status.values() if r["status"]=="AUTO"),
            "expiry_review_count":sum(1 for r in expiry_status.values() if r["status"]=="REVIEW")}

def _version_bump_for_trade(db,fund,actor,reason):
    if fund=="fortress":
        config_bump(db,actor,reason)
    else:
        _bump_opportunity_config(db,actor,reason)

def _intraday_panel(db,fund):
    report,baseline=_base_for_trade(db,fund)
    today=_ny_now().date().isoformat()
    active=_active_trades(db,fund,report,today)
    history=list(reversed(_trade_history(db,fund)))[:35]
    active_positions=_overlay_positions(baseline,active)
    position_lines=[]
    for p in sorted(active_positions,key=lambda x:_opp_position_label(x)):
        position_lines.append(f'<option value="{html.escape(_opp_position_label(p))}">{html.escape(_opp_position_label(p))}: {float(p.get("quantity") or 0):,.0f}</option>')
    trade_rows=[]
    for t in history:
        status="Active" if t.get("status")=="ACTIVE" else "Cancelled"
        if t.get("status")=="ACTIVE" and t.get("trade_date","")<=report:
            status="Included in Nirvana baseline"
        remove=(f'<form method="post" action="/admin/intraday/{fund}/cancel" onsubmit="return confirm(\'Cancel this trade?\');"><input type="hidden" name="trade_id" value="{html.escape(str(t["id"]))}"><button type="submit">Cancel</button></form>' if status=="Active" else "—")
        trade_rows.append(f'<tr><td>{html.escape(str(t.get("trade_date")))}</td><td>{html.escape(_opp_position_label(t))}</td><td>{html.escape(str(t.get("action")))}</td><td>{float(t.get("signed_quantity") or 0):+,.0f}</td><td>${float(t.get("execution_price") or 0):.4f}</td><td>{status}</td><td>{remove}</td></tr>')
    exp_rows=[]
    now=_ny_now()
    assessments=_auto_expiry_assessment(db,fund,report,baseline,active)
    for pos in _overlay_positions(baseline,active):
        if str(pos.get("instrument_type") or "OPTION").upper()!="OPTION":continue
        expiry=str(pos.get("expiry") or "")
        if not expiry or expiry<=report or expiry>today:continue
        k=_dr_position_key(pos);a=assessments.get(k) or {}
        status=a.get("status","REVIEW")
        params=''.join(f'<input type="hidden" name="{f}" value="{html.escape(str(pos.get(f) or ""))}">' for f in ("ticker","expiry","option_type","strike"))
        label=html.escape(_opp_position_label(pos))
        price=('—' if a.get("underlying") is None else f'${a["underlying"]:,.2f}')
        distance=('—' if a.get("otm_pct") is None else f'{a["otm_pct"]:+.2f}%')
        detail=html.escape(str(a.get("reason") or ""))
        if status=="CONFIRMED":
            action=f'<form method="post" action="/admin/intraday/{fund}/expiration">{params}<input type="hidden" name="operation" value="clear"><button>Undo / exclude auto</button></form>'
            label_status="Verified manually"
        elif status=="BLOCKED":
            action=f'<form method="post" action="/admin/intraday/{fund}/expiration">{params}<input type="hidden" name="operation" value="resume"><button>Resume auto</button></form>'
            label_status="Excluded by admin"
        elif status=="AUTO":
            action=f'<form method="post" action="/admin/intraday/{fund}/expiration">{params}<input type="hidden" name="operation" value="clear"><button>Exclude auto</button></form>'
            label_status="Estimated worthless · automatic"
        elif status=="WAIT":
            action="—";label_status="Expiring today · monitoring"
        else:
            action=(f'<form method="post" action="/admin/intraday/{fund}/expiration" onsubmit="return confirm(\'Only confirm if no exercise/assignment for {label}.\');">{params}<input type="hidden" name="operation" value="confirm"><button>Confirm manually</button></form>')
            label_status="Review needed"
        exp_rows.append(f'<tr><td>{label}</td><td>{price}</td><td>{distance}</td><td>{html.escape(label_status)}<br><span class="muted">{detail}</span></td><td>{action}</td></tr>')
    sel=''.join(position_lines)
    return f'''<div class="card"><h3>Intraday Trade Management · {fund.title()} (v6.15.1)</h3>
<div class="notice"><strong>Baseline report: {html.escape(report or "UNKNOWN")}</strong>. Trades dated AFTER this baseline are added to Daily Return and Delta. Signed quantity +BUY / -SELL. Execution price per share/contract unit, NOT total premium. Record each fill. New uploads absorb already-reported trades without deleting history.</div>
<form method="post" action="/admin/intraday/{fund}/add" class="row">
<label>Action<select name="action" required><option value="OPEN">Open Position</option><option value="INCREASE">Increase Position</option><option value="REDUCE">Reduce Position</option><option value="CLOSE">Close Position</option></select></label>
<label>Existing (for Increase / Reduce / Close)<select name="existing"><option value="">— New instrument —</option>{sel}</select></label>
<label>Ticker<input name="ticker" placeholder="ORCL" maxlength="15"></label>
<label>Instrument<select name="instrument_type"><option value="OPTION">Option</option><option value="EQUITY">Equity</option></select></label>
<label>Expiry<input name="expiry" type="date"></label><label>Type<select name="option_type"><option value="P">Put</option><option value="C">Call</option></select></label>
<label>Strike<input name="strike" type="number" step="any" min="0"></label>
<label>Open Side<select name="side"><option value="SELL">Sell / Short</option><option value="BUY">Buy / Long</option></select></label>
<label>Contracts / Shares<input name="quantity" type="number" step="1" min="1" placeholder="Auto for Close"></label>
<label>Execution Price ($)<input name="execution_price" type="number" min="0" step="0.000001" required></label>
<label>Trade Date (NY)<input type="date" name="trade_date" value="{today}" required></label>
<label>Note<input name="note" maxlength="200"></label><button type="submit">Save Trade &amp; Publish Collector Config</button></form>
<p class="muted">For existing positions, select the instrument from the dropdown. For Open, fill the instrument fields. Quantity is absolute; Close automatically uses the entire remaining position. Default option multiplier: 100, equity: 1. Admin-only and audit logged.</p>
<h4>Recent trades</h4><div style="overflow-x:auto"><table><thead><tr><th>Trade Date</th><th>Instrument</th><th>Action</th><th>Signed Qty</th><th>Execution</th><th>Status</th><th>Correction</th></tr></thead><tbody>{''.join(trade_rows) or '<tr><td colspan="7">No intraday trades.</td></tr>'}</tbody></table></div>
<h4>Expiration Reconciliation · automatic OTM estimate (v6.15.1)</h4>
<p class="muted">On expiration day before the New York close: monitoring only. After the Bloomberg MARKET CLOSED snapshot at 16:00–16:45 ET, positions that are clearly OTM (at least 2% of the underlying price AND $0.10 from strike) are marked at estimated $0 automatically for <strong>both</strong> Daily Returns. Near-strike, ITM, missing or stale-price options are NOT automatically zeroed. This is an <strong>estimated</strong> return, NOT proof of no exercise/assignment. No Nirvana position or baseline is deleted. Admin can exclude auto for a position. The Bloomberg Collector must include underlying_spot and spot_source (v4.2).</p>
<table><thead><tr><th>Expiring Option</th><th>Bloomberg Underlying</th><th>OTM %</th><th>Estimate / Risk</th><th>Action</th></tr></thead><tbody>{''.join(exp_rows) or '<tr><td colspan="5">No expiration candidates after active baseline.</td></tr>'}</tbody></table></div>'''

@app.post("/admin/intraday/{fund}/add")
async def intraday_trade_add(request:Request,fund:str):
    user=require_admin(request)
    if fund not in ("fortress","opportunity"):raise HTTPException(404)
    form=await request.form()
    action=str(form.get("action") or "").upper()
    if action not in ("OPEN","INCREASE","REDUCE","CLOSE"):
        raise HTTPException(400,"Unsupported trade action")
    trade_date=str(form.get("trade_date") or "")
    try:dt=date.fromisoformat(trade_date)
    except Exception:raise HTTPException(400,"Invalid trade date")
    ny_today=_ny_now().date()
    if dt>ny_today or (ny_today-dt).days>7:
        raise HTTPException(400,"Trade date must be today or within 7 past calendar days (New York)")
    try:price=float(form.get("execution_price"))
    except Exception:raise HTTPException(400,"Execution price is required")
    if not math.isfinite(price) or price<0 or price>10000000:
        raise HTTPException(400,"Invalid execution price")
    with SessionLocal() as db:
        baseline_day,baseline=_base_for_trade(db,fund)
        if not baseline_day or trade_date<=baseline_day:
            raise HTTPException(400,"Trade date must be AFTER the active P&L baseline report date")
        # Existing positions are resolved by canonical instrument label, never by free-text guessing.
        existing=str(form.get("existing") or "").strip()
        active=_active_trades(db,fund,baseline_day,trade_date)
        effective=_overlay_positions(baseline,active)
        found=next((p for p in effective if _opp_position_label(p)==existing),None) if existing else None
        if action!="OPEN":
            if found is None:raise HTTPException(400,"Select a current position to increase, reduce or close")
            p=dict(found)
        else:
            ticker=str(form.get("ticker") or "").strip().upper()
            if not re.fullmatch(r"[A-Z][A-Z0-9.\-]{0,14}",ticker):
                raise HTTPException(400,"Invalid ticker")
            inst=str(form.get("instrument_type") or "OPTION").upper()
            if inst not in ("OPTION","EQUITY"):raise HTTPException(400,"Invalid instrument type")
            expiry=typ="";strike=None
            if inst=="OPTION":
                expiry=str(form.get("expiry") or "");typ=str(form.get("option_type") or "").upper()
                try:date.fromisoformat(expiry);strike=float(form.get("strike"))
                except Exception:raise HTTPException(400,"Option requires valid expiration and strike")
                if typ not in ("P","C") or not math.isfinite(strike) or strike<=0:
                    raise HTTPException(400,"Invalid option")
                if expiry<trade_date:raise HTTPException(400,"Option expiry precedes trade date")
            p={"ticker":ticker,"instrument_type":inst,"expiry":expiry,"option_type":typ,"strike":strike,
               "multiplier":100.0 if inst=="OPTION" else 1.0,"quantity":0.0}
            if any(_dr_position_key(x)==_dr_position_key(p) for x in effective):
                raise HTTPException(400,"Position already exists: use Increase instead of Open")
        current=float(p.get("quantity") or 0)
        if action=="OPEN" and current!=0:raise HTTPException(400,"Open requires zero current quantity")
        if action!="OPEN" and current==0:raise HTTPException(400,"Position is flat")
        if action=="CLOSE":qty=abs(current)
        else:
            try:qty=float(form.get("quantity"))
            except Exception:raise HTTPException(400,"Quantity required")
            if not qty.is_integer() or qty<1 or qty>1e9:raise HTTPException(400,"Quantity must be a positive integer")
        if action=="REDUCE" and qty>=abs(current):
            raise HTTPException(400,"Reduce cannot fully close or reverse: use Close")
        if action=="OPEN":
            side=str(form.get("side") or "SELL").upper()
            if side not in ("BUY","SELL"):raise HTTPException(400,"Invalid side")
            signed=qty if side=="BUY" else -qty
        elif action=="INCREASE":signed=qty if current>0 else -qty
        else:signed=-qty if current>0 else qty
        if not math.isfinite(qty) or abs(signed)>1e9:raise HTTPException(400,"Invalid size")
        # Do not retroactively reorder fills after a newer trade for this instrument.
        later=[t for t in _trade_history(db,fund) if t.get("status")=="ACTIVE"
               and t.get("trade_date","")>trade_date and _dr_position_key(t)==_dr_position_key(p)]
        if later:raise HTTPException(400,"A newer trade exists for this instrument. Enter trades in date order.")
        record={k:p.get(k) for k in ("ticker","instrument_type","expiry","option_type","strike","multiplier")}
        record.update({"id":uuid.uuid4().hex,"action":action,"trade_date":trade_date,
            "signed_quantity":int(signed),"execution_price":price,"status":"ACTIVE",
            "created_at":utcnow().isoformat(),"actor":user["username"],"note":str(form.get("note") or "")[:200]})
        history=_trade_history(db,fund)
        if len(history)>=2000:raise HTTPException(400,"Trade ledger limit reached: export/archive required")
        history.append(record)
        _set_json_setting(db,_ledger_key(fund),history)
        _version_bump_for_trade(db,fund,user["username"],f'Intraday {action}: {_opp_position_label(record)} {signed:+g} @ {price}')
        audit(db,user["username"],"INTRADAY_TRADE_ADDED",f"{fund}: {record['id']}; {action}; {_opp_position_label(record)}; qty {signed:+g}; price {price}; trade_date {trade_date}")
        db.commit()
    return RedirectResponse('/admin/fortress' if fund=='fortress' else '/admin/opportunity',303)

@app.post("/admin/intraday/{fund}/cancel")
async def intraday_trade_cancel(request:Request,fund:str):
    u=require_admin(request);form=await request.form();trade_id=str(form.get("trade_id") or "")
    with SessionLocal() as db:
        baseline_day,_=_base_for_trade(db,fund)
        history=_trade_history(db,fund)
        trade=next((t for t in history if t.get("id")==trade_id),None)
        if not trade or trade.get("status")!="ACTIVE":raise HTTPException(404,"Active trade not found")
        if trade.get("trade_date","")<=baseline_day:
            raise HTTPException(400,"Trade already absorbed by Nirvana baseline; cannot cancel. Correct it through the next report.")
        # Preserve trade-sequence consistency; undo later fills first.
        idx=history.index(trade)
        if any(x.get("status")=="ACTIVE" and x.get("trade_date","")>baseline_day
               and _dr_position_key(x)==_dr_position_key(trade) for x in history[idx+1:]):
            raise HTTPException(400,"Cancel later trades in this instrument first.")
        trade["status"]="CANCELLED";trade["cancelled_by"]=u["username"];trade["cancelled_at"]=utcnow().isoformat()
        _set_json_setting(db,_ledger_key(fund),history)
        _version_bump_for_trade(db,fund,u["username"],f'Cancelled intraday trade {trade_id}')
        audit(db,u["username"],"INTRADAY_TRADE_CANCELLED",f"{fund} {trade_id}");db.commit()
    return RedirectResponse('/admin/fortress' if fund=='fortress' else '/admin/opportunity',303)

@app.post("/admin/intraday/{fund}/expiration")
async def intraday_expiration(request:Request,fund:str):
    u=require_admin(request);form=await request.form()
    if fund not in ("fortress","opportunity"):raise HTTPException(404)
    expiry=str(form.get("expiry") or "");now=_ny_now()
    if not re.fullmatch(r"\d{4}-\d{2}-\d{2}",expiry) or expiry>now.date().isoformat() or (expiry==now.date().isoformat() and now.hour<16):
        raise HTTPException(400,"Expiration confirmation allowed only after 4 PM New York on expiry day")
    typ=str(form.get("option_type") or "").upper()
    try:strike=float(form.get("strike"))
    except Exception:raise HTTPException(400,"Invalid strike")
    ticker=str(form.get("ticker") or "").upper()
    if typ not in ("P","C") or not re.fullmatch(r"[A-Z][A-Z0-9.\-]{0,14}",ticker) or not math.isfinite(strike) or strike<=0:
        raise HTTPException(400,"Invalid option")
    operation=str(form.get("operation") or "")
    if operation not in ("confirm","clear","resume"):raise HTTPException(400,"Invalid operation")
    with SessionLocal() as db:
        baseline_day,pos=_base_for_trade(db,fund)
        record={"ticker":ticker,"instrument_type":"OPTION","expiry":expiry,"option_type":typ,"strike":strike}
        active=_active_trades(db,fund,baseline_day)
        effective=_overlay_positions(pos,active)
        if expiry<=baseline_day or not any(_dr_position_key(p)==_dr_position_key(record) for p in effective):
            raise HTTPException(400,"Option is not present in the active daily position set")
        old=_expiry_history(db,fund)
        old=[x for x in old if _dr_position_key(x)!=_dr_position_key(record)]
        old.append({**record,"status":"CONFIRMED_WORTHLESS" if operation=="confirm" else ("RESUMED" if operation=="resume" else "CLEARED"),
                    "actor":u["username"],"updated_at":utcnow().isoformat()})
        _set_json_setting(db,f"expiration_confirm_v615_{fund}",old)
        _version_bump_for_trade(db,fund,u["username"],f'Expiration {operation}: {_opp_position_label(record)}')
        audit(db,u["username"],"EXPIRATION_CONFIRMATION",f"{fund}: {operation} {_opp_position_label(record)}")
        db.commit()
    return RedirectResponse('/admin/fortress' if fund=='fortress' else '/admin/opportunity',303)
# === End v6.15 add-on ===

def _opp_limit(db):
    raw=str(get_setting(db,"opportunity_limit_pct","") or "").strip()
    try:return float(raw) if raw else None
    except Exception:return None

def _bump_opportunity_config(db,actor,reason):
    v=int(get_setting(db,"opportunity_config_version","1") or 1)+1
    set_setting(db,"opportunity_config_version",v)
    audit(db,actor,"OPPORTUNITY_CONFIG_PUBLISHED",f"v{v}: {reason}")
    return v

def _parse_opportunity_exposure_xlsx(data:bytes,filename:str="Exposure.xlsx"):
    """
    Opportunity-only Exposure by Underlying parser.

    Important:
    - This function is used ONLY for Opportunity Delta / Exposure import.
    - It does not change Fortress parsing.
    - It does not change Opportunity/Fortress Daily Return baselines, Bid/Ask
      logic, Yahoo snapshots, Bloomberg Daily Return, or collector behavior.
    """
    wb=load_workbook(BytesIO(data),read_only=True,data_only=True)

    # Nirvana exports can change the workbook's active sheet depending on
    # report filters. Prefer the named Exposure Summary sheet, then scan all
    # sheets as a fallback instead of trusting wb.active.
    sheets=[]
    for s in wb.worksheets:
        if str(s.title or "").strip().lower()=="exposure summary":
            sheets.insert(0,s)
        else:
            sheets.append(s)

    ws=None; vals=None; header_i=None
    required_headers={"symbol","nav","position","delta"}
    for candidate in sheets:
        candidate_vals=list(candidate.iter_rows(values_only=True))
        # Search generously because Nirvana may add/remove report metadata rows.
        for i,row in enumerate(candidate_vals[:60]):
            norm=[str(x or "").strip() for x in (row or [])]
            norm_lower={x.lower() for x in norm if x}
            if required_headers.issubset(norm_lower):
                ws=candidate; vals=candidate_vals; header_i=i
                break
        if header_i is not None:
            break

    if ws is None or vals is None or header_i is None:
        raise ValueError("Could not locate Opportunity Exposure Summary header row (Symbol/NAV/Position/Delta).")

    headers=[str(x or "").strip() for x in vals[header_i]]
    hm={h.lower():i for i,h in enumerate(headers) if h}
    for req in ("symbol","nav","position","delta"):
        if req not in hm:
            raise ValueError(f"Opportunity Exposure is missing required column: {req.title()}")

    report_date=""
    for row in vals[:header_i]:
        line=" ".join(str(x or "") for x in row)
        m=re.search(r"Report Date:\s*(\d{1,2}/\d{1,2}/\d{4})",line,re.I)
        if m:
            report_date=datetime.strptime(m.group(1),"%m/%d/%Y").date().isoformat();break

    nav=None; positions=[]; seen=set()

    # Cash / money-market instruments must NOT contribute to Opportunity
    # Bloomberg Delta Monitor. NAV is still the report's Grand Total NAV.
    # Keeping the denominator unchanged is intentional.
    excluded_symbols={"FGTXX","MPFXX","USD","CASH"}

    exp_col=hm.get("expiration date")

    for row in vals[header_i+1:]:
        try:
            sym=str(row[hm["symbol"]] or "").strip()
        except Exception:
            continue
        if not sym:
            continue

        if sym.lower().startswith("grand total"):
            try: nav=float(row[hm["nav"]])
            except Exception: nav=None
            continue

        # Opportunity Delta-only filter. This does not affect Daily Return.
        clean_sym=sym.upper().strip()
        if clean_sym in excluded_symbols:
            continue

        try: qty=float(row[hm["position"]])
        except Exception: continue
        try: delta=float(row[hm["delta"]])
        except Exception: continue
        if abs(qty)<1e-12:
            continue

        if sym.startswith("O:"):
            m=re.match(r"^O:([^ ]+)\s+(\d{2})([A-X])([0-9.]+)D(\d{1,2})$",sym,re.I)
            if not m:
                continue
            ticker=m.group(1).upper(); code=m.group(3).upper(); strike=float(m.group(4))
            typ="C" if "A"<=code<="L" else "P"
            if exp_col is None:
                continue
            expv=row[exp_col]
            if hasattr(expv,"date"):
                expiry=expv.date().isoformat()
            else:
                try: expiry=datetime.strptime(str(expv)[:10],"%Y-%m-%d").date().isoformat()
                except Exception: continue
            key=("OPTION",ticker,expiry,typ,round(strike,6),round(qty,6))
            if key in seen:
                continue
            seen.add(key)
            positions.append({
                "id":len(positions)+1,
                "instrument_type":"OPTION",
                "ticker":ticker,
                "expiry":expiry,
                "option_type":typ,
                "strike":strike,
                "quantity":qty,
                "multiplier":100.0,
                "bloomberg_security":option_security(ticker,expiry,typ,strike),
                "underlying_security":f"{ticker} US Equity",
                "nirvana_delta":delta
            })

        elif abs(delta-1.0)<1e-9:
            ticker=clean_sym
            # Only real equities/ETFs remain here. Money-market/cash symbols
            # were excluded above.
            key=("EQUITY",ticker,round(qty,6))
            if key in seen:
                continue
            seen.add(key)
            positions.append({
                "id":len(positions)+1,
                "instrument_type":"EQUITY",
                "ticker":ticker,
                "expiry":"",
                "option_type":"",
                "strike":None,
                "quantity":qty,
                "multiplier":1.0,
                "bloomberg_security":f"{ticker} US Equity",
                "underlying_security":f"{ticker} US Equity",
                "nirvana_delta":delta
            })

    if nav is None or nav<=0:
        raise ValueError("Could not read Grand Total NAV from Opportunity exposure report.")
    if not positions:
        raise ValueError("No detailed Opportunity exposure positions were parsed.")

    return {
        "filename":filename,
        "report_date":report_date,
        "nav_usd":nav,
        "positions":positions
    }

def _opp_baseline_positions_for_calc(b):
    return list((b or {}).get("positions") or [])

def _market_config_position_from_daily(p):
    inst=str(p.get("instrument_type") or "OPTION").upper(); ticker=str(p.get("ticker") or "").upper()
    if inst=="OPTION":
        return {"instrument_type":"OPTION","ticker":ticker,"expiry":p.get("expiry") or "","option_type":p.get("option_type") or "","strike":p.get("strike"),"quantity":p.get("quantity"),"multiplier":p.get("multiplier",100),"bloomberg_security":option_security(ticker,p.get("expiry") or "",p.get("option_type") or "P",p.get("strike")),"underlying_security":f"{ticker} US Equity"}
    return {"instrument_type":"EQUITY","ticker":ticker,"expiry":"","option_type":"","strike":None,"quantity":p.get("quantity"),"multiplier":1.0,"bloomberg_security":f"{ticker} US Equity","underlying_security":f"{ticker} US Equity"}

def _merge_market_positions(delta_positions,daily_positions):
    merged={}
    def key(p):
        inst=str(p.get("instrument_type") or "OPTION").upper()
        if inst=="OPTION":return ("OPTION",str(p.get("ticker") or "").upper(),str(p.get("expiry") or ""),str(p.get("option_type") or "").upper(),round(float(p.get("strike") or 0),6))
        return (inst,str(p.get("ticker") or "").upper())
    for p in delta_positions or []:
        x=dict(p);x["use_for_delta"]=True;x["use_for_daily_return"]=False;merged[key(x)]=x
    for p in daily_positions or []:
        x=_market_config_position_from_daily(p);k=key(x)
        if k in merged: merged[k]["use_for_daily_return"]=True
        else:
            x["use_for_delta"]=False;x["use_for_daily_return"]=True;merged[k]=x
    out=[]
    # Existing DB IDs stay untouched; synthetic intraday instruments get distinct IDs.
    ids=[int(x["id"]) for x in merged.values() if x.get("id") is not None]
    next_id=max(ids,default=0)+1
    for x in merged.values():
        x=dict(x)
        if x.get("id") is None:
            x["id"]=next_id;next_id+=1
        out.append(x)
    return out

def _report_export_links(fund_key,report_key):
    return (
        f'<div class="report-export-bar">'
        f'<span>Export:</span>'
        f'<a class="export-btn" href="/export/{fund_key}/{report_key}/xlsx">Excel</a>'
        f'<a class="export-btn" href="/export/{fund_key}/{report_key}/pdf">PDF</a>'
        f'</div>'
    )

def _export_position_label(p):
    inst=str(p.get("instrument_type") or "OPTION").upper()
    if inst=="OPTION":
        strike=p.get("strike")
        st="" if strike is None else f"{float(strike):g}"
        return f'{str(p.get("ticker") or "").upper()} {p.get("expiry") or ""} {str(p.get("option_type") or "").upper()}{st}'.strip()
    return str(p.get("ticker") or p.get("security_name") or "").upper()

def _fortress_export_baseline(db):
    b=_active_daily_baseline(db)
    if not b:
        return None,None
    dbpos=db.query(DailyReturnPosition).filter_by(baseline_id=b.id).order_by(DailyReturnPosition.id).all()
    positions=[{
        "id":p.id,
        "security_name":p.security_name,
        "instrument_type":p.instrument_type,
        "ticker":p.ticker,
        "expiry":p.expiry,
        "option_type":p.option_type,
        "strike":p.strike,
        "quantity":float(p.quantity),
        "multiplier":float(p.multiplier or 100),
        "baseline_price":float(p.baseline_price) if p.baseline_price is not None else None,
        "yahoo_manual_price":float(p.yahoo_manual_price) if p.yahoo_manual_price is not None else None,
        "bloomberg_manual_price":float(p.bloomberg_manual_price) if p.bloomberg_manual_price is not None else None,
    } for p in dbpos]
    baseline={
        "report_date":b.report_date,
        "run_date":b.run_date,
        "filename":b.filename,
        "nav_usd":float(b.nav_usd),
    }
    return baseline,positions

def _collect_delta_export(fund_key):
    with SessionLocal() as db:
        if fund_key=="fortress":
            snap,age=latest_snapshot(db)
            cfg=int(get_setting(db,"config_version","1"))
            limit=float(get_setting(db,"limit_pct","15"))
            if not snap:
                raise HTTPException(404,"No Fortress Bloomberg Delta snapshot is available.")
            raw_positions=snap.get("positions",[]) or []
            positions=sorted(raw_positions,key=lambda x:abs(float(x.get("contribution_pct") or 0)),reverse=True)
            buffer=snap.get("buffer_pct")
            title="RPD Fortress Fund - Bloomberg Delta Monitor"
        else:
            snap,age=_latest_opportunity_snapshot(db)
            cfg=int(get_setting(db,"opportunity_config_version","1") or 1)
            limit=_opp_limit(db)
            if not snap:
                raise HTTPException(404,"No Opportunity Bloomberg Delta snapshot is available.")
            raw_positions=[x for x in (snap.get("positions",[]) or []) if x.get("use_for_delta")]
            positions=sorted(raw_positions,key=lambda x:abs(float(x.get("contribution_pct") or 0)),reverse=True)
            buffer=snap.get("buffer_pct")
            title="RPD Opportunity Fund - Bloomberg Delta Monitor"

    mode=str(snap.get("mode") or "LIVE")
    closed=(mode=="MARKET_CLOSED")
    status="MARKET CLOSED" if closed else mode
    cov=snap.get("coverage",{}) or {}
    rows=[]
    for r in positions:
        rows.append([
            _export_position_label(r),
            float(r.get("quantity") or 0),
            r.get("delta"),
            r.get("underlying_spot"),
            r.get("multiplier"),
            float(r.get("delta_exposure_usd") or 0),
            float(r.get("contribution_pct") or 0)/100.0,
        ])

    summary=[
        ("Feed",status),
        ("Delta Exposure %",float(snap.get("delta_exposure_pct") or 0)/100.0),
        ("Delta Exposure $",float(snap.get("delta_exposure_usd") or 0)),
        ("Limit %",None if limit is None else float(limit)/100.0),
        ("Buffer %",None if buffer is None else float(buffer)/100.0),
        ("Coverage",f'{cov.get("valid",0)}/{cov.get("total",0)}'),
        ("Collector / Portal",f'{snap.get("config_version","-")} / {cfg}'),
        ("Snapshot (New York)",format_new_york_time(snap.get("timestamp_utc",""))),
        ("Collector",str(snap.get("collector_id") or "")),
    ]
    return {
        "title":title,
        "summary":summary,
        "headers":["Position","Qty","Bloomberg Delta","Underlying Price","Multiplier","Delta Exposure $","Contribution % NAV"],
        "rows":rows,
        "formats":["text","integer","raw4","price","integer","currency0","percent3"],
        "fund":fund_key,
        "report_key":"delta",
    }

def _collect_daily_export(fund_key,source_key):
    with SessionLocal() as db:
        if fund_key=="fortress":
            baseline,positions=_fortress_export_baseline(db)
            if not baseline:
                raise HTTPException(404,"No Fortress P&L baseline is available.")
            if source_key=="bloomberg":
                snap,_=latest_snapshot(db)
                calc=_daily_with_ledger(db,"fortress","bloomberg",positions,baseline["nav_usd"],snap,baseline["report_date"])
            else:
                snap=_latest_yahoo_collector_snapshot(db)
                calc=_daily_with_ledger(db,"fortress","yahoo",positions,baseline["nav_usd"],snap,baseline["report_date"])
            fund_name="RPD Fortress Fund"
        else:
            baseline=_opportunity_baseline(db)
            if not baseline:
                raise HTTPException(404,"No Opportunity P&L baseline is available.")
            positions=_opp_baseline_positions_for_calc(baseline)
            if source_key=="bloomberg":
                snap,_=_latest_opportunity_snapshot(db)
                calc=_daily_with_ledger(db,"opportunity","bloomberg",positions,float(baseline.get("nav_usd") or 0),snap,baseline.get("report_date"))
            else:
                snap=_latest_yahoo_v2_fund_snapshot(db,"opportunity")
                calc=_daily_with_ledger(db,"opportunity","yahoo",positions,float(baseline.get("nav_usd") or 0),snap,baseline.get("report_date"))
            fund_name="RPD Opportunity Fund"

    nav=float(baseline.get("nav_usd") or 0)
    status="FULL" if calc.get("complete") else "PARTIAL"
    source_name="Bloomberg" if source_key=="bloomberg" else "Yahoo"
    title=f"{fund_name} - {source_name} Daily Return"
    snapshot_ts=""
    if isinstance(snap,dict):
        snapshot_ts=str(snap.get("timestamp_utc") or snap.get("collected_at_utc") or "")

    summary=[
        ("Status",status),
        ("Estimated Daily Return",float(calc.get("estimated_return_pct") or 0)/100.0),
        ("Estimated P&L",float(calc.get("estimated_pnl") or 0)),
        ("Baseline NAV",nav),
        ("Baseline Report Date",str(baseline.get("report_date") or "")),
        ("Baseline File",str(baseline.get("filename") or "")),
        ("Coverage",f'{calc.get("coverage_valid",0)}/{calc.get("coverage_total",0)}'),
        ("Auto OTM estimate count",int(calc.get("auto_expiry_count") or 0)),
        ("OTM estimates note","PROVISIONAL — not an assignment confirmation" if calc.get("auto_expiry_count") else "None"),
        ("Snapshot (New York)",format_new_york_time(snapshot_ts) if snapshot_ts else "-"),
    ]

    rows=[]
    for r in calc.get("rows",[]) or []:
        if source_key=="bloomberg":
            bid=r.get("bid"); ask=r.get("ask"); market=r.get("market_mark")
            manual=r.get("bloomberg_manual_price")
            current=r.get("current_mark")
            effective_source=r.get("effective_source") or ""
        else:
            bid=r.get("bid"); ask=r.get("ask"); market=r.get("market_mark")
            manual=r.get("manual_price")
            current=r.get("current_mark")
            effective_source=r.get("effective_source") or ""

        rows.append([
            _export_position_label(r),
            float(r.get("quantity") or 0),
            r.get("baseline_price",r.get("previous_mark")),
            bid,
            ask,
            market,
            manual,
            current,
            r.get("change"),
            r.get("estimated_pnl"),
            None if r.get("contribution_pct") is None else float(r.get("contribution_pct"))/100.0,
            effective_source,
        ])

    headers=[
        "Position","Qty","Previous Mark",
        f"{source_name} Bid",f"{source_name} Ask",f"{source_name} Market",
        "Manual Override","Effective Price","Change",
        "Estimated P&L","Contribution","Effective Source",
    ]
    return {
        "title":title,
        "summary":summary,
        "headers":headers,
        "rows":rows,
        "formats":["text","integer","price","price","price","price","price","price","price","currency2","percent4","text"],
        "fund":fund_key,
        "report_key":f"{source_key}-daily-return",
    }

def _xlsx_report_bytes(report):
    wb=Workbook()
    ws=wb.active
    ws.title="Report"

    navy="17365D"; blue="D9EAF7"; light="EEF4FB"; white="FFFFFF"; gray="667085"
    from openpyxl.styles import Font,PatternFill,Alignment,Border,Side
    thin=Side(style="thin",color="D9E1EA")

    ncols=max(4,len(report["headers"]))
    ws.merge_cells(start_row=1,start_column=1,end_row=1,end_column=ncols)
    c=ws.cell(1,1,report["title"])
    c.font=Font(size=16,bold=True,color=navy)
    c.alignment=Alignment(vertical="center")
    ws.row_dimensions[1].height=26

    ny_now=datetime.now(ZoneInfo("America/New_York"))
    ws.merge_cells(start_row=2,start_column=1,end_row=2,end_column=ncols)
    c=ws.cell(2,1,f'Exported: {ny_now.strftime("%Y-%m-%d %I:%M:%S %p")} New York')
    c.font=Font(size=9,color=gray)

    row=4
    for label,value in report["summary"]:
        ws.cell(row,1,label).font=Font(bold=True,color=navy)
        ws.cell(row,2,value)
        if label.endswith("%") or label=="Estimated Daily Return":
            ws.cell(row,2).number_format="0.0000%"
        elif "$" in label or label in ("Estimated P&L","Baseline NAV"):
            ws.cell(row,2).number_format='$#,##0.00;[Red]-$#,##0.00'
        row+=1

    row+=1
    header_row=row
    for col,h in enumerate(report["headers"],1):
        cell=ws.cell(row,col,h)
        cell.fill=PatternFill("solid",fgColor=navy)
        cell.font=Font(bold=True,color=white)
        cell.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True)
        cell.border=Border(bottom=thin)
    ws.row_dimensions[row].height=28

    for data in report["rows"]:
        row+=1
        for col,val in enumerate(data,1):
            cell=ws.cell(row,col,val)
            cell.border=Border(bottom=thin)
            fmt=report["formats"][col-1]
            if fmt=="integer": cell.number_format='#,##0;[Red]-#,##0'
            elif fmt=="currency0": cell.number_format='$#,##0;[Red]-$#,##0'
            elif fmt=="currency2": cell.number_format='$#,##0.00;[Red]-$#,##0.00'
            elif fmt=="price": cell.number_format='$0.0000'
            elif fmt=="raw4": cell.number_format='0.0000'
            elif fmt=="percent3": cell.number_format='0.000%'
            elif fmt=="percent4": cell.number_format='0.0000%'
        if row%2==0:
            for col in range(1,ncols+1):
                ws.cell(row,col).fill=PatternFill("solid",fgColor="F8FAFC")

    widths=[32,12,16,14,14,16,16,16,14,18,16,22]
    for i in range(1,ncols+1):
        ws.column_dimensions[chr(64+i) if i<=26 else "A"].width=widths[i-1] if i<=len(widths) else 16
    # Use get_column_letter for correctness beyond Z.
    from openpyxl.utils import get_column_letter
    for i in range(1,ncols+1):
        ws.column_dimensions[get_column_letter(i)].width=widths[i-1] if i<=len(widths) else 16

    ws.freeze_panes=ws.cell(header_row+1,1)
    ws.auto_filter.ref=f"A{header_row}:{get_column_letter(ncols)}{row}"
    ws.sheet_view.showGridLines=False
    ws.page_setup.orientation="landscape"
    ws.page_setup.fitToWidth=1
    ws.sheet_properties.pageSetUpPr.fitToPage=True

    bio=BytesIO()
    wb.save(bio)
    return bio.getvalue()

def _pdf_safe_text(value):
    s="" if value is None else str(value)
    s=s.replace("—","-").replace("−","-").replace("×","x").replace("÷","/")
    return s.encode("latin-1","replace").decode("latin-1").replace("\\","\\\\").replace("(","\\(").replace(")","\\)")

def _pdf_fit(value,width,font_size):
    s="" if value is None else str(value)
    max_chars=max(3,int(width/(font_size*0.54)))
    return s if len(s)<=max_chars else s[:max_chars-1]+"~"

def _pdf_report_bytes(report):
    # Dependency-free landscape Letter PDF using standard Helvetica fonts.
    page_w,page_h=792.0,612.0
    margin=24.0
    font_size=6.7 if len(report["headers"])>6 else 8.5

    if len(report["headers"])==7:
        widths=[175,70,90,105,72,120,110]
    elif len(report["headers"])<=4:
        widths=[260,90,170,170]
    else:
        widths=[112,48,52,45,45,50,55,52,45,62,52,65]

    total=sum(widths)
    usable=page_w-2*margin
    scale=min(1.0,usable/total)
    widths=[w*scale for w in widths]

    def text_cmd(x,y,text,size=8,bold=False):
        font="/F2" if bold else "/F1"
        return f"BT {font} {size:.2f} Tf {x:.2f} {y:.2f} Td ({_pdf_safe_text(text)}) Tj ET\n"

    def line_cmd(x1,y1,x2,y2):
        return f"0.82 0.86 0.91 RG 0.5 w {x1:.2f} {y1:.2f} m {x2:.2f} {y2:.2f} l S\n"

    pages=[]
    rows=list(report["rows"])

    def make_page(page_no,start_index):
        cmds=[]
        y=page_h-margin
        cmds.append(text_cmd(margin,y,report["title"],15,True)); y-=18
        ny_now=datetime.now(ZoneInfo("America/New_York"))
        cmds.append(text_cmd(margin,y,f'Exported {ny_now.strftime("%Y-%m-%d %I:%M:%S %p")} New York',7.5,False)); y-=16

        # Summary in two columns.
        sx=[margin,margin+370]
        sy=y
        for i,(label,value) in enumerate(report["summary"]):
            col=i%2
            if i and col==0: sy-=13
            if isinstance(value,float):
                if label.endswith("%") or label=="Estimated Daily Return":
                    txt=f"{value*100:.4f}%"
                elif "$" in label or label in ("Estimated P&L","Baseline NAV"):
                    txt=f"${value:,.2f}"
                else:
                    txt=f"{value:,.4f}"
            else:
                txt=str(value)
            cmds.append(text_cmd(sx[col],sy,f"{label}: {txt}",7.5,False))
        y=sy-18
        cmds.append(line_cmd(margin,y,page_w-margin,y)); y-=14

        # Headers
        x=margin
        for h,w in zip(report["headers"],widths):
            cmds.append(text_cmd(x+2,y,_pdf_fit(h,w-4,font_size),font_size,True))
            x+=w
        y-=8
        cmds.append(line_cmd(margin,y,page_w-margin,y)); y-=12

        idx=start_index
        max_rows=30 if len(report["headers"])>6 else 34
        count=0
        while idx<len(rows) and count<max_rows and y>margin+18:
            data=rows[idx]
            x=margin
            for val,w,fmt in zip(data,widths,report["formats"]):
                if val is None:
                    txt="-"
                elif fmt=="integer":
                    txt=f"{float(val):,.0f}"
                elif fmt=="currency0":
                    txt=f"${float(val):,.0f}"
                elif fmt=="currency2":
                    txt=f"${float(val):+,.2f}"
                elif fmt=="price":
                    txt=f"${float(val):.4f}"
                elif fmt=="raw4":
                    txt=f"{float(val):+.4f}"
                elif fmt=="percent3":
                    txt=f"{float(val)*100:.3f}%"
                elif fmt=="percent4":
                    txt=f"{float(val)*100:+.4f}%"
                else:
                    txt=str(val)
                cmds.append(text_cmd(x+2,y,_pdf_fit(txt,w-4,font_size),font_size,False))
                x+=w
            y-=14
            cmds.append(line_cmd(margin,y+3,page_w-margin,y+3))
            idx+=1; count+=1

        cmds.append(text_cmd(page_w-margin-70,12,f"Page {page_no}",7,False))
        return "".join(cmds).encode("latin-1","replace"),idx

    next_idx=0; page_no=1
    while next_idx<len(rows) or page_no==1:
        content,next_idx=make_page(page_no,next_idx)
        pages.append(content)
        page_no+=1
        if next_idx>=len(rows): break

    # Build PDF objects.
    objects=[]
    objects.append(b"<< /Type /Catalog /Pages 2 0 R >>")
    # Pages object filled after page object ids known.
    objects.append(b"")
    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica-Bold >>")

    page_ids=[]
    content_ids=[]
    for content in pages:
        page_id=len(objects)+1
        content_id=page_id+1
        page_ids.append(page_id); content_ids.append(content_id)
        objects.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 {page_w:.0f} {page_h:.0f}] "
            f"/Resources << /Font << /F1 3 0 R /F2 4 0 R >> >> /Contents {content_id} 0 R >>".encode()
        )
        objects.append(b"<< /Length "+str(len(content)).encode()+b" >>\nstream\n"+content+b"\nendstream")

    kids=" ".join(f"{pid} 0 R" for pid in page_ids)
    objects[1]=f"<< /Type /Pages /Kids [{kids}] /Count {len(page_ids)} >>".encode()

    output=bytearray(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n")
    offsets=[0]
    for i,obj in enumerate(objects,1):
        offsets.append(len(output))
        output.extend(f"{i} 0 obj\n".encode())
        output.extend(obj)
        output.extend(b"\nendobj\n")
    xref=len(output)
    output.extend(f"xref\n0 {len(objects)+1}\n".encode())
    output.extend(b"0000000000 65535 f \n")
    for off in offsets[1:]:
        output.extend(f"{off:010d} 00000 n \n".encode())
    output.extend(
        f"trailer\n<< /Size {len(objects)+1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    )
    return bytes(output)


def _fund_reports(fund_key):
    return [
        _collect_delta_export(fund_key),
        _collect_daily_export(fund_key,"bloomberg"),
        _collect_daily_export(fund_key,"yahoo"),
    ]

def _xlsx_fund_report_bytes(fund_key):
    reports=_fund_reports(fund_key)
    from openpyxl.styles import Font,PatternFill,Alignment,Border,Side
    from openpyxl.utils import get_column_letter

    wb=Workbook()
    # remove default; create one sheet per report
    default=wb.active
    wb.remove(default)

    sheet_names=[
        "Bloomberg Delta",
        "Bloomberg Daily Return",
        "Yahoo Daily Return",
    ]
    navy="17365D"; white="FFFFFF"; gray="667085"
    thin=Side(style="thin",color="D9E1EA")

    for report,sheet_name in zip(reports,sheet_names):
        ws=wb.create_sheet(title=sheet_name)
        ncols=max(4,len(report["headers"]))
        ws.merge_cells(start_row=1,start_column=1,end_row=1,end_column=ncols)
        c=ws.cell(1,1,report["title"])
        c.font=Font(size=16,bold=True,color=navy)
        ws.row_dimensions[1].height=26

        ny_now=datetime.now(ZoneInfo("America/New_York"))
        ws.merge_cells(start_row=2,start_column=1,end_row=2,end_column=ncols)
        c=ws.cell(2,1,f'Exported: {ny_now.strftime("%Y-%m-%d %I:%M:%S %p")} New York')
        c.font=Font(size=9,color=gray)

        row=4
        for label,value in report["summary"]:
            ws.cell(row,1,label).font=Font(bold=True,color=navy)
            ws.cell(row,2,value)
            if label.endswith("%") or label=="Estimated Daily Return":
                ws.cell(row,2).number_format="0.0000%"
            elif "$" in label or label in ("Estimated P&L","Baseline NAV"):
                ws.cell(row,2).number_format='$#,##0.00;[Red]-$#,##0.00'
            row+=1

        row+=1
        header_row=row
        for col,h in enumerate(report["headers"],1):
            cell=ws.cell(row,col,h)
            cell.fill=PatternFill("solid",fgColor=navy)
            cell.font=Font(bold=True,color=white)
            cell.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True)
            cell.border=Border(bottom=thin)
        ws.row_dimensions[row].height=28

        for data in report["rows"]:
            row+=1
            for col,val in enumerate(data,1):
                cell=ws.cell(row,col,val)
                cell.border=Border(bottom=thin)
                fmt=report["formats"][col-1]
                if fmt=="integer": cell.number_format='#,##0;[Red]-#,##0'
                elif fmt=="currency0": cell.number_format='$#,##0;[Red]-$#,##0'
                elif fmt=="currency2": cell.number_format='$#,##0.00;[Red]-$#,##0.00'
                elif fmt=="price": cell.number_format='$0.0000'
                elif fmt=="percent3": cell.number_format='0.000%'
                elif fmt=="percent4": cell.number_format='0.0000%'
            if row%2==0:
                for col in range(1,ncols+1):
                    ws.cell(row,col).fill=PatternFill("solid",fgColor="F8FAFC")

        widths=[32,12,16,14,14,16,16,16,14,18,16,22]
        for i in range(1,ncols+1):
            ws.column_dimensions[get_column_letter(i)].width=widths[i-1] if i<=len(widths) else 16

        ws.freeze_panes=ws.cell(header_row+1,1)
        ws.auto_filter.ref=f"A{header_row}:{get_column_letter(ncols)}{row}"
        ws.sheet_view.showGridLines=False
        ws.page_setup.orientation="landscape"
        ws.page_setup.fitToWidth=1
        ws.sheet_properties.pageSetUpPr.fitToPage=True

    bio=BytesIO()
    wb.save(bio)
    return bio.getvalue()

def _pdf_fund_report_bytes(fund_key):
    # Reuse the already-tested per-report PDF generator and merge the
    # three sections into one PDF document.
    reports=_fund_reports(fund_key)
    writer=PdfWriter()
    for report in reports:
        part=_pdf_report_bytes(report)
        reader=PdfReader(BytesIO(part))
        for page in reader.pages:
            writer.add_page(page)
    bio=BytesIO()
    writer.write(bio)
    return bio.getvalue()

def _fund_export_buttons(fund_key):
    return (
        f'<div class="report-export-bar">'
        f'<span>Export Full Fund Report:</span>'
        f'<a class="export-btn" href="/export/{fund_key}/xlsx">Excel</a>'
        f'<a class="export-btn" href="/export/{fund_key}/pdf">PDF</a>'
        f'</div>'
    )

@app.get("/export/{fund_key}/{fmt}")
def export_full_fund_report(request:Request,fund_key:str,fmt:str):
    require_user(request)
    fund_key=str(fund_key).lower()
    fmt=str(fmt).lower()
    if fund_key not in ("fortress","opportunity"):
        raise HTTPException(404,"Unknown fund.")
    if fmt not in ("xlsx","pdf"):
        raise HTTPException(404,"Unknown export format.")

    date_txt=datetime.now(ZoneInfo("America/New_York")).strftime("%Y-%m-%d")
    fund_label="Fortress" if fund_key=="fortress" else "Opportunity"
    filename=f"RPD_{fund_label}_Full_Report_{date_txt}.{fmt}"

    if fmt=="xlsx":
        data=_xlsx_fund_report_bytes(fund_key)
        media="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    else:
        data=_pdf_fund_report_bytes(fund_key)
        media="application/pdf"

    return Response(
        content=data,
        media_type=media,
        headers={"Content-Disposition":f'attachment; filename="{filename}"'}
    )


@app.get("/export/{fund_key}/{report_key}/{fmt}")
def export_report(request:Request,fund_key:str,report_key:str,fmt:str):
    require_user(request)
    fund_key=str(fund_key).lower()
    report_key=str(report_key).lower()
    fmt=str(fmt).lower()
    if fund_key not in ("fortress","opportunity"):
        raise HTTPException(404,"Unknown fund.")
    if fmt not in ("xlsx","pdf"):
        raise HTTPException(404,"Unknown export format.")

    if report_key=="delta":
        report=_collect_delta_export(fund_key)
        report_label="Bloomberg_Delta"
    elif report_key=="bloomberg-daily-return":
        report=_collect_daily_export(fund_key,"bloomberg")
        report_label="Bloomberg_Daily_Return"
    elif report_key=="yahoo-daily-return":
        report=_collect_daily_export(fund_key,"yahoo")
        report_label="Yahoo_Daily_Return"
    else:
        raise HTTPException(404,"Unknown report.")

    date_txt=datetime.now(ZoneInfo("America/New_York")).strftime("%Y-%m-%d")
    fund_label="Fortress" if fund_key=="fortress" else "Opportunity"
    filename=f"RPD_{fund_label}_{report_label}_{date_txt}.{fmt}"

    if fmt=="xlsx":
        data=_xlsx_report_bytes(report)
        media="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    else:
        data=_pdf_report_bytes(report)
        media="application/pdf"
    return Response(content=data,media_type=media,headers={"Content-Disposition":f'attachment; filename="{filename}"'})


@app.get("/overview")
def overview(request:Request):
    u=require_user(request)
    with SessionLocal() as db:
        # Fortress delta
        fp,fage=latest_snapshot(db); fcfg=int(get_setting(db,"config_version","1")); flimit=float(get_setting(db,"limit_pct","15"))
        if fp:
            fm=fp.get("mode","LIVE"); fclosed=fm=="MARKET_CLOSED"; fstale=(fage>STALE_AFTER_SECONDS) and not fclosed
            fdelta_status="MARKET CLOSED" if fclosed else ("OFFLINE / STALE" if fstale else fm); fdelta_sub="Final closing snapshot" if fclosed else f"Age {fage:.1f}s"
            fdelta_pct=float(fp.get("delta_exposure_pct",0) or 0); fdelta_usd=float(fp.get("delta_exposure_usd",0) or 0); fcov=fp.get("coverage",{}) or {}; fdelta_cov=f'{fcov.get("valid",0)}/{fcov.get("total",0)}'; fcollector=fp.get("config_version","—")
            fbuffer=fp.get("buffer_pct"); fbuffer=float(fbuffer) if fbuffer is not None else flimit-fdelta_pct
        else:
            fdelta_status="WAITING";fdelta_sub="No collector snapshot";fdelta_pct=fdelta_usd=0.0;fdelta_cov="0/0";fcollector="—";fbuffer=flimit

        # Fortress daily summaries
        fb=_active_daily_baseline(db)
        if fb:
            fdb=db.query(DailyReturnPosition).filter_by(baseline_id=fb.id).order_by(DailyReturnPosition.id).all()
            fpos=[{"id":x.id,"security_name":x.security_name,"instrument_type":x.instrument_type,"ticker":x.ticker,"expiry":x.expiry,"option_type":x.option_type,"strike":x.strike,"quantity":x.quantity,"multiplier":x.multiplier,"baseline_price":x.baseline_price,"baseline_market_value":x.baseline_market_value,"yahoo_manual_price":x.yahoo_manual_price,"bloomberg_manual_price":x.bloomberg_manual_price} for x in fdb]
            fnav=float(fb.nav_usd or 0); freport=str(fb.report_date or "—")
            fys=_latest_yahoo_v2_fund_snapshot(db,"fortress") or _latest_yahoo_collector_snapshot(db)
            fycalc=_daily_with_ledger(db,"fortress","yahoo",fpos,fnav,fys,freport); fy_pnl=float(fycalc.get("estimated_pnl") or 0);fy_ret=float(fycalc.get("estimated_return_pct") or 0);fy_cov=f'{fycalc.get("coverage_valid",0)}/{fycalc.get("coverage_total",0)}';fy_status="FULL" if fycalc.get("complete") else "PARTIAL"
            fbcalc=_daily_with_ledger(db,"fortress","bloomberg",fpos,fnav,fp,freport);fb_pnl=float(fbcalc["estimated_pnl"]);fb_ret=float(fbcalc["estimated_return_pct"]);fb_cov=f'{fbcalc["coverage_valid"]}/{fbcalc["coverage_total"]}';fb_status="FULL" if fbcalc["complete"] else "PARTIAL"
        else:
            fnav=0;freport="—";fy_pnl=fy_ret=fb_pnl=fb_ret=0;fy_cov=fb_cov="0/0";fy_status=fb_status="NO BASELINE"

        # Opportunity delta
        op,oage=_latest_opportunity_snapshot(db); ocfg=int(get_setting(db,"opportunity_config_version","1") or 1); olimit=_opp_limit(db)
        if op:
            om=op.get("mode","LIVE"); oclosed=om=="MARKET_CLOSED"; ostale=(oage>STALE_AFTER_SECONDS) and not oclosed
            odelta_status="MARKET CLOSED" if oclosed else ("OFFLINE / STALE" if ostale else om); odelta_sub="Final closing snapshot" if oclosed else f"Age {oage:.1f}s"
            odelta_pct=float(op.get("delta_exposure_pct",0) or 0);odelta_usd=float(op.get("delta_exposure_usd",0) or 0);ocov=op.get("coverage",{}) or {};odelta_cov=f'{ocov.get("valid",0)}/{ocov.get("total",0)}';ocollector=op.get("config_version","—"); obuffer=op.get("buffer_pct")
            if obuffer is not None: obuffer=float(obuffer)
        else:
            odelta_status="WAITING";odelta_sub="No collector snapshot";odelta_pct=odelta_usd=0.0;odelta_cov="0/0";ocollector="—";obuffer=None

        ob=_opportunity_baseline(db)
        if ob:
            opos=_opp_baseline_positions_for_calc(ob);onav=float(ob.get("nav_usd") or 0);oreport=str(ob.get("report_date") or "—")
            oys=_latest_yahoo_v2_fund_snapshot(db,"opportunity"); oycalc=_daily_with_ledger(db,"opportunity","yahoo",opos,onav,oys,oreport);oy_pnl=float(oycalc.get("estimated_pnl") or 0);oy_ret=float(oycalc.get("estimated_return_pct") or 0);oy_cov=f'{oycalc.get("coverage_valid",0)}/{oycalc.get("coverage_total",0)}';oy_status="FULL" if oycalc.get("complete") else "PARTIAL"
            obcalc=_daily_with_ledger(db,"opportunity","bloomberg",opos,onav,op,oreport);ob_pnl=float(obcalc["estimated_pnl"]);ob_ret=float(obcalc["estimated_return_pct"]);ob_cov=f'{obcalc["coverage_valid"]}/{obcalc["coverage_total"]}';ob_status="FULL" if obcalc["complete"] else "PARTIAL"
        else:
            onav=0;oreport="—";oy_pnl=oy_ret=ob_pnl=ob_ret=0;oy_cov=ob_cov="0/0";oy_status=ob_status="NO BASELINE"

    def cls(v): return "ok" if v>=0 else "bad"
    def limit_txt(x): return "—" if x is None else f"{x:.3f}%"
    def buffer_txt(x): return "—" if x is None else f"{x:+.3f}%"
    def fund_column(title,accent,delta_link,dstatus,dsub,dpct,dusd,limit,buffer,dcov,collector,portalv,bd_link,bd_status,bd_ret,bd_pnl,nav,report,bd_cov,y_link,y_status,y_ret,y_pnl,y_cov):
        return f'''<section class="fundcol"><div class="fundtitle {accent}"><h2>{title}</h2></div>
        <div class="summarypanel"><div class="panelhead"><div class="headgroup"><h3>Bloomberg Delta Monitor</h3><span class="delay-badge bloomberg">◷&nbsp; 2-min delayed</span></div><a href="{delta_link}">Open full monitor →</a></div><div class="mini-grid">
        <div class="mini"><span>Feed</span><b class="ok">{html.escape(dstatus)}</b><small>{html.escape(dsub)}</small></div><div class="mini"><span>Delta Exposure</span><b>{dpct:.3f}%</b><small>${dusd:,.0f}</small></div><div class="mini"><span>Limit</span><b>{limit_txt(limit)}</b></div><div class="mini"><span>Buffer</span><b class="{cls(buffer or 0) if buffer is not None else ''}">{buffer_txt(buffer)}</b></div><div class="mini"><span>Coverage</span><b>{dcov}</b><small>Collector v{collector} / Portal v{portalv}</small></div></div></div>
        <div class="summarypanel"><div class="panelhead"><div class="headgroup"><h3>Bloomberg Daily Return</h3><span class="delay-badge bloomberg">◷&nbsp; 2-min delayed</span></div><a href="{bd_link}">Open details →</a></div><div class="mini-grid">
        <div class="mini"><span>Status</span><b class="ok">{bd_status}</b><small>Bloomberg marks</small></div><div class="mini"><span>Estimated Daily Return</span><b class="{cls(bd_ret)}">{bd_ret:+.4f}%</b></div><div class="mini"><span>Estimated P&amp;L</span><b class="{cls(bd_pnl)}">${bd_pnl:+,.2f}</b></div><div class="mini"><span>Baseline NAV</span><b>${nav:,.2f}</b><small>Report {html.escape(report)}</small></div><div class="mini"><span>Coverage</span><b>{bd_cov}</b><small>Bloomberg marks</small></div></div></div>
        <div class="summarypanel"><div class="panelhead"><div class="headgroup"><h3>Yahoo Daily Return</h3><span class="delay-badge yahoo">◷&nbsp; 15-min delayed</span></div><a href="{y_link}">Open details →</a></div><div class="mini-grid">
        <div class="mini"><span>Status</span><b class="ok">{y_status}</b><small>Mark-to-market estimate</small></div><div class="mini"><span>Estimated Daily Return</span><b class="{cls(y_ret)}">{y_ret:+.4f}%</b></div><div class="mini"><span>Estimated P&amp;L</span><b class="{cls(y_pnl)}">${y_pnl:+,.2f}</b></div><div class="mini"><span>Baseline NAV</span><b>${nav:,.2f}</b><small>Report {html.escape(report)}</small></div><div class="mini"><span>Coverage</span><b>{y_cov}</b><small>Yahoo local collector marks</small></div></div></div></section>'''

    body=f'''<style>.funds{{display:grid;grid-template-columns:1fr 1fr;gap:18px;align-items:stretch}}.fundcol{{min-width:0;display:grid;grid-template-rows:auto repeat(3,minmax(0,1fr));row-gap:14px;align-self:stretch}}.fundtitle{{padding:12px 18px;border-radius:12px 12px 0 0;margin-bottom:0}}.fundtitle h2{{margin:0;font-size:24px}}.fundtitle.fortress{{background:#eef6ff}}.fundtitle.opportunity{{background:#f5efff}}.summarypanel{{background:#fff;border:1px solid #dfe5ef;border-radius:13px;padding:14px;margin-bottom:0;display:flex;flex-direction:column;height:100%;box-sizing:border-box}}.panelhead{{display:flex;justify-content:space-between;align-items:center;gap:10px;margin-bottom:12px}}.headgroup{{display:flex;align-items:center;gap:9px;min-width:0;flex-wrap:wrap}}.panelhead h3{{margin:0;font-size:19px}}.panelhead a{{font-size:13px;white-space:nowrap}}.delay-badge{{display:inline-flex;align-items:center;white-space:nowrap;border-radius:999px;padding:4px 10px;font-size:11.5px;font-weight:600;line-height:1;border:1px solid transparent}}.delay-badge.bloomberg{{background:#eef3fb;color:#2d4771;border-color:#dde7f5}}.delay-badge.yahoo{{background:#fff1c9;color:#805b00;border-color:#f6dda0}}.mini-grid{{display:grid;grid-template-columns:.95fr 1fr 1.18fr 1.25fr .95fr;gap:9px;flex:1;align-items:stretch}}.mini{{border:1px solid #e4e9f1;border-radius:10px;padding:10px;min-height:98px;min-width:0;overflow:hidden;display:flex;flex-direction:column;height:100%;box-sizing:border-box}}.mini span{{display:block;color:#73819a;font-size:12px;line-height:1.25;min-height:30px}}.mini small{{display:block;color:#73819a;font-size:12px;line-height:1.25}}.mini b{{display:block;font-size:14.5px;line-height:1.15;margin:6px 0 4px;white-space:nowrap;letter-spacing:-.35px;font-variant-numeric:tabular-nums}}.mini:first-child b{{white-space:normal;font-size:14px}}.mini:nth-child(3) b,.mini:nth-child(4) b{{font-size:14px}}@media(min-width:1700px){{.mini b{{font-size:15.5px}}.mini:first-child b{{font-size:15px}}.mini:nth-child(3) b,.mini:nth-child(4) b{{font-size:15px}}}}@media(max-width:1250px){{.funds{{grid-template-columns:1fr}}.mini-grid{{grid-template-columns:.95fr 1fr 1.18fr 1.25fr .95fr}}.mini b{{font-size:18px}}.mini:first-child b{{font-size:17px}}.mini:nth-child(3) b,.mini:nth-child(4) b{{font-size:17px}}}}@media(max-width:800px){{.mini-grid{{grid-template-columns:1fr 1fr}}.mini b,.mini:first-child b,.mini:nth-child(3) b,.mini:nth-child(4) b{{font-size:17px}}.delay-badge{{font-size:11px;padding:4px 8px}}}}</style>
    <div class="funds">
    {fund_column("RPD Opportunity Fund","opportunity","/opportunity",odelta_status,odelta_sub,odelta_pct,odelta_usd,olimit,obuffer,odelta_cov,ocollector,ocfg,"/opportunity/bloomberg-daily-return",ob_status,ob_ret,ob_pnl,onav,oreport,ob_cov,"/opportunity/yahoo-daily-return",oy_status,oy_ret,oy_pnl,oy_cov)}
    {fund_column("RPD Fortress Fund","fortress","/",fdelta_status,fdelta_sub,fdelta_pct,fdelta_usd,flimit,fbuffer,fdelta_cov,fcollector,fcfg,"/bloomberg-daily-return",fb_status,fb_ret,fb_pnl,fnav,freport,fb_cov,"/daily-return",fy_status,fy_ret,fy_pnl,fy_cov)}
    </div>'''

    return page("Overview",body,u,60)


def _delta_input_cell(value,source,kind="number"):
    if value is None:
        main="—"
    else:
        try:
            v=float(value)
            if kind=="price":
                main=f"${v:,.2f}"
            elif kind=="delta":
                main=f"{v:+.4f}"
            elif kind=="multiplier":
                main=f"{v:,.0f}"
            else:
                main=f"{v:,.4f}"
        except Exception:
            main=html.escape(str(value))
    src=html.escape(str(source or ""))
    sub=f'<div class="muted" style="font-size:10px;margin-top:2px">{src}</div>' if src else ""
    return f'<div>{main}</div>{sub}'

@app.get("/")
def dashboard(request:Request):
    u=current_user(request)
    if not u:return RedirectResponse("/login",303)
    with SessionLocal() as db:
        p,age=latest_snapshot(db);cfg=int(get_setting(db,"config_version","1"));limit=float(get_setting(db,"limit_pct","15"))
        if not p:
            return page("Delta Monitor",f'<div class="card"><h1>RPD Fortress Fund — Bloomberg Delta Monitor</h1><div class="notice">WAITING — No collector snapshot yet.</div><p>Published config v{cfg}</p></div>',u,5)
        mode=p.get("mode","LIVE")
        closed_mode=(mode=="MARKET_CLOSED")
        stale=(age>STALE_AFTER_SECONDS) and not closed_mode
        status="MARKET CLOSED" if closed_mode else ("OFFLINE / STALE" if stale else mode)
        pct=p.get("delta_exposure_pct",0) or 0;buf=p.get("buffer_pct",0) or 0;exp=p.get("delta_exposure_usd",0) or 0;cov=p.get("coverage",{})
        rows=""
        for r in sorted(p.get("positions",[]),key=lambda x:abs(x.get("contribution_pct",0)),reverse=True):
            rows+=f'<tr><td>{html.escape(r.get("ticker",""))} {html.escape(r.get("expiry",""))} {html.escape(r.get("option_type",""))}{r.get("strike")}</td><td>{r.get("quantity",0):,.0f}</td><td>{_delta_input_cell(r.get("delta"),r.get("delta_source"),"delta")}</td><td>{_delta_input_cell(r.get("underlying_spot"),r.get("spot_source"),"price")}</td><td>{_delta_input_cell(r.get("multiplier"),None,"multiplier")}</td><td>${r.get("delta_exposure_usd",0):,.0f}</td><td>{r.get("contribution_pct",0):.3f}%</td></tr>'
        mismatch=p.get("config_version")!=cfg
        body=f"""<h1>RPD Fortress Fund — Bloomberg Delta Monitor</h1>{_fund_export_buttons("fortress")}<div class="grid">
<div class="card metric"><div class="muted">Feed</div><div class="v {'bad' if stale else 'ok'}">{status}</div><div class="muted">{'Final closing snapshot' if closed_mode else f'Age {age:.1f}s'}</div></div>
<div class="card metric"><div class="muted">Delta Exposure</div><div class="v">{pct:.3f}%</div><div class="muted">${exp:,.0f}</div></div>
<div class="card metric"><div class="muted">Limit</div><div class="v">{limit:.3f}%</div></div>
<div class="card metric"><div class="muted">Buffer</div><div class="v">{buf:.3f}%</div></div>
<div class="card metric"><div class="muted">Coverage</div><div class="v">{cov.get("valid",0)}/{cov.get("total",0)}</div><div class="muted">Collector v{p.get("config_version")} / Portal v{cfg}</div></div></div>
{('<div class="notice">A new portfolio config is published; collector has not reported it yet.</div>' if mismatch else '')}
<div class="card"><div class="muted">Last update (New York)</div>{html.escape(format_new_york_time(p.get("timestamp_utc","")))} · Collector: {html.escape(p.get("collector_id",""))}</div>
<div class="card"><div class="muted" style="margin-bottom:10px"><strong>Formula:</strong> Qty × Multiplier × Bloomberg Delta × Underlying Price = Delta Exposure $</div><table><thead><tr><th>Position</th><th>Qty</th><th>Bloomberg Delta</th><th>Underlying Price</th><th>Multiplier</th><th>Delta Exposure $</th><th>Contribution % NAV</th></tr></thead><tbody>{rows}</tbody></table></div>"""
        return page("Delta Monitor",body,u,3)


@app.get("/opportunity")
def opportunity_dashboard(request:Request):
    u=require_user(request)
    with SessionLocal() as db:
        p,age=_latest_opportunity_snapshot(db);cfg=int(get_setting(db,"opportunity_config_version","1") or 1);limit=_opp_limit(db)
        if not p:
            return page("Opportunity Bloomberg Delta Monitor",f'<div class="card"><h1>RPD Opportunity Fund — Bloomberg Delta Monitor</h1><div class="notice">WAITING — No Opportunity collector snapshot yet.</div><p>Published config v{cfg}</p></div>',u,5)
        mode=p.get("mode","LIVE");closed=mode=="MARKET_CLOSED";stale=(age>STALE_AFTER_SECONDS) and not closed;status="MARKET CLOSED" if closed else ("OFFLINE / STALE" if stale else mode)
        pct=float(p.get("delta_exposure_pct",0) or 0);exp=float(p.get("delta_exposure_usd",0) or 0);buf=p.get("buffer_pct");cov=p.get("coverage",{}) or {};rows=""
        for r in sorted([x for x in (p.get("positions",[]) or []) if x.get("use_for_delta")],key=lambda x:abs(float(x.get("contribution_pct") or 0)),reverse=True):
            inst=str(r.get("instrument_type") or "OPTION").upper(); label=(f'{r.get("ticker","")} {r.get("expiry","")} {r.get("option_type","")}{r.get("strike")}' if inst=="OPTION" else str(r.get("ticker") or ""))
            rows+=f'<tr><td>{html.escape(label)}</td><td>{float(r.get("quantity") or 0):,.0f}</td><td>{_delta_input_cell(r.get("delta"),r.get("delta_source"),"delta")}</td><td>{_delta_input_cell(r.get("underlying_spot"),r.get("spot_source"),"price")}</td><td>{_delta_input_cell(r.get("multiplier"),None,"multiplier")}</td><td>${float(r.get("delta_exposure_usd") or 0):,.0f}</td><td>{float(r.get("contribution_pct") or 0):.3f}%</td></tr>'
        limit_html='—' if limit is None else f'{limit:.3f}%';buf_html='—' if buf is None else f'{float(buf):+.3f}%'
        body=f'''<h1>RPD Opportunity Fund — Bloomberg Delta Monitor</h1>{_fund_export_buttons("opportunity")}<div class="grid"><div class="card metric"><div class="muted">Feed</div><div class="v {'bad' if stale else 'ok'}">{status}</div><div class="muted">{'Final closing snapshot' if closed else f'Age {age:.1f}s'}</div></div><div class="card metric"><div class="muted">Delta Exposure</div><div class="v">{pct:.3f}%</div><div class="muted">${exp:,.0f}</div></div><div class="card metric"><div class="muted">Limit</div><div class="v">{limit_html}</div></div><div class="card metric"><div class="muted">Buffer</div><div class="v">{buf_html}</div></div><div class="card metric"><div class="muted">Coverage</div><div class="v">{cov.get('valid',0)}/{cov.get('total',0)}</div><div class="muted">Collector v{p.get('config_version')} / Portal v{cfg}</div></div></div><div class="card"><div class="muted">Last update (New York)</div>{html.escape(format_new_york_time(p.get('timestamp_utc','')))} · Collector: {html.escape(str(p.get('collector_id','')))}</div><div class="card"><div class="muted" style="margin-bottom:10px"><strong>Formula:</strong> Qty × Multiplier × Bloomberg Delta × Underlying Price = Delta Exposure $</div><table><thead><tr><th>Position</th><th>Qty</th><th>Bloomberg Delta</th><th>Underlying Price</th><th>Multiplier</th><th>Delta Exposure $</th><th>Contribution % NAV</th></tr></thead><tbody>{rows}</tbody></table></div>'''
        return page("Opportunity Bloomberg Delta Monitor",body,u,3)

def _render_opportunity_yahoo(request:Request):
    u=require_user(request)
    with SessionLocal() as db:
        b=_opportunity_baseline(db);snap=_latest_yahoo_v2_fund_snapshot(db,"opportunity")
    if not b:
        extra=' <a href="/admin/opportunity">Upload Opportunity P&amp;L Report</a>' if u["username"]==ADMIN_USER else ''
        return page("Opportunity Yahoo Daily Return",f'<h1>RPD Opportunity Fund — Yahoo Daily Return</h1><div class="card"><div class="notice">No Opportunity P&amp;L baseline loaded.{extra}</div></div>',u)
    positions=_opp_baseline_positions_for_calc(b)
    with SessionLocal() as db: calc=_daily_with_ledger(db,"opportunity","yahoo",positions,float(b.get("nav_usd") or 0),snap,b.get("report_date"))
    status="FULL" if calc["complete"] else "PARTIAL";ret=float(calc.get("estimated_return_pct") or 0);pnl=float(calc.get("estimated_pnl") or 0);rows=""
    for r in calc["rows"]:
        strike="" if r.get("strike") is None else f'{r["strike"]:g}';posname=(f'{r["ticker"]} {r["expiry"]} {r["option_type"]}{strike}' if r["instrument_type"]=="OPTION" else r["ticker"]);manual=r.get("manual_price");manual_val='' if manual is None else f'{float(manual):.4f}'
        if u["username"]==ADMIN_USER and r.get("id") is not None:
            mh=f'<form method="post" action="/admin/opportunity/manual-price" style="display:flex;gap:4px"><input type="hidden" name="position_id" value="{r["id"]}"><input type="hidden" name="source" value="YAHOO"><input type="number" step="0.0001" min="0" name="price" value="{manual_val}" style="width:82px"><button>Save</button><button name="clear" value="1">Clear</button></form>'
        else: mh='—' if not manual_val else f'${float(manual_val):.4f}'
        fmt=lambda v:'—' if v is None else f'${float(v):.2f}'
        rows+=f'<tr><td>{html.escape(posname)}</td><td>{float(r["quantity"]):,.0f}</td><td>{fmt(r.get("previous_mark"))}</td><td>{fmt(r.get("bid"))}</td><td>{fmt(r.get("ask"))}</td><td>{fmt(r.get("market_mark"))}</td><td>{mh}</td><td>{fmt(r.get("current_mark"))}</td><td>{"—" if r.get("change") is None else f"{r["change"]:+.2f}"}</td><td>{"—" if r.get("estimated_pnl") is None else f"${r["estimated_pnl"]:+,.2f}"}</td><td>{"—" if r.get("contribution_pct") is None else f"{r["contribution_pct"]:+.4f}%"}</td><td>{html.escape(r.get("effective_source") or "")}</td></tr>'
    warn='' if calc["complete"] else '<div class="notice"><strong>PARTIAL:</strong> Yahoo local collector has not supplied a usable current mark for every Opportunity position.</div><br>'
    body=f'''<h1>RPD Opportunity Fund — Yahoo Daily Return</h1>{_fund_export_buttons("opportunity")}<div class="grid"><div class="card metric"><div class="muted">Status</div><div class="v {'ok' if calc['complete'] else 'bad'}">{status}</div></div><div class="card metric"><div class="muted">Estimated Daily Return</div><div class="v">{ret:+.4f}%</div></div><div class="card metric"><div class="muted">Estimated P&amp;L</div><div class="v">${pnl:+,.2f}</div></div><div class="card metric"><div class="muted">Baseline NAV</div><div class="v">${float(b.get('nav_usd') or 0):,.2f}</div><div class="muted">Report {html.escape(str(b.get('report_date') or ''))}</div></div><div class="card metric"><div class="muted">Coverage</div><div class="v">{calc['coverage_valid']}/{calc['coverage_total']}</div><div class="muted">Yahoo local collector marks</div></div></div>{warn}<div class="card"><strong>{html.escape(str(b.get('filename') or 'Opportunity PNL'))}</strong> · Formula: (Current Mark − Previous P&amp;L Report Price) × Quantity × Multiplier. v6.15 intraday adjustments: Signed Trade Qty × (Current Mark − Execution Price) × Multiplier; confirmed worthless and automatically estimated OTM expirations use a zero mark.</div><div class="card"><table><thead><tr><th>Position</th><th>Qty</th><th>Previous Mark</th><th>Yahoo Bid</th><th>Yahoo Ask</th><th>Yahoo Market</th><th>Manual Override</th><th>Effective Price</th><th>Change</th><th>Estimated P&amp;L</th><th>Contribution</th><th>Effective Source</th></tr></thead><tbody>{rows}</tbody></table></div>'''
    return page("Opportunity Yahoo Daily Return",body,u,60 if auto_refresh_allowed() else None)

@app.get("/opportunity/yahoo-daily-return")
def opportunity_yahoo_daily_return(request:Request): return _render_opportunity_yahoo(request)

@app.get("/opportunity/bloomberg-daily-return")
def opportunity_bloomberg_daily_return(request:Request):
    u=require_user(request)
    with SessionLocal() as db:
        b=_opportunity_baseline(db);snap,age=_latest_opportunity_snapshot(db)
    if not b:
        extra=' <a href="/admin/opportunity">Upload Opportunity P&amp;L Report</a>' if u["username"]==ADMIN_USER else ''
        return page("Opportunity Bloomberg Daily Return",f'<h1>RPD Opportunity Fund — Bloomberg Daily Return</h1><div class="card"><div class="notice">No Opportunity P&amp;L baseline loaded.{extra}</div></div>',u)
    pos=_opp_baseline_positions_for_calc(b)
    with SessionLocal() as db: calc=_daily_with_ledger(db,"opportunity","bloomberg",pos,float(b.get("nav_usd") or 0),snap,b.get("report_date"))
    ret=float(calc["estimated_return_pct"]);pnl=float(calc["estimated_pnl"]);rows=""
    for r in calc["rows"]:
        strike="" if r.get("strike") is None else f'{r["strike"]:g}';name=(f'{r["ticker"]} {r["expiry"]} {r["option_type"]}{strike}' if r["instrument_type"]=="OPTION" else r["ticker"]);manual=r.get("bloomberg_manual_price");mv='' if manual is None else f'{float(manual):.4f}'
        if u["username"]==ADMIN_USER and r.get("id") is not None:
            mh=f'<form method="post" action="/admin/opportunity/manual-price" style="display:flex;gap:4px"><input type="hidden" name="position_id" value="{r["id"]}"><input type="hidden" name="source" value="BLOOMBERG"><input type="number" step="0.0001" min="0" name="price" value="{mv}" style="width:82px"><button>Save</button><button name="clear" value="1">Clear</button></form>'
        else: mh='—' if not mv else f'${float(mv):.4f}'
        fmt=lambda v:'—' if v is None else f'${float(v):.2f}'
        rows+=f'<tr><td>{html.escape(name)}</td><td>{float(r["quantity"]):,.0f}</td><td>{fmt(r.get("baseline_price"))}</td><td>{fmt(r.get("bid"))}</td><td>{fmt(r.get("ask"))}</td><td>{fmt(r.get("market_mark"))}</td><td>{mh}</td><td>{fmt(r.get("current_mark"))}</td><td>{"—" if r.get("change") is None else f"{r["change"]:+.2f}"}</td><td>{"—" if r.get("estimated_pnl") is None else f"${r["estimated_pnl"]:+,.2f}"}</td><td>{"—" if r.get("contribution_pct") is None else f"{r["contribution_pct"]:+.4f}%"}</td><td>{html.escape(r.get("effective_source") or "")}</td></tr>'
    status="FULL" if calc["complete"] else "PARTIAL";warn='' if calc["complete"] else '<div class="notice"><strong>PARTIAL:</strong> Bloomberg has not supplied a usable current mark for every Opportunity position.</div><br>'
    body=f'''<h1>RPD Opportunity Fund — Bloomberg Daily Return</h1>{_fund_export_buttons("opportunity")}<div class="grid"><div class="card metric"><div class="muted">Status</div><div class="v {'ok' if calc['complete'] else 'bad'}">{status}</div></div><div class="card metric"><div class="muted">Estimated Daily Return</div><div class="v">{ret:+.4f}%</div></div><div class="card metric"><div class="muted">Estimated P&amp;L</div><div class="v">${pnl:+,.2f}</div></div><div class="card metric"><div class="muted">Baseline NAV</div><div class="v">${float(b.get('nav_usd') or 0):,.2f}</div><div class="muted">Report {html.escape(str(b.get('report_date') or ''))}</div></div><div class="card metric"><div class="muted">Coverage</div><div class="v">{calc['coverage_valid']}/{calc['coverage_total']}</div><div class="muted">Bloomberg marks</div></div></div>{warn}<div class="card"><div class="muted">Collector snapshot</div>{html.escape(format_new_york_time((snap or {}).get('timestamp_utc','')))}</div><div class="card"><table><thead><tr><th>Position</th><th>Qty</th><th>Previous Mark</th><th>Bloomberg Bid</th><th>Bloomberg Ask</th><th>Bloomberg Market</th><th>Manual Override</th><th>Effective Price</th><th>Change</th><th>Estimated P&amp;L</th><th>Contribution</th><th>Effective Source</th></tr></thead><tbody>{rows}</tbody></table></div>'''
    return page("Opportunity Bloomberg Daily Return",body,u,60 if auto_refresh_allowed() else None)

@app.get("/admin/opportunity")
def opportunity_admin(request:Request):
    u=require_admin(request)
    with SessionLocal() as db:
        exp=_opportunity_exposure(db);base=_opportunity_baseline(db);limit=get_setting(db,"opportunity_limit_pct","");ver=get_setting(db,"opportunity_config_version","1")
        hist=_opp_exposure_history(db,exp)
        exp_changes=_json_setting(db,"opportunity_exposure_last_changes",None)
        pnl_changes=_json_setting(db,"opportunity_pnl_last_changes",None)
    exp_txt='<p>No Opportunity Exposure report loaded.</p>' if not exp else f'<div class="notice"><strong>Exposure:</strong> {html.escape(str(exp.get("filename") or ""))} · Report {html.escape(str(exp.get("report_date") or ""))} · NAV ${float(exp.get("nav_usd") or 0):,.2f} · Delta positions {len(exp.get("positions") or [])}</div>'
    base_txt='<p>No Opportunity P&amp;L baseline loaded.</p>' if not base else f'<div class="notice"><strong>P&amp;L baseline:</strong> {html.escape(str(base.get("filename") or ""))} · Report {html.escape(str(base.get("report_date") or ""))} · NAV ${float(base.get("nav_usd") or 0):,.2f} · Daily Return positions {len(base.get("positions") or [])}</div>'
    exp_changes_html=_opp_changes_html(exp_changes,"Changes Detected — Exposure Upload")
    pnl_changes_html=_opp_changes_html(pnl_changes,"Changes Detected — P&amp;L Upload")
    portfolio_html=_opp_portfolio_table(hist);pnl_positions_html=_opp_pnl_positions_table(base)
    body=admin_tabs()+f"""<h1>Opportunity Setup</h1>
<div class="card"><h3>Opportunity Delta Monitor</h3><p>Upload the latest Nirvana Exposure by Underlying XLSX. This updates only Opportunity.</p><form method="post" action="/admin/opportunity/exposure-upload" enctype="multipart/form-data"><input type="file" name="file" accept=".xlsx" required><button>Upload Opportunity Exposure</button></form><br>{exp_txt}<br>{exp_changes_html}</div>
<div class="card"><h3>Opportunity Portfolio</h3><p class="muted">Upload-driven history. ACTIVE positions are in the latest Exposure file; positions no longer present are ARCHIVED, and past-expiry options are shown as EXPIRED. This is display/history only and does not change Delta calculations.</p>{portfolio_html}</div>
<div class="card"><h3>Opportunity Daily Return Baseline</h3><p>Upload the previous trading day's Opportunity PNL Report as PDF or Excel (XLSX). It supplies positions, previous marks and NAV for both Bloomberg and Yahoo Daily Return. Excel is recommended when available.</p><form method="post" action="/admin/opportunity/pnl-upload" enctype="multipart/form-data"><input type="file" name="file" accept=".pdf,.xlsx,.xlsm,application/pdf,application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" required><button>Upload Opportunity P&amp;L</button></form><br>{base_txt}<br>{pnl_changes_html}<br><h4>Daily Return Baseline Positions</h4>{pnl_positions_html}</div>
<div class="card"><h3>Opportunity Settings — Config v{html.escape(str(ver))}</h3><form method="post" action="/admin/opportunity/settings" class="row"><label>Delta Limit % (optional)<input name="limit_pct" value="{html.escape(str(limit or ''))}" placeholder="Leave blank if no limit"></label><button>Publish</button></form></div>"""
    with SessionLocal() as db:
        body+=_intraday_panel(db,"opportunity")
    return page("Opportunity Setup",body,u)

@app.post("/admin/opportunity/exposure-upload")
async def opportunity_exposure_upload(request:Request,file:UploadFile=File(...)):
    u=require_admin(request);filename=file.filename or "Opportunity Exposure.xlsx"
    if not filename.lower().endswith(".xlsx"): raise HTTPException(400,"Please upload the Opportunity Exposure by Underlying XLSX.")
    data=await file.read()
    try: parsed=_parse_opportunity_exposure_xlsx(data,filename)
    except Exception as e: raise HTTPException(400,f"Could not parse Opportunity Exposure: {e}")
    with SessionLocal() as db:
        changes=_opp_update_exposure_history(db,parsed,filename);_set_json_setting(db,"opportunity_exposure_config",parsed);audit(db,u["username"],"OPPORTUNITY_EXPOSURE_CHANGES",f'added={len(changes.get("added") or [])}; archived={len(changes.get("archived") or [])}; expired={len(changes.get("expired") or [])}; qty_changed={len(changes.get("quantity_changed") or [])}; reactivated={len(changes.get("reactivated") or [])}');_bump_opportunity_config(db,u["username"],f'Exposure {filename}; report={parsed.get("report_date")}; positions={len(parsed.get("positions") or [])}');db.commit()
    return RedirectResponse("/admin/opportunity",303)

@app.post("/admin/opportunity/pnl-upload")
async def opportunity_pnl_upload(request:Request,file:UploadFile=File(...)):
    u=require_admin(request);filename=file.filename or "Opportunity PNL.pdf"
    if not filename.lower().endswith((".pdf",".xlsx",".xlsm")): raise HTTPException(400,"Please upload the Opportunity PNL Report as PDF or XLSX.")
    data=await file.read()
    if len(data)>12*1024*1024: raise HTTPException(400,"PNL report file is too large.")
    try: parsed=_parse_pnl_upload(data,filename)
    except Exception as e: raise HTTPException(400,f"Could not parse Opportunity PNL Report: {e}")
    positions=[]
    for i,x in enumerate(parsed["positions"],1):
        positions.append({**x,"id":i,"yahoo_manual_price":None,"bloomberg_manual_price":None})
    stored={"filename":filename,"report_date":parsed["report_date"],"run_date":parsed["run_date"],"nav_usd":parsed["nav_usd"],"positions":positions,"version":int(datetime.now(timezone.utc).timestamp())}
    with SessionLocal() as db:
        changes=_opp_pnl_changes(db,parsed,filename);_set_json_setting(db,"opportunity_daily_baseline",stored);audit(db,u["username"],"OPPORTUNITY_PNL_CHANGES",f'added={len(changes.get("added") or [])}; archived={len(changes.get("archived") or [])}; expired={len(changes.get("expired") or [])}; qty_changed={len(changes.get("quantity_changed") or [])}; reactivated={len(changes.get("reactivated") or [])}');_bump_opportunity_config(db,u["username"],f'PNL baseline {filename}; report={parsed.get("report_date")}; positions={len(positions)}');db.commit()
    return RedirectResponse("/admin/opportunity",303)

@app.post("/admin/opportunity/settings")
async def opportunity_settings(request:Request,limit_pct:str=Form("")):
    u=require_admin(request);raw=str(limit_pct or "").strip()
    if raw:
        try:
            val=float(raw)
            if val<=0: raise ValueError
        except Exception: raise HTTPException(400,"Opportunity limit must be positive or blank.")
    with SessionLocal() as db:
        set_setting(db,"opportunity_limit_pct",raw);_bump_opportunity_config(db,u["username"],f"Limit={raw or 'N/A'}");db.commit()
    return RedirectResponse("/admin/opportunity",303)

@app.post("/admin/opportunity/manual-price")
async def opportunity_manual_price(request:Request):
    u=require_admin(request);form=await request.form();source=str(form.get("source") or "").upper();clear=str(form.get("clear") or "")=="1"
    try: pid=int(form.get("position_id"))
    except Exception: raise HTTPException(400,"Invalid position id.")
    val=None
    if not clear:
        try:
            val=float(str(form.get("price") or "").strip())
            if val<0: raise ValueError
        except Exception: raise HTTPException(400,"Manual price must be zero or greater.")
    with SessionLocal() as db:
        b=_opportunity_baseline(db)
        if not b: raise HTTPException(404,"Opportunity baseline not found.")
        found=False
        for p in b.get("positions",[]):
            if int(p.get("id") or 0)==pid:
                if source=="YAHOO": p["yahoo_manual_price"]=val;redirect="/opportunity/yahoo-daily-return"
                elif source=="BLOOMBERG": p["bloomberg_manual_price"]=val;redirect="/opportunity/bloomberg-daily-return"
                else: raise HTTPException(400,"Unknown source.")
                found=True;break
        if not found: raise HTTPException(404,"Position not found.")
        _set_json_setting(db,"opportunity_daily_baseline",b);audit(db,u["username"],"OPPORTUNITY_MANUAL_PRICE",f"{source} position_id={pid}; price={val}");db.commit()
    return RedirectResponse(redirect,303)

@app.get("/admin")
def admin_home(request:Request):
    require_admin(request)
    return RedirectResponse("/admin/fortress",303)

@app.get("/admin/fortress")
def fortress_setup(request:Request):
    u=require_admin(request)
    with SessionLocal() as db:
        nav=float(get_setting(db,"nav_usd","0") or 0)
        ver=get_setting(db,"config_version","1")
        active_count=db.query(Position).filter_by(active=True).count()
        b=_active_daily_baseline(db)
        current='<p class="muted">No P&amp;L baseline loaded.</p>'
        if b:
            pos=db.query(DailyReturnPosition).filter_by(baseline_id=b.id).order_by(DailyReturnPosition.id).all()
            prows=''
            for p in pos:
                strike='' if p.strike is None else f"{p.strike:g}"
                name=f"{p.ticker} {p.expiry} {p.option_type}{strike}" if p.instrument_type=="OPTION" else p.ticker
                prows+=f'<tr><td>{html.escape(name)}</td><td>{p.quantity:,.0f}</td><td>${p.baseline_price:.2f}</td></tr>'
            current=f'''<div class="notice"><strong>Active baseline:</strong> {html.escape(b.filename)} · Report {html.escape(b.report_date)} · NAV ${b.nav_usd:,.2f} · Positions {len(pos)}</div><br><table><thead><tr><th>Position</th><th>Qty</th><th>Previous Mark</th></tr></thead><tbody>{prows}</tbody></table>'''

    body=admin_tabs()+f'''<h1>Fortress Setup</h1>
<div class="card"><h3>Fortress Delta Monitor — Exposure Upload</h3>
<p>Upload the latest Nirvana <strong>Exposure Summary by Underlying</strong> XLSX. The portal reads option positions, quantities and NAV, then shows a change preview before anything is published.</p>
<form method="post" action="/admin/import" enctype="multipart/form-data"><input type="file" name="file" accept=".xlsx,.csv" required><button>Upload Fortress Exposure</button></form><br>
<div class="notice"><strong>Current Delta config:</strong> v{html.escape(str(ver))} · NAV ${nav:,.2f} · Active Delta positions {active_count}</div></div>
<div class="card"><h3>Fortress Daily Return Baseline — P&amp;L Upload</h3>
<p>Upload the previous trading day's Nirvana <strong>PNL Report</strong> as PDF or Excel (XLSX). It supplies positions, quantities, previous marks and NAV for both Bloomberg and Yahoo Daily Return.</p>
<form method="post" action="/admin/daily-return/upload" enctype="multipart/form-data"><input type="file" name="file" accept=".pdf,.xlsx,.xlsm,application/pdf,application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" required><button>Upload Fortress P&amp;L</button></form><br>{current}</div>'''
    with SessionLocal() as db:
        today=_ny_now().date().isoformat()
        portfolio=db.query(Position).order_by(Position.ticker,Position.expiry,Position.strike).all()
        prows=[]
        for item in portfolio:
            status=("EXPIRED" if str(item.expiry or "")<today else ("ACTIVE" if item.active else "ARCHIVED"))
            sec=html.escape(str(item.bloomberg_security or ""))
            prows.append(f'<tr><td>{html.escape(str(item.ticker))}</td><td>{html.escape(str(item.expiry))}</td><td>{html.escape(str(item.option_type))}</td><td>{float(item.strike):g}</td><td>{float(item.quantity):,.0f}</td><td>{status}</td><td>{sec}</td></tr>')
        body+='''<div class="card"><h3>Fortress Portfolio</h3><p class="muted">Read-only Exposure portfolio. Positions come from the current Nirvana upload and database; EXPIRED / ARCHIVED are displayed for reference only. Intraday trades are managed below, not here.</p><div style="overflow-x:auto"><table><thead><tr><th>Ticker</th><th>Expiry</th><th>Type</th><th>Strike</th><th>Qty</th><th>Status</th><th>Bloomberg Security</th></tr></thead><tbody>{''.join(prows) or '<tr><td colspan="7">No portfolio positions.</td></tr>'}</tbody></table></div></div>'''
        body+=_intraday_panel(db,"fortress")
    return page("Fortress Setup",body,u)

@app.post("/admin/settings")
def settings_update(request:Request,nav_usd:str=Form(...),limit_pct:str=Form(...)):
    u=require_admin(request);nav=float(nav_usd.replace(",",""));limit=float(limit_pct)
    if nav<=0 or not 0<limit<100:raise HTTPException(400,"Invalid NAV/limit")
    with SessionLocal() as db:
        set_setting(db,"nav_usd",nav);set_setting(db,"limit_pct",limit);config_bump(db,u["username"],f"NAV={nav}; limit={limit}%");db.commit()
    return RedirectResponse("/admin",303)

@app.post("/admin/positions/add")
def pos_add(request:Request,ticker:str=Form(...),expiry:str=Form(...),option_type:str=Form(...),strike:str=Form(...),quantity:str=Form(...),multiplier:str=Form("100")):
    u=require_admin(request);t=ticker.strip().upper();typ=option_type.strip().upper()[0];k=float(strike);q=float(quantity);m=float(multiplier);datetime.strptime(expiry,"%Y-%m-%d")
    if typ not in ("P","C"):raise HTTPException(400)
    with SessionLocal() as db:
        db.add(Position(ticker=t,expiry=expiry,option_type=typ,strike=k,quantity=q,multiplier=m,bloomberg_security=option_security(t,expiry,typ,k),underlying_security=f"{t} US Equity",active=True,source="manual"))
        audit(db,u["username"],"POSITION_ADDED",f"{t} {expiry} {typ}{k} qty={q}");config_bump(db,u["username"],"Position added");db.commit()
    return RedirectResponse("/admin",303)

@app.post("/admin/positions/{pid}/quantity")
def pos_qty(request:Request,pid:int,quantity:str=Form(...)):
    u=require_admin(request);q=float(quantity)
    with SessionLocal() as db:
        p=db.get(Position,pid)
        if not p:raise HTTPException(404)
        old=p.quantity;p.quantity=q;p.updated_at=utcnow();audit(db,u["username"],"QUANTITY_CHANGED",f"{p.ticker} {p.expiry} {p.option_type}{p.strike}: {old}->{q}");config_bump(db,u["username"],"Quantity changed");db.commit()
    return RedirectResponse("/admin",303)

@app.post("/admin/positions/{pid}/archive")
def pos_archive(request:Request,pid:int):
    u=require_admin(request)
    with SessionLocal() as db:
        p=db.get(Position,pid)
        if not p:raise HTTPException(404)
        if p.active:
            p.active=False;p.updated_at=utcnow();audit(db,u["username"],"POSITION_ARCHIVED",f"{p.ticker} {p.expiry} {p.option_type}{p.strike}");config_bump(db,u["username"],"Position archived");db.commit()
    return RedirectResponse("/admin",303)

@app.get("/admin/users")
def users_page(request:Request):
    u=require_admin(request)
    with SessionLocal() as db:
        users=db.query(User).order_by(User.username).all();rows=""
        for x in users:
            if x.username==ADMIN_USER:
                rows+=f"<tr><td>{x.username}</td><td>ADMIN (fixed)</td><td>Active</td><td>Protected</td></tr>"
            else:
                rows+=f"""<tr><td>{html.escape(x.username)}</td><td>Viewer</td><td>{'Active' if x.active else 'Disabled'}</td><td>
<form method="post" action="/admin/users/{x.id}/toggle" style="display:inline"><button class="secondary">{'Disable' if x.active else 'Enable'}</button></form>
<form method="post" action="/admin/users/{x.id}/reset" style="display:inline"><input type="password" name="password" placeholder="New password" required><button>Reset</button></form>
<form method="post" action="/admin/users/{x.id}/delete" style="display:inline"><button class="danger">Delete</button></form></td></tr>"""
        body=admin_tabs()+f"""<div class="card"><h2>Viewer Users</h2><div class="notice">Fuat is the only admin. New accounts are viewer-only.</div><br>
<form method="post" action="/admin/users/add" class="row"><label>Username<input name="username" required></label><label>Temporary password<input type="password" name="password" required></label><button>Add Viewer</button></form></div>
<div class="card"><table><thead><tr><th>User</th><th>Role</th><th>Status</th><th>Actions</th></tr></thead><tbody>{rows}</tbody></table></div>"""
        return page("Users",body,u)

@app.post("/admin/users/add")
def user_add(request:Request,username:str=Form(...),password:str=Form(...)):
    u=require_admin(request);name=username.strip().lower()
    if name==ADMIN_USER or not re.match(r"^[a-z0-9._-]{3,40}$",name):raise HTTPException(400,"Invalid/reserved username")
    if len(password)<8:raise HTTPException(400,"Password must be at least 8 characters")
    with SessionLocal() as db:
        if db.query(User).filter_by(username=name).first():raise HTTPException(400,"User exists")
        db.add(User(username=name,password_hash=hash_password(password),role="viewer",active=True));audit(db,u["username"],"VIEWER_USER_ADDED",name);db.commit()
    return RedirectResponse("/admin/users",303)

def viewer(db,uid):
    x=db.get(User,uid)
    if not x or x.username==ADMIN_USER:raise HTTPException(400,"Protected user")
    return x

@app.post("/admin/users/{uid}/toggle")
def user_toggle(request:Request,uid:int):
    u=require_admin(request)
    with SessionLocal() as db:
        x=viewer(db,uid);x.active=not x.active;x.updated_at=utcnow();audit(db,u["username"],"VIEWER_STATUS_CHANGED",f"{x.username}: active={x.active}");db.commit()
    return RedirectResponse("/admin/users",303)

@app.post("/admin/users/{uid}/reset")
def user_reset(request:Request,uid:int,password:str=Form(...)):
    u=require_admin(request)
    if len(password)<8:raise HTTPException(400)
    with SessionLocal() as db:
        x=viewer(db,uid);x.password_hash=hash_password(password);x.updated_at=utcnow();audit(db,u["username"],"VIEWER_PASSWORD_RESET",x.username);db.commit()
    return RedirectResponse("/admin/users",303)

@app.post("/admin/users/{uid}/delete")
def user_delete(request:Request,uid:int):
    u=require_admin(request)
    with SessionLocal() as db:
        x=viewer(db,uid);name=x.username;db.delete(x);audit(db,u["username"],"VIEWER_USER_DELETED",name);db.commit()
    return RedirectResponse("/admin/users",303)

@app.get("/admin/audit")
def audit_page(request:Request):
    u=require_admin(request)
    with SessionLocal() as db:
        items=db.query(Audit).order_by(Audit.id.desc()).limit(300).all()
        rows="".join(f"<tr><td>{x.ts}</td><td>{html.escape(x.actor)}</td><td>{html.escape(x.action)}</td><td>{html.escape(x.detail)}</td></tr>" for x in items)
    return page("Audit",admin_tabs()+f'<div class="card"><h2>Audit Log</h2><table><thead><tr><th>Time</th><th>Actor</th><th>Action</th><th>Detail</th></tr></thead><tbody>{rows}</tbody></table></div>',u)

@app.get("/admin/import")
def import_get(request:Request):
    require_admin(request)
    return RedirectResponse("/admin/fortress",303)

def _compare_import(parsed,current):
    cur={keypos(x):x for x in current};new={keypos(x):x for x in parsed}
    added=[x for k,x in new.items() if k not in cur]
    changed=[(cur[k],x) for k,x in new.items() if k in cur and (float(cur[k]["quantity"])!=float(x["quantity"]) or float(cur[k]["multiplier"])!=float(x["multiplier"]))]
    missing=[x for k,x in cur.items() if k not in new]
    return added,changed,missing

def _preview_html(jid,parsed,errors,current,mapping,report_nav=None,current_nav=None,report_date=None,auto=False):
    added,changed,missing=_compare_import(parsed,current)
    def fmt(x):return f'{x["ticker"]} {x["expiry"]} {x["option_type"]}{x["strike"]:g} qty={x["quantity"]:,.0f}'
    detail="<h3>Added</h3><ul>"+"".join(f"<li>{html.escape(fmt(x))}</li>" for x in added[:50])+"</ul>"
    detail+="<h3>Changed</h3><ul>"+"".join(f'<li>{html.escape(fmt(a))} → qty={b["quantity"]:,.0f}</li>' for a,b in changed[:50])+"</ul>"
    detail+="<h3>Missing from uploaded file</h3><ul>"+"".join(f"<li>{html.escape(fmt(x))}</li>" for x in missing[:50])+"</ul>"
    if errors:detail+="<h3>Parse errors</h3><pre>"+html.escape("\n".join(errors[:30]))+"</pre>"
    navline=""
    if report_nav:
        navline=f'<div class="notice"><strong>NAV from Nirvana:</strong> ${report_nav:,.2f}'
        if current_nav is not None:navline+=f' &nbsp; | &nbsp; Current portal NAV: ${current_nav:,.2f}'
        navline+=' &nbsp; — NAV will be updated when you Confirm & Publish.</div><br>'
    source='Nirvana Exposure Summary auto-detected' if auto else 'Generic mapped import'
    if report_date:source+=f' · Report Date {html.escape(str(report_date))}'
    mj=quote(json.dumps(mapping))
    return admin_tabs()+f"""<div class="card"><h2>Import Preview</h2><p><strong>{source}</strong></p>{navline}
<p><strong>Parsed option positions:</strong> {len(parsed)} · <strong>Added:</strong> {len(added)} · <strong>Changed:</strong> {len(changed)} · <strong>Missing:</strong> {len(missing)} · <strong>Parse errors:</strong> {len(errors)}</p>{detail}
<form method="post" action="/admin/import/{jid}/apply"><input type="hidden" name="mapping_json" value="{html.escape(mj)}"><label>Mode<select name="mode"><option value="replace">Replace active portfolio — missing positions archived</option><option value="merge">Merge/update only — missing positions untouched</option></select></label><br><br><button>Confirm & Publish</button> <a href="/admin/fortress">Cancel</a></form></div>"""

@app.post("/admin/import")
async def import_upload(request:Request,file:UploadFile=File(...)):
    u=require_admin(request)
    data=await file.read()
    headers,rows,meta=parse_upload(data,file.filename or "upload")
    if not rows:raise HTTPException(400,"No rows found")
    jid=uuid.uuid4().hex
    envelope={"headers":headers,"meta":meta}
    with SessionLocal() as db:
        db.add(ImportJob(id=jid,filename=file.filename or "upload",headers_json=json.dumps(envelope),rows_json=json.dumps(rows,default=str)))
        audit(db,u["username"],"NIRVANA_FILE_UPLOADED",f"{file.filename}; rows={len(rows)}; format={meta.get('format')}")
        db.commit()

    if meta.get("format")=="nirvana_exposure_by_underlying":
        parsed,errors,_=parse_nirvana_exposure_rows(rows)
        if not parsed:raise HTTPException(400,"Nirvana report detected but no option rows could be parsed.")
        with SessionLocal() as db:
            current=[{"ticker":p.ticker,"expiry":p.expiry,"option_type":p.option_type,"strike":p.strike,"quantity":p.quantity,"multiplier":p.multiplier} for p in db.query(Position).filter_by(active=True).all()]
            current_nav=float(get_setting(db,"nav_usd","0") or 0)
        mapping={"format":"nirvana_exposure_by_underlying"}
        body=_preview_html(jid,parsed,errors,current,mapping,meta.get("nav_usd"),current_nav,meta.get("report_date"),True)
        return page("Nirvana Import Preview",body,u)

    mapping=detect_map(headers)
    def opts(sel=""):
        return "".join(f'<option value="{html.escape(h)}" {"selected" if h==sel else ""}>{html.escape(h)}</option>' for h in headers)
    body=admin_tabs()+f"""<div class="card"><h2>Map Columns</h2><p>{html.escape(file.filename or "")}: {len(rows)} rows</p><form method="post" action="/admin/import/{jid}/preview"><div class="row">
<label>Ticker<select name="ticker_col">{opts(mapping.get("ticker",""))}</select></label><label>Expiry<select name="expiry_col">{opts(mapping.get("expiry",""))}</select></label>
<label>Put/Call<select name="type_col">{opts(mapping.get("option_type",""))}</select></label><label>Strike<select name="strike_col">{opts(mapping.get("strike",""))}</select></label>
<label>Quantity<select name="qty_col">{opts(mapping.get("quantity",""))}</select></label><label>Multiplier<select name="mult_col"><option value="">Default 100</option>{opts(mapping.get("multiplier",""))}</select></label></div><br><button>Build Preview</button></form></div>"""
    return page("Map Import",body,u)

@app.post("/admin/import/{jid}/preview")
def import_preview(request:Request,jid:str,ticker_col:str=Form(...),expiry_col:str=Form(...),type_col:str=Form(...),strike_col:str=Form(...),qty_col:str=Form(...),mult_col:str=Form("")):
    u=require_admin(request);mapping={"ticker":ticker_col,"expiry":expiry_col,"type":type_col,"strike":strike_col,"qty":qty_col,"mult":mult_col}
    with SessionLocal() as db:
        job=db.get(ImportJob,jid)
        if not job:raise HTTPException(404)
        parsed,errors=normalized_import(json.loads(job.rows_json),mapping)
        current=[{"ticker":p.ticker,"expiry":p.expiry,"option_type":p.option_type,"strike":p.strike,"quantity":p.quantity,"multiplier":p.multiplier} for p in db.query(Position).filter_by(active=True).all()]
    body=_preview_html(jid,parsed,errors,current,mapping)
    return page("Import Preview",body,u)

@app.post("/admin/import/{jid}/apply")
def import_apply(request:Request,jid:str,mapping_json:str=Form(...),mode:str=Form(...)):
    u=require_admin(request);mapping=json.loads(unquote(mapping_json))
    with SessionLocal() as db:
        job=db.get(ImportJob,jid)
        if not job:raise HTTPException(404)
        rows=json.loads(job.rows_json)
        try:envelope=json.loads(job.headers_json)
        except Exception:envelope={"headers":[],"meta":{}}
        meta=envelope.get("meta",{}) if isinstance(envelope,dict) else {}
        if mapping.get("format")=="nirvana_exposure_by_underlying":
            parsed,errors,_=parse_nirvana_exposure_rows(rows)
        else:
            parsed,errors=normalized_import(rows,mapping)
        if not parsed:raise HTTPException(400,"No valid option positions parsed; nothing was changed.")
        if errors and len(errors)>max(3,len(parsed)//10):raise HTTPException(400,"Too many parse errors; nothing was changed.")
        active=db.query(Position).filter_by(active=True).all();cur={(p.ticker,p.expiry,p.option_type,round(p.strike,6)):p for p in active};incoming={keypos(x):x for x in parsed}
        added=changed=archived=0
        for k,x in incoming.items():
            p=cur.get(k)
            if p:
                if p.quantity!=x["quantity"] or p.multiplier!=x["multiplier"]:
                    p.quantity=x["quantity"];p.multiplier=x["multiplier"];p.updated_at=utcnow();changed+=1
            else:
                db.add(Position(ticker=x["ticker"],expiry=x["expiry"],option_type=x["option_type"],strike=x["strike"],quantity=x["quantity"],multiplier=x["multiplier"],bloomberg_security=option_security(x["ticker"],x["expiry"],x["option_type"],x["strike"]),underlying_security=f'{x["ticker"]} US Equity',active=True,source="nirvana"));added+=1
        if mode=="replace":
            for k,p in cur.items():
                if k not in incoming:p.active=False;p.updated_at=utcnow();archived+=1
        nav_note=""
        report_nav=meta.get("nav_usd") if mapping.get("format")=="nirvana_exposure_by_underlying" else None
        if report_nav:
            old_nav=float(get_setting(db,"nav_usd","0") or 0)
            set_setting(db,"nav_usd",float(report_nav))
            nav_note=f"; NAV {old_nav}->{float(report_nav)}"
        audit(db,u["username"],"NIRVANA_IMPORT_APPLIED",f"{job.filename}; mode={mode}; added={added}; changed={changed}; archived={archived}; parse_errors={len(errors)}{nav_note}")
        report_day=str(meta.get("report_date") or "")
        if re.fullmatch(r"\d{1,2}/\d{1,2}/\d{4}",report_day):
            report_day=datetime.strptime(report_day,"%m/%d/%Y").date().isoformat()
        if not re.fullmatch(r"\d{4}-\d{2}-\d{2}",report_day):
            current_baseline=_active_daily_baseline(db)
            report_day=str(current_baseline.report_date or "") if current_baseline else ""
        if re.fullmatch(r"\d{4}-\d{2}-\d{2}",report_day):
            set_setting(db,"fortress_exposure_report_date",report_day)
        config_bump(db,u["username"],f"Nirvana import {job.filename}")
        db.delete(job);db.commit()
    return RedirectResponse("/admin/fortress",303)

def _active_daily_baseline(db):
    return db.query(DailyReturnBaseline).filter_by(active=True).order_by(DailyReturnBaseline.id.desc()).first()


def _latest_yahoo_collector_snapshot(db):
    """Return Fortress Yahoo snapshot; prefer the multi-fund local collector."""
    v2=_latest_yahoo_v2_fund_snapshot(db,"fortress")
    if v2:
        return v2
    raw=get_setting(db,"yahoo_collector_snapshot","")
    if not raw:
        return None
    try:
        data=json.loads(raw)
        return data if isinstance(data,dict) else None
    except Exception:
        return None

def _yahoo_quote_data_from_snapshot(snapshot):
    """Convert local collector rows into calculate_mark_to_market_return quote_data."""
    out={}
    if not snapshot:
        return out
    for r in snapshot.get("positions",[]) or []:
        symbol=str(r.get("symbol") or "").upper().strip()
        if not symbol:
            continue
        out[symbol]={
            "mark":r.get("mark"),
            "bid":r.get("bid"),
            "ask":r.get("ask"),
            "last":r.get("last"),
            "source":r.get("source") or "",
            "note":r.get("note") or "",
        }
    return out

@app.get("/daily-return")
def daily_return_page(request:Request):
    u=require_user(request)
    with SessionLocal() as db:
        b=_active_daily_baseline(db)
        if not b:
            extra=' <a href="/admin/daily-return">Upload P&amp;L Report</a>' if u["username"]==ADMIN_USER else ''
            return page("Estimated Daily Return",f'<h1>Estimated Daily Return</h1><div class="card"><div class="notice">No P&amp;L baseline has been uploaded yet.{extra}</div></div>',u)

        dbpos=db.query(DailyReturnPosition).filter_by(baseline_id=b.id).order_by(DailyReturnPosition.id).all()
        positions=[{
            "id":p.id,
            "security_name":p.security_name,
            "instrument_type":p.instrument_type,
            "ticker":p.ticker,
            "expiry":p.expiry,
            "option_type":p.option_type,
            "strike":p.strike,
            "quantity":p.quantity,
            "multiplier":p.multiplier,
            "baseline_price":p.baseline_price,
            "baseline_market_value":p.baseline_market_value,
            "yahoo_manual_price":p.yahoo_manual_price,
            "bloomberg_manual_price":p.bloomberg_manual_price,
        } for p in dbpos]
        baseline={
            "report_date":b.report_date,
            "run_date":b.run_date,
            "filename":b.filename,
            "nav_usd":b.nav_usd,
        }

    with SessionLocal() as db:
        yahoo_snapshot=_latest_yahoo_collector_snapshot(db)
    yahoo_quote_data=_yahoo_quote_data_from_snapshot(yahoo_snapshot)
    with SessionLocal() as db:
        calc=_daily_with_ledger(db,"fortress","yahoo",positions,float(baseline["nav_usd"]),yahoo_snapshot,baseline["report_date"])
    calc["yahoo_error"]=""
    status="FULL" if calc["complete"] else "PARTIAL"
    ret=calc["estimated_return_pct"] or 0.0
    pnl=calc["estimated_pnl"] or 0.0

    rows=""
    for r in calc["rows"]:
        strike="" if r.get("strike") is None else f'{r["strike"]:g}'
        posname=(f'{r["ticker"]} {r["expiry"]} {r["option_type"]}{strike}' if r["instrument_type"]=="OPTION" else r["ticker"])
        previous='—' if r.get("previous_mark") is None else f'${r["previous_mark"]:.2f}'
        bid='—' if r.get("bid") is None else f'${r["bid"]:.2f}'
        ask='—' if r.get("ask") is None else f'${r["ask"]:.2f}'
        market='—' if r.get("market_mark") is None else f'${r["market_mark"]:.2f}'
        current='—' if r.get("current_mark") is None else f'${r["current_mark"]:.2f}'
        manual_val='' if r.get("manual_price") is None else f'{float(r["manual_price"]):.4f}'
        change='—' if r.get("change") is None else f'{r["change"]:+.2f}'
        epnl='—' if r.get("estimated_pnl") is None else f'${r["estimated_pnl"]:+,.2f}'
        contrib='—' if r.get("contribution_pct") is None else f'{r["contribution_pct"]:+.4f}%'
        source=(r.get("effective_source") or "")
        if u["username"]==ADMIN_USER and r.get("id") is not None:
            manual_html=(f'<form method="post" action="/admin/daily-return/manual-price" style="display:flex;gap:4px;align-items:center">'
                         f'<input type="hidden" name="position_id" value="{r["id"]}">'
                         f'<input type="hidden" name="source" value="YAHOO">'
                         f'<input type="number" step="0.0001" min="0" name="price" value="{manual_val}" style="width:82px">'
                         f'<button>Save</button><button name="clear" value="1">Clear</button></form>')
        else:
            manual_html='—' if not manual_val else f'${float(manual_val):.4f}'
        rows+=f'<tr><td>{html.escape(posname)}</td><td>{r["quantity"]:,.0f}</td><td>{previous}</td><td>{bid}</td><td>{ask}</td><td>{market}</td><td>{manual_html}</td><td>{current}</td><td>{change}</td><td>{epnl}</td><td>{contrib}</td><td>{html.escape(source)}</td></tr>'

    warning=""
    if not calc["complete"]:
        warning='<div class="notice"><strong>PARTIAL:</strong> Yahoo local collector has not supplied a usable current mark for every position. The return shown includes only covered positions.</div><br>'

    if yahoo_snapshot:
        yahoo_collector_name=str(yahoo_snapshot.get("collector_name") or "Yahoo-PC")
        yahoo_collected_at=str(yahoo_snapshot.get("collected_at_utc") or "")
        yahoo_collector_error=str(yahoo_snapshot.get("error") or "")
        try:
            ts=datetime.fromisoformat(yahoo_collected_at.replace("Z","+00:00"))
            age_seconds=max(0,(datetime.now(timezone.utc)-ts).total_seconds())
        except Exception:
            age_seconds=None
        if age_seconds is not None and age_seconds > 1800:
            warning+=f'<div class="notice"><strong>Yahoo Collector:</strong> Snapshot is stale ({age_seconds/60:.0f} minutes old) · {html.escape(yahoo_collector_name)}</div><br>'
        if yahoo_collector_error:
            warning+=f'<div class="notice"><strong>Yahoo Collector:</strong> {html.escape(yahoo_collector_error)}</div><br>'
    else:
        warning+='<div class="notice"><strong>Yahoo Collector:</strong> No local collector snapshot has been received yet.</div><br>'

    quality_class='ok' if calc['complete'] else 'bad'
    body=f'''<h1>RPD Fortress Fund — Yahoo Daily Return</h1>{_fund_export_buttons("fortress")}
<div class="grid">
<div class="card metric"><div class="muted">Status</div><div class="v {quality_class}">{status}</div><div class="muted">Mark-to-market estimate</div></div>
<div class="card metric"><div class="muted">Estimated Daily Return</div><div class="v">{ret:+.4f}%</div></div>
<div class="card metric"><div class="muted">Estimated P&amp;L</div><div class="v">${pnl:+,.2f}</div></div>
<div class="card metric"><div class="muted">Baseline NAV</div><div class="v">${baseline['nav_usd']:,.2f}</div><div class="muted">Report {html.escape(baseline['report_date'])}</div></div>
<div class="card metric"><div class="muted">Coverage</div><div class="v">{calc['coverage_valid']}/{calc['coverage_total']}</div><div class="muted">Yahoo local collector marks</div></div>
</div>
{warning}
<div class="card">
<div class="muted">Daily Return baseline — completely separate from Bloomberg</div>
<strong>{html.escape(baseline['filename'])}</strong> · Report Date {html.escape(baseline['report_date'])} · NAV ${baseline['nav_usd']:,.2f}
<br><div class="muted" style="margin-top:6px">Calculated {html.escape(ny_time_label(calc['calculated_at_utc']))}. Current Mark is supplied by the local Yahoo Collector: Bid/Ask midpoint when both are available; if Bid is zero/unavailable and Ask is available, Ask/2; otherwise Yahoo Last. Formula: (Current Mark − Previous P&amp;L Report Price) × Quantity × Multiplier. v6.15 intraday adjustments: Signed Trade Qty × (Current Mark − Execution Price) × Multiplier; confirmed worthless and automatically estimated OTM expirations use a zero mark. Daily Return = Total Estimated P&amp;L ÷ Baseline NAV.</div>
</div>
<div class="card"><table><thead><tr><th>Position</th><th>Qty</th><th>Previous Mark</th><th>Yahoo Bid</th><th>Yahoo Ask</th><th>Yahoo Market</th><th>Manual Override</th><th>Effective Price</th><th>Change</th><th>Estimated P&amp;L</th><th>Contribution</th><th>Effective Source</th></tr></thead><tbody>{rows}</tbody></table></div>'''
    return page("Yahoo Daily Return",body,u,60 if auto_refresh_allowed() else None)


@app.get("/bloomberg-daily-return")
def bloomberg_daily_return_page(request:Request):
    u=require_user(request)

    with SessionLocal() as db:
        b=_active_daily_baseline(db)
        if not b:
            extra=' <a href="/admin/daily-return">Upload P&amp;L Report</a>' if u["username"]==ADMIN_USER else ''
            return page(
                "Bloomberg Daily Return",
                f'<h1>RPD Fortress Fund — Bloomberg Daily Return</h1>'
                f'<div class="card"><div class="notice">No P&amp;L baseline has been uploaded yet.{extra}</div></div>',
                u,60
            )

        dbpos=db.query(DailyReturnPosition).filter_by(baseline_id=b.id).order_by(DailyReturnPosition.id).all()
        baseline={
            "report_date":b.report_date,
            "run_date":b.run_date,
            "filename":b.filename,
            "nav_usd":float(b.nav_usd),
        }
        snap,age=latest_snapshot(db)

    positions=[{
        "security_name":p.security_name,
        "instrument_type":p.instrument_type,
        "ticker":p.ticker,
        "expiry":p.expiry,
        "option_type":p.option_type,
        "strike":p.strike,
        "quantity":float(p.quantity),
        "multiplier":float(p.multiplier or 100),
        "baseline_price":float(p.baseline_price) if p.baseline_price is not None else None,
        "position_id":p.id,
        "bloomberg_manual_price":float(p.bloomberg_manual_price) if p.bloomberg_manual_price is not None else None,
    } for p in dbpos]

    with SessionLocal() as db:
        calc=_daily_with_ledger(db,"fortress","bloomberg",positions,baseline["nav_usd"],snap,baseline["report_date"])
    snapshot_ts=(snap or {}).get("timestamp_utc","")
    collector_id=(snap or {}).get("collector_id","")
    collector_mode=(snap or {}).get("mode","")
    rows=""
    total_pnl=float(calc["estimated_pnl"])
    valid=calc["coverage_valid"]
    total=calc["coverage_total"]
    ret=calc["estimated_return_pct"]
    complete=calc["complete"]
    for r in calc["rows"]:
        posname=_opp_position_label(r)
        fmt=lambda value: "—" if value is None else f'${float(value):.2f}'
        prev_txt=fmt(r.get("baseline_price"));bid_txt=fmt(r.get("bid"));ask_txt=fmt(r.get("ask"))
        market_txt=fmt(r.get("market_mark"));cur_txt=fmt(r.get("current_mark"))
        chg_txt="—" if r.get("change") is None else f'{float(r["change"]):+.2f}'
        pnl_txt="—" if r.get("estimated_pnl") is None else f'${float(r["estimated_pnl"]):+,.2f}'
        con_txt="—" if r.get("contribution_pct") is None else f'{float(r["contribution_pct"]):+.4f}%'
        manual=r.get("bloomberg_manual_price")
        manual_val="" if manual is None else f'{float(manual):.4f}'
        pid=r.get("position_id")
        if u["username"]==ADMIN_USER and pid is not None:
            manual_html=(f'<form method="post" action="/admin/daily-return/manual-price" style="display:flex;gap:4px">'
                f'<input type="hidden" name="position_id" value="{pid}">'
                f'<input type="hidden" name="source" value="BLOOMBERG">'
                f'<input type="number" step="0.0001" min="0" name="price" value="{manual_val}" style="width:82px">'
                f'<button>Save</button><button name="clear" value="1">Clear</button></form>')
        else:manual_html="—" if not manual_val else f'${float(manual):.4f}'
        rows+=(f'<tr><td>{html.escape(posname)}</td><td>{float(r.get("quantity") or 0):,.0f}</td>'
               f'<td>{prev_txt}</td><td>{bid_txt}</td><td>{ask_txt}</td><td>{market_txt}</td>'
               f'<td>{manual_html}</td><td>{cur_txt}</td><td>{chg_txt}</td><td>{pnl_txt}</td>'
               f'<td>{con_txt}</td><td>{html.escape(str(r.get("effective_source") or ""))}</td></tr>')

    if not snap:
        status="WAITING"; status_class="bad"; feed_note="No Bloomberg Collector snapshot yet."
    elif valid==0:
        status="WAITING"; status_class="bad"; feed_note="The current snapshot does not yet contain Bloomberg option marks. The updated Collector must publish a new full snapshot."
    elif complete:
        status="FULL"; status_class="ok"; feed_note=f'Bloomberg snapshot age {age:.1f}s · {html.escape(collector_mode or "LIVE")}'
    else:
        status="PARTIAL"; status_class="bad"; feed_note=f'Bloomberg marks {valid}/{total} · snapshot age {age:.1f}s'

    warning=""
    if not complete:
        warning='<div class="notice"><strong>'+status+':</strong> '+feed_note+' The displayed P&amp;L includes covered positions only.</div><br>'

    body=f'''<h1>RPD Fortress Fund — Bloomberg Daily Return</h1>{_fund_export_buttons("fortress")}
<div class="grid">
<div class="card metric"><div class="muted">Status</div><div class="v {status_class}">{status}</div><div class="muted">Bloomberg marks</div></div>
<div class="card metric"><div class="muted">Estimated Daily Return</div><div class="v">{ret:+.4f}%</div></div>
<div class="card metric"><div class="muted">Estimated P&amp;L</div><div class="v">${total_pnl:+,.2f}</div></div>
<div class="card metric"><div class="muted">Baseline NAV</div><div class="v">${baseline['nav_usd']:,.2f}</div><div class="muted">Report {html.escape(baseline['report_date'])}</div></div>
<div class="card metric"><div class="muted">Coverage</div><div class="v">{valid}/{total}</div><div class="muted">Bloomberg option marks</div></div>
</div>
{warning}
<div class="card">
<div class="muted">Shared P&amp;L baseline</div>
<strong>{html.escape(baseline['filename'])}</strong> · Report Date {html.escape(baseline['report_date'])} · NAV ${baseline['nav_usd']:,.2f}
<br><div class="muted" style="margin-top:6px">
Bloomberg snapshot: {html.escape(format_new_york_time(snapshot_ts)) if snapshot_ts else '—'}
{(' · Collector: '+html.escape(collector_id)) if collector_id else ''}
<br>Current Mark = Bloomberg option Bid/Ask midpoint when both are available; otherwise Bloomberg Last.
Formula: (Current Mark − Previous P&amp;L Report Price) × Quantity × Multiplier. v6.15 intraday adjustments: Signed Trade Qty × (Current Mark − Execution Price) × Multiplier; confirmed worthless and automatically estimated OTM expirations use a zero mark.
Daily Return = Total Estimated P&amp;L ÷ Baseline NAV.
</div></div>
<div class="card"><table><thead><tr><th>Position</th><th>Qty</th><th>Previous Mark</th><th>Bloomberg Bid</th><th>Bloomberg Ask</th><th>Bloomberg Market</th><th>Manual Override</th><th>Effective Price</th><th>Change</th><th>Estimated P&amp;L</th><th>Contribution</th><th>Effective Source</th></tr></thead><tbody>{rows}</tbody></table></div>'''
    return page("Bloomberg Daily Return",body,u,60)

@app.get("/admin/daily-return")
def daily_return_admin(request:Request):
    require_admin(request)
    return RedirectResponse("/admin/fortress",303)


@app.post("/admin/daily-return/manual-price")
async def daily_return_manual_price(request:Request):
    u=require_admin(request)
    form=await request.form()
    try:
        position_id=int(form.get("position_id"))
    except Exception:
        raise HTTPException(400,"Invalid position id.")
    source=str(form.get("source") or "").upper()
    clear=str(form.get("clear") or "")=="1"
    raw=str(form.get("price") or "").strip()

    value=None
    if not clear:
        try:
            value=float(raw)
            if value < 0:
                raise ValueError
        except Exception:
            raise HTTPException(400,"Manual price must be zero or greater.")

    with SessionLocal() as db:
        p=db.query(DailyReturnPosition).filter_by(id=position_id).first()
        if not p:
            raise HTTPException(404,"Position not found.")
        if source=="YAHOO":
            p.yahoo_manual_price=value
            action="YAHOO_MANUAL_PRICE_CLEARED" if clear else "YAHOO_MANUAL_PRICE_SET"
            redirect="/daily-return"
        elif source=="BLOOMBERG":
            p.bloomberg_manual_price=value
            action="BLOOMBERG_MANUAL_PRICE_CLEARED" if clear else "BLOOMBERG_MANUAL_PRICE_SET"
            redirect="/bloomberg-daily-return"
        else:
            raise HTTPException(400,"Unknown source.")
        audit(db,u["username"],action,f"position_id={position_id}; price={value}")
        db.commit()

    return RedirectResponse(redirect,303)

@app.post("/admin/daily-return/upload")
async def daily_return_upload(request:Request,file:UploadFile=File(...)):
    u=require_admin(request)
    filename=file.filename or "PNL.pdf"
    if not filename.lower().endswith((".pdf",".xlsx",".xlsm")):
        raise HTTPException(400,"Please upload the Nirvana PNL Report as PDF or XLSX.")
    data=await file.read()
    if len(data)>12*1024*1024:
        raise HTTPException(400,"PNL report file is too large.")
    try:
        parsed=_parse_pnl_upload(data,filename)
    except Exception as e:
        raise HTTPException(400,f"Could not parse PNL Report: {e}")

    with SessionLocal() as db:
        for old in db.query(DailyReturnBaseline).filter_by(active=True).all():
            old.active=False
        b=DailyReturnBaseline(report_date=parsed["report_date"],run_date=parsed["run_date"],filename=filename,
                              nav_usd=parsed["nav_usd"],active=True,uploaded_by=u["username"])
        db.add(b);db.flush()
        for x in parsed["positions"]:
            db.add(DailyReturnPosition(
                baseline_id=b.id,security_name=x["security_name"],instrument_type=x["instrument_type"],ticker=x["ticker"],
                expiry=x.get("expiry","") or "",option_type=x.get("option_type","") or "",strike=x.get("strike"),
                quantity=x["quantity"],multiplier=x["multiplier"],baseline_price=x.get("baseline_price"),
                baseline_market_value=x["baseline_market_value"],manual_price=None,yahoo_manual_price=None,bloomberg_manual_price=None))
        audit(db,u["username"],"DAILY_RETURN_PNL_BASELINE_UPLOADED",
              f'{filename}; report_date={parsed["report_date"]}; nav={parsed["nav_usd"]}; positions={len(parsed["positions"])}')
        db.commit()

    return RedirectResponse("/admin/fortress",303)

def collector_auth(request):
    if not COLLECTOR_TOKEN or not hmac.compare_digest(request.headers.get("authorization",""),f"Bearer {COLLECTOR_TOKEN}"):
        raise HTTPException(401,"Bad collector token")


@app.get("/api/yahoo-collector/ping")
def yahoo_collector_ping(request:Request):
    collector_auth(request)
    return {"ok":True,"service":"yahoo-local-collector"}

@app.get("/api/yahoo-collector/config")
def yahoo_collector_config(request:Request):
    """Return the active Daily Return baseline positions to the local Yahoo collector."""
    collector_auth(request)
    with SessionLocal() as db:
        b=_active_daily_baseline(db)
        if not b:
            return {
                "version":0,
                "baseline_id":None,
                "report_date":None,
                "positions":[],
            }

        dbpos=db.query(DailyReturnPosition).filter_by(
            baseline_id=b.id
        ).order_by(DailyReturnPosition.id).all()

        positions=[]
        for p in dbpos:
            positions.append({
                "id":p.id,
                "instrument_type":str(p.instrument_type or "OPTION").upper(),
                "ticker":p.ticker,
                "expiry":p.expiry,
                "option_type":p.option_type,
                "strike":p.strike,
                "quantity":p.quantity,
                "multiplier":p.multiplier,
            })

        positions=_overlay_positions(positions,_active_trades(db,"fortress",str(b.report_date or "")))
        return {
            "version":f"{b.id}-C{get_setting(db,'config_version','1')}",
            "baseline_id":b.id,
            "report_date":b.report_date,
            "nav_usd":float(b.nav_usd or 0),
            "positions":positions,
        }

@app.post("/api/yahoo-collector/snapshot")
async def yahoo_collector_snapshot(request:Request):
    """Store the latest local Yahoo option-chain snapshot in Settings."""
    collector_auth(request)
    data=await request.json()

    if not isinstance(data,dict):
        raise HTTPException(400,"Invalid snapshot payload.")

    positions=data.get("positions")
    if not isinstance(positions,list):
        raise HTTPException(400,"positions must be a list.")

    # Keep payload intentionally small and derived only.
    cleaned=[]
    for r in positions[:200]:
        if not isinstance(r,dict):
            continue
        cleaned.append({
            "symbol":str(r.get("symbol") or "")[:64],
            "ticker":str(r.get("ticker") or "")[:30],
            "expiry":str(r.get("expiry") or "")[:10],
            "option_type":str(r.get("option_type") or "")[:4],
            "strike":r.get("strike"),
            "bid":r.get("bid"),
            "ask":r.get("ask"),
            "last":r.get("last"),
            "mark":r.get("mark"),
            "source":str(r.get("source") or "")[:80],
            "note":str(r.get("note") or "")[:160],
        })

    stored={
        "collector_name":str(data.get("collector_name") or "Yahoo-PC")[:80],
        "baseline_id":data.get("baseline_id"),
        "config_version":data.get("config_version"),
        "collected_at_utc":str(data.get("collected_at_utc") or utcnow().isoformat()),
        "coverage_valid":int(data.get("coverage_valid") or 0),
        "coverage_total":int(data.get("coverage_total") or len(cleaned)),
        "error":str(data.get("error") or "")[:1000],
        "positions":cleaned,
    }

    with SessionLocal() as db:
        set_setting(db,"yahoo_collector_snapshot",json.dumps(stored))
        set_setting(db,"yahoo_collector_updated_at",stored["collected_at_utc"])
        db.commit()

    return {"ok":True,"positions_received":len(cleaned)}


def _fortress_daily_positions(db):
    b=_active_daily_baseline(db)
    if not b:return None,[]
    dbpos=db.query(DailyReturnPosition).filter_by(baseline_id=b.id).order_by(DailyReturnPosition.id).all()
    return b,[{"id":p.id,"security_name":p.security_name,"instrument_type":p.instrument_type,"ticker":p.ticker,"expiry":p.expiry,"option_type":p.option_type,"strike":p.strike,"quantity":p.quantity,"multiplier":p.multiplier,"baseline_price":p.baseline_price,"baseline_market_value":p.baseline_market_value,"yahoo_manual_price":p.yahoo_manual_price,"bloomberg_manual_price":p.bloomberg_manual_price} for p in dbpos]

def _fortress_delta_positions(db):
    today=date.today().isoformat();out=[]
    for p in db.query(Position).filter_by(active=True).order_by(Position.ticker,Position.expiry,Position.strike).all():
        if p.expiry<today:continue
        out.append({"id":p.id,"instrument_type":"OPTION","ticker":p.ticker,"expiry":p.expiry,"option_type":p.option_type,"strike":p.strike,"quantity":p.quantity,"multiplier":p.multiplier,"bloomberg_security":p.bloomberg_security,"underlying_security":p.underlying_security})
    return out

@app.get("/api/collector/config-v2")
def collector_config_v2(request:Request):
    collector_auth(request)
    with SessionLocal() as db:
        fb,fdaily=_fortress_daily_positions(db);fdelta=_fortress_delta_positions(db);fcfg=int(get_setting(db,"config_version","1"));flimit=float(get_setting(db,"limit_pct","15"));fnav=float(get_setting(db,"nav_usd","0"))
        ob=_opportunity_baseline(db);oe=_opportunity_exposure(db);ocfg=int(get_setting(db,"opportunity_config_version","1") or 1);olimit=_opp_limit(db)
        opos_daily=list((ob or {}).get("positions") or []);opos_delta=list((oe or {}).get("positions") or [])
        # Delta and Daily baseline dates can differ; overlay each independently.
        fdelta_date=str(get_setting(db,"fortress_exposure_report_date","") or (fb.report_date if fb else ""))
        fdelta=_overlay_positions(fdelta,_active_trades(db,"fortress",fdelta_date))
        fdaily=_overlay_positions(fdaily,_active_trades(db,"fortress",str(fb.report_date or "") if fb else ""))
        odelta_date=str((oe or {}).get("report_date") or (ob or {}).get("report_date") or "")
        opos_delta=_overlay_positions(opos_delta,_active_trades(db,"opportunity",odelta_date))
        opos_daily=_overlay_positions(opos_daily,_active_trades(db,"opportunity",str((ob or {}).get("report_date") or "")))
        # Once expiration has been verified, avoid showing risk as an OPEN delta position.
        fexpired=_confirmed_expiry_keys(db,"fortress",fdelta_date)
        oexpired=_confirmed_expiry_keys(db,"opportunity",odelta_date)
        fdelta=[p for p in fdelta if _dr_position_key(p) not in fexpired]
        opos_delta=[p for p in opos_delta if _dr_position_key(p) not in oexpired]
        fmerged=_merge_market_positions(fdelta,fdaily);omerged=_merge_market_positions(opos_delta,opos_daily)
        version=f'F{fcfg}-FB{getattr(fb,"id",0) if fb else 0}-O{ocfg}-OB{(ob or {}).get("version",0)}'
        return {"version":version,"poll_seconds":120,"funds":{"fortress":{"fund_key":"fortress","portfolio_name":"RPD Fortress Fund","version":fcfg,"nav_usd":fnav,"limit_pct":flimit,"positions":fmerged},"opportunity":{"fund_key":"opportunity","portfolio_name":"RPD Opportunity Fund","version":ocfg,"nav_usd":float((oe or {}).get("nav_usd") or 0),"limit_pct":olimit,"positions":omerged}}}

@app.post("/api/collector/snapshot-v2")
async def collector_snapshot_v2(request:Request):
    collector_auth(request);data=await request.json();funds=(data or {}).get("funds") or {}
    with SessionLocal() as db:
        fp=funds.get("fortress")
        if isinstance(fp,dict):
            db.add(Snapshot(payload=json.dumps(fp)));db.flush();count=db.query(Snapshot).count()
            if count>250:
                for x in db.query(Snapshot).order_by(Snapshot.id.asc()).limit(count-250).all():db.delete(x)
        op=funds.get("opportunity")
        if isinstance(op,dict): set_setting(db,"opportunity_bloomberg_snapshot",json.dumps(op))
        db.commit()
    return {"ok":True,"funds_received":[k for k,v in funds.items() if isinstance(v,dict)]}

@app.get("/api/yahoo-collector/config-v2")
def yahoo_collector_config_v2(request:Request):
    collector_auth(request)
    with SessionLocal() as db:
        fb,fdaily=_fortress_daily_positions(db);ob=_opportunity_baseline(db);odaily=list((ob or {}).get("positions") or [])
        def clean(rows):
            out=[]
            for p in rows:
                inst=str(p.get("instrument_type") or "OPTION").upper()
                if inst not in ("OPTION","EQUITY"):continue
                out.append({"id":p.get("id"),"instrument_type":inst,"ticker":p.get("ticker"),"expiry":p.get("expiry") or "","option_type":p.get("option_type") or "","strike":p.get("strike"),"quantity":p.get("quantity"),"multiplier":p.get("multiplier",100 if inst=="OPTION" else 1)})
            return out
        fdaily=_overlay_positions(fdaily,_active_trades(db,"fortress",str(fb.report_date or "") if fb else ""))
        odaily=_overlay_positions(odaily,_active_trades(db,"opportunity",str((ob or {}).get("report_date") or "")))
        version=f'FB{getattr(fb,"id",0) if fb else 0}-FC{get_setting(db,"config_version","1")}-OB{(ob or {}).get("version",0)}-OC{get_setting(db,"opportunity_config_version","1")}'
        return {"version":version,"funds":{"fortress":{"fund_key":"fortress","portfolio_name":"RPD Fortress Fund","baseline_id":getattr(fb,"id",None) if fb else None,"report_date":getattr(fb,"report_date",None) if fb else None,"positions":clean(fdaily)},"opportunity":{"fund_key":"opportunity","portfolio_name":"RPD Opportunity Fund","baseline_id":None,"report_date":(ob or {}).get("report_date"),"positions":clean(odaily)}}}

@app.post("/api/yahoo-collector/snapshot-v2")
async def yahoo_collector_snapshot_v2(request:Request):
    collector_auth(request);data=await request.json()
    if not isinstance(data,dict) or not isinstance(data.get("funds"),dict): raise HTTPException(400,"Invalid multi-fund Yahoo snapshot.")
    with SessionLocal() as db:
        set_setting(db,"yahoo_collector_snapshot_v2",json.dumps(data));set_setting(db,"yahoo_collector_v2_updated_at",str(data.get("collected_at_utc") or utcnow().isoformat()));db.commit()
    return {"ok":True,"funds_received":list(data.get("funds",{}).keys())}

@app.get("/api/collector/ping")
def collector_ping(request:Request):
    collector_auth(request);return {"ok":True}

@app.get("/api/collector/config")
def collector_config(request:Request):
    collector_auth(request)
    with SessionLocal() as db:
        today=date.today().isoformat();positions=[]
        for p in db.query(Position).filter_by(active=True).order_by(Position.ticker,Position.expiry,Position.strike).all():
            if p.expiry<today:continue
            positions.append({"id":p.id,"ticker":p.ticker,"expiry":p.expiry,"option_type":p.option_type,"strike":p.strike,"quantity":p.quantity,"multiplier":p.multiplier,"bloomberg_security":p.bloomberg_security,"underlying_security":p.underlying_security})
        fdate=str(get_setting(db,"fortress_exposure_report_date","") or _daily_report_date(db,"fortress"))
        positions=_overlay_positions(positions,_active_trades(db,"fortress",fdate))
        confirmed=_confirmed_expiry_keys(db,"fortress",fdate)
        positions=[p for p in positions if _dr_position_key(p) not in confirmed]
        return {"version":int(get_setting(db,"config_version","1")),"updated_at":get_setting(db,"config_updated_at",""),"portfolio_name":get_setting(db,"portfolio_name","RPD Fortress Fund"),"nav_usd":float(get_setting(db,"nav_usd","0")),"limit_pct":float(get_setting(db,"limit_pct","15")),"positions":positions}

@app.post("/api/collector/snapshot")
async def collector_snapshot(request:Request):
    collector_auth(request);data=await request.json()
    with SessionLocal() as db:
        db.add(Snapshot(payload=json.dumps(data)));db.flush()
        count=db.query(Snapshot).count()
        if count>250:
            for x in db.query(Snapshot).order_by(Snapshot.id.asc()).limit(count-250).all():db.delete(x)
        db.commit()
    return {"ok":True}
