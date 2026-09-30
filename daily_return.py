from __future__ import annotations

import io
import json
import math
import re
import threading
import time
from datetime import datetime, timezone, time as dtime
from urllib.parse import urlencode, quote
from urllib.request import Request, urlopen
from urllib.error import HTTPError, URLError
from zoneinfo import ZoneInfo

from openpyxl import load_workbook
from pypdf import PdfReader

from helpers import parse_date, parse_nirvana_option_symbol

NY = ZoneInfo("America/New_York")
OPTION_RE = re.compile(
    r"^(?P<ticker>[A-Z0-9.\-]+)\s+(?P<expiry>\d{1,2}/\d{1,2}/\d{2})\s+(?P<strike>\d+(?:\.\d+)?)\s+(?P<type>PUT|CALL)$",
    re.I,
)

_spot_lock = threading.Lock()
_spot_cache = {"ts": 0.0, "symbols": (), "data": {}, "error": ""}
CACHE_SECONDS = 55


def _num(v):
    if v is None:
        return None
    s = str(v).strip()
    if not s:
        return None
    neg = s.startswith("(") and s.endswith(")")
    s = s.strip("()").replace(",", "").replace("$", "").strip()
    if s in {"", "-", "—"}:
        return None
    try:
        x = float(s)
    except Exception:
        return None
    if not math.isfinite(x):
        return None
    return -x if neg else x


def _positive(v):
    try:
        x = float(v)
        return x if math.isfinite(x) and x > 0 else None
    except Exception:
        return None


def _extract_layout_text(data: bytes) -> str:
    reader = PdfReader(io.BytesIO(data))
    pages = []
    for page in reader.pages:
        try:
            pages.append(page.extract_text(extraction_mode="layout") or "")
        except TypeError:
            pages.append(page.extract_text() or "")
    return "\n".join(pages)


def _column_slices(header: str):
    labels = [
        "Security Name", "Price", "Change", "%", "Quantity", "Unit Cost",
        "Market Value", "Day P&L", "MTD P&L", "YTD P&L", "ITD P&L",
    ]
    starts = []
    cursor = 0
    for lab in labels:
        idx = header.find(lab, cursor)
        if idx < 0:
            raise ValueError(f"PNL report column not found: {lab}")
        starts.append((lab, idx))
        cursor = idx + len(lab)
    slices = {}
    for i, (lab, start) in enumerate(starts):
        end = starts[i + 1][1] if i + 1 < len(starts) else None
        slices[lab] = (start, end)
    return slices


def _cell(line: str, col_slices, name: str) -> str:
    start, end = col_slices[name]
    if len(line) <= start:
        return ""
    return line[start:end].strip() if end is not None else line[start:].strip()


def parse_nirvana_pnl_pdf(data: bytes) -> dict:
    """Parse Nirvana PNL PDF as the previous-close fund/NAV baseline."""
    text = _extract_layout_text(data)
    lines = text.splitlines()
    if not any("PNL Report" in x for x in lines[:20]):
        raise ValueError("This does not look like a Nirvana PNL Report PDF.")

    report_date = None
    run_date = None
    for line in lines[:30]:
        m = re.search(r"Report Date:\s*(\d{1,2}/\d{1,2}/\d{4})", line, re.I)
        if m:
            report_date = datetime.strptime(m.group(1), "%m/%d/%Y").date().isoformat()
        m2 = re.search(r"Run Date:\s*([^\n]+)$", line, re.I)
        if m2:
            run_date = m2.group(1).strip()

    header_idx = None
    for i, line in enumerate(lines):
        if "Security Name" in line and "Market Value" in line and "Day P&L" in line:
            header_idx = i
            break
    if header_idx is None:
        raise ValueError("Could not locate the PNL table header.")
    cols = _column_slices(lines[header_idx])

    positions = []
    nav = None
    ignored = []
    for line in lines[header_idx + 1:]:
        if not line.strip():
            continue
        name = _cell(line, cols, "Security Name")
        if not name or name.startswith("Page "):
            continue
        if name.startswith("Grand Total"):
            gm = re.search(r"Grand Total\s*:\s*([0-9,]+(?:\.[0-9]+)?)", line, re.I)
            nav = _num(gm.group(1)) if gm else _num(_cell(line, cols, "Market Value"))
            continue
        if name.upper() == "CLOSED":
            continue

        price = _num(_cell(line, cols, "Price"))
        qty = _num(_cell(line, cols, "Quantity"))
        market_value = _num(_cell(line, cols, "Market Value"))
        if qty is None or market_value is None:
            ignored.append(name)
            continue

        upper = name.upper().strip()
        if upper in {"FGTXX", "USD", "CASH"}:
            continue

        om = OPTION_RE.match(upper)
        if om:
            expiry = datetime.strptime(om.group("expiry"), "%m/%d/%y").date().isoformat()
            typ = "P" if om.group("type").upper() == "PUT" else "C"
            strike = float(om.group("strike"))
            multiplier = 100.0
            baseline_mark = (market_value / (qty * multiplier)) if qty else price
            if baseline_mark is None or baseline_mark <= 0:
                baseline_mark = price
            positions.append({
                "security_name": name.strip(),
                "instrument_type": "OPTION",
                "ticker": om.group("ticker").upper(),
                "expiry": expiry,
                "option_type": typ,
                "strike": strike,
                "quantity": qty,
                "multiplier": multiplier,
                "baseline_price": baseline_mark,
                "display_price": price,
                "baseline_market_value": market_value,
            })
            continue

        ticker = upper.split()[0]
        if re.fullmatch(r"[A-Z][A-Z0-9.\-]{0,14}", ticker):
            multiplier = 1.0
            baseline_mark = (market_value / qty) if qty else price
            positions.append({
                "security_name": name.strip(),
                "instrument_type": "EQUITY",
                "ticker": ticker,
                "expiry": "",
                "option_type": "",
                "strike": None,
                "quantity": qty,
                "multiplier": multiplier,
                "baseline_price": baseline_mark,
                "display_price": price,
                "baseline_market_value": market_value,
            })
        else:
            ignored.append(name)

    if nav is None or nav <= 0:
        raise ValueError("Could not read Grand Total / NAV from the PNL report.")
    if not positions:
        raise ValueError("No live option/equity positions were parsed from the PNL report.")

    return {
        "report_date": report_date or "",
        "run_date": run_date or "",
        "nav_usd": float(nav),
        "positions": positions,
        "ignored": ignored,
    }


def _norm_header(v):
    return re.sub(r"[^a-z0-9]+", " ", str(v or "").strip().lower()).strip()


def parse_nirvana_exposure_xlsx(data: bytes) -> dict:
    """Parse Nirvana Exposure by Underlying XLSX for Daily Return only."""
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    ws = wb.active
    vals = list(ws.iter_rows(values_only=True))
    if not vals:
        raise ValueError("Exposure workbook is empty.")

    header_idx = None
    for i, row in enumerate(vals[:30]):
        n = {_norm_header(x) for x in row if x is not None}
        if {"symbol", "expiration date", "position", "delta", "delta adjusted position", "net exposure"}.issubset(n):
            header_idx = i
            break
    if header_idx is None:
        raise ValueError("Could not locate the Nirvana Exposure Summary header.")

    headers = [str(x).strip() if x is not None else "" for x in vals[header_idx]]
    idx = {_norm_header(h): i for i, h in enumerate(headers)}
    required = ["symbol", "position", "delta", "delta adjusted position", "net exposure"]
    for col in required:
        if col not in idx:
            raise ValueError(f"Exposure report column not found: {col}")

    preamble = "\n".join(" ".join(str(x) for x in r if x is not None) for r in vals[:header_idx])
    m = re.search(r"Report Date:\s*([0-9/\-]+)", preamble, re.I)
    report_date = parse_date(m.group(1)) if m else None

    positions = []
    nav = None
    for row in vals[header_idx + 1:]:
        if not any(x is not None and str(x).strip() for x in row):
            continue
        def val(name):
            j = idx.get(name)
            return row[j] if j is not None and j < len(row) else None
        sym = str(val("symbol") or "").strip()
        if sym.lower().startswith("grand total"):
            nav = _num(val("nav")) if "nav" in idx else None
            continue
        if not sym.startswith("O:"):
            continue
        parsed = parse_nirvana_option_symbol(sym)
        if not parsed:
            continue
        qty = _num(val("position"))
        delta = _num(val("delta"))
        dap = _num(val("delta adjusted position"))
        net = _num(val("net exposure"))
        net_pct = _num(val("net exposure %")) if "net exposure %" in idx else None
        if qty is None or delta is None or dap is None or net is None:
            continue
        baseline_spot = abs(net / dap) if abs(dap) > 1e-12 else None
        if baseline_spot is None or not math.isfinite(baseline_spot) or baseline_spot <= 0:
            continue
        positions.append({
            **parsed,
            "quantity": qty,
            "delta": delta,
            "delta_adjusted_position": dap,
            "net_exposure": net,
            "net_exposure_pct": net_pct,
            "baseline_underlying_price": baseline_spot,
            "source_symbol": sym,
        })

    if not report_date:
        raise ValueError("Could not read Report Date from the Exposure report.")
    if not positions:
        raise ValueError("No option rows with usable Delta / Delta Adjusted Position / Net Exposure were found.")
    return {"report_date": report_date, "nav_usd": nav, "positions": positions}


def _http_json(url: str, timeout: int = 10):
    req = Request(url, headers={
        "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 Chrome/142 Safari/537.36",
        "Accept": "application/json,text/plain,*/*",
        "Accept-Language": "en-US,en;q=0.9",
        "Connection": "close",
    })
    try:
        with urlopen(req, timeout=timeout) as r:
            return json.loads(r.read().decode("utf-8")), ""
    except HTTPError as e:
        if e.code == 429:
            return None, "Yahoo rate limited (HTTP 429)"
        return None, f"Yahoo HTTP {e.code}"
    except URLError as e:
        return None, f"Yahoo network error: {getattr(e, 'reason', e)}"
    except Exception as e:
        return None, f"Yahoo error: {e}"


def _last_number(values):
    if not isinstance(values, list):
        return None
    for v in reversed(values):
        x = _positive(v)
        if x is not None:
            return x
    return None


def _parse_spark(payload, wanted):
    out = {}
    try:
        results = payload.get("spark", {}).get("result", [])
        for item in results:
            symbol = str(item.get("symbol") or "").upper()
            responses = item.get("response") or []
            if not responses:
                continue
            r = responses[0] or {}
            meta = r.get("meta") or {}
            px = _positive(meta.get("regularMarketPrice"))
            source = "Yahoo regularMarketPrice"
            if px is None:
                quote_block = ((r.get("indicators") or {}).get("quote") or [{}])[0]
                px = _last_number(quote_block.get("close"))
                source = "Yahoo 1m close"
            if symbol in wanted and px is not None:
                out[symbol] = (px, source, "")
    except Exception:
        pass
    return out


def _chart_quote(symbol):
    enc = quote(symbol, safe="")
    last_err = "Yahoo price unavailable"
    for host in ("query1.finance.yahoo.com", "query2.finance.yahoo.com"):
        url = f"https://{host}/v8/finance/chart/{enc}?interval=1m&range=1d&includePrePost=false"
        payload, err = _http_json(url)
        if err:
            last_err = err
            if "429" in err or "rate limited" in err.lower():
                return None, "UNAVAILABLE", err
            continue
        try:
            results = payload.get("chart", {}).get("result") or []
            if not results:
                last_err = "Yahoo chart returned no result"
                continue
            r = results[0]
            meta = r.get("meta") or {}
            px = _positive(meta.get("regularMarketPrice"))
            if px is not None:
                return px, "Yahoo regularMarketPrice", ""
            quote_block = ((r.get("indicators") or {}).get("quote") or [{}])[0]
            px = _last_number(quote_block.get("close"))
            if px is not None:
                return px, "Yahoo 1m close", ""
        except Exception as e:
            last_err = f"Yahoo chart parse error: {e}"
    return None, "UNAVAILABLE", last_err


def yahoo_underlying_spots(tickers):
    """Fetch only underlying stock prices; no Yahoo option-chain calls."""
    symbols = tuple(sorted({str(x).upper().strip() for x in tickers if str(x).strip()}))
    now = time.time()
    with _spot_lock:
        if _spot_cache["symbols"] == symbols and now - _spot_cache["ts"] < CACHE_SECONDS:
            return dict(_spot_cache["data"]), _spot_cache["error"]

    data = {}
    global_err = ""
    if symbols:
        qs = urlencode({
            "symbols": ",".join(symbols),
            "range": "1d",
            "interval": "1m",
            "indicators": "close",
            "includePrePost": "false",
        })
        spark_rate_limited = False
        for host in ("query1.finance.yahoo.com", "query2.finance.yahoo.com"):
            payload, err = _http_json(f"https://{host}/v7/finance/spark?{qs}")
            if err:
                global_err = err
                if "429" in err or "rate limited" in err.lower():
                    spark_rate_limited = True
                continue
            data.update(_parse_spark(payload or {}, set(symbols)))
            if data:
                break

        if not spark_rate_limited:
            for symbol in symbols:
                if symbol not in data:
                    px, src, err = _chart_quote(symbol)
                    data[symbol] = (px, src, err)
        else:
            for symbol in symbols:
                data.setdefault(symbol, (None, "UNAVAILABLE", global_err or "Yahoo rate limited"))

    with _spot_lock:
        _spot_cache.update({"ts": now, "symbols": symbols, "data": dict(data), "error": global_err})
    return data, global_err


def _key(p):
    return (
        str(p.get("ticker") or "").upper(),
        str(p.get("expiry") or ""),
        str(p.get("option_type") or "").upper(),
        round(float(p.get("strike") or 0.0), 6),
    )


def calculate_delta_estimated_return(pnl_positions: list[dict], exposure_positions: list[dict], nav_usd: float) -> dict:
    """First-order, delta-based intraday return estimate using Yahoo underlying prices."""
    exposure_map = {_key(p): p for p in exposure_positions}
    tickers = {p.get("ticker") for p in pnl_positions if p.get("ticker")}
    spots, yahoo_error = yahoo_underlying_spots(tickers)

    rows = []
    total_pnl = 0.0
    valid = 0
    quote_tickers_valid = {t for t, v in spots.items() if v and v[0] is not None}

    for p in pnl_positions:
        ticker = str(p.get("ticker") or "").upper()
        current_spot, source, quote_note = spots.get(ticker, (None, "UNAVAILABLE", "Yahoo quote not returned"))
        pnl = None
        contrib = None
        delta = None
        dap = None
        baseline_spot = None
        match_note = ""

        if p.get("instrument_type") == "OPTION":
            e = exposure_map.get(_key(p))
            if e is None:
                match_note = "No same-date Exposure baseline match"
            else:
                delta = e.get("delta")
                dap = e.get("delta_adjusted_position")
                baseline_spot = e.get("baseline_underlying_price")
                if current_spot is not None and dap is not None and baseline_spot is not None:
                    pnl = float(dap) * (float(current_spot) - float(baseline_spot))
        else:
            baseline_spot = p.get("baseline_price")
            dap = p.get("quantity")
            if current_spot is not None and baseline_spot is not None:
                pnl = float(p.get("quantity") or 0.0) * (float(current_spot) - float(baseline_spot))

        if pnl is not None:
            contrib = pnl / nav_usd * 100.0 if nav_usd else None
            total_pnl += pnl
            valid += 1

        note_parts = [x for x in [match_note, quote_note] if x]
        rows.append({
            **p,
            "nirvana_delta": delta,
            "delta_adjusted_position": dap,
            "baseline_underlying_price": baseline_spot,
            "current_underlying_price": current_spot,
            "price_source": source,
            "note": " · ".join(note_parts),
            "estimated_pnl": pnl,
            "contribution_pct": contrib,
        })

    total = len(pnl_positions)
    return {
        "rows": rows,
        "estimated_pnl": total_pnl,
        "estimated_return_pct": (total_pnl / nav_usd * 100.0) if nav_usd else None,
        "coverage_valid": valid,
        "coverage_total": total,
        "ticker_coverage_valid": len(quote_tickers_valid),
        "ticker_coverage_total": len(tickers),
        "complete": bool(total and valid == total),
        "yahoo_error": yahoo_error,
        "calculated_at_utc": datetime.now(timezone.utc).isoformat(),
    }


def auto_refresh_allowed(now_utc=None):
    now_utc = now_utc or datetime.now(timezone.utc)
    ny = now_utc.astimezone(NY)
    t = ny.time().replace(tzinfo=None)
    return ny.weekday() < 5 and dtime(9, 30) <= t < dtime(16, 15)


def ny_time_label(timestamp_utc: str):
    try:
        ts = datetime.fromisoformat(str(timestamp_utc).replace("Z", "+00:00"))
        return ts.astimezone(NY).strftime("%d %b %Y, %H:%M:%S %Z")
    except Exception:
        return str(timestamp_utc or "")
